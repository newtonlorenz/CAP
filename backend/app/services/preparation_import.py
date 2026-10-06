"""Read spreadsheet form text without persisting uploads or evaluating formulas."""

import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.parsers import expat

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 50
MAX_CELL_CHARS = 10000
MAX_PREVIEW_CHARS = 2_000_000
MAX_ENTRIES = 512


def invalid(message: str) -> None:
    raise HTTPException(422, message)


def check_xml(data: bytes, worksheet: bool) -> None:
    # Validate before openpyxl parses any XML: its optional defusedxml dependency
    # is not required by this project. Reject DTDs in every supported encoding.
    parser = expat.ParserCreate(namespace_separator="}")
    row_count = 0
    cell_count = 0

    def reject_doctype(*_args):
        invalid("Workbook XML must not contain a document type or entities")

    def check_element(name, attributes):
        nonlocal row_count, cell_count
        if not worksheet:
            return
        local = name.rsplit("}", 1)[-1]
        if local == "row":
            row_count += 1
            cell_count = 0
            if row_count > MAX_ROWS:
                invalid(f"Workbook exceeds the {MAX_ROWS:,} row limit per sheet")
        if local == "c":
            cell_count += 1
            if cell_count > MAX_COLUMNS:
                invalid(f"Workbook exceeds the {MAX_COLUMNS} column limit")
        if local == "row" and "r" in attributes and not 1 <= int(attributes["r"]) <= MAX_ROWS:
            invalid(f"Workbook exceeds the {MAX_ROWS:,} row limit per sheet")
        if local == "c" and "r" in attributes:
            match = re.fullmatch(r"([A-Za-z]{1,3})([1-9][0-9]{0,6})", attributes["r"])
            if (
                match is None
                or column_index_from_string(match[1]) > MAX_COLUMNS
                or int(match[2]) > MAX_ROWS
            ):
                invalid(f"Workbook exceeds the {MAX_ROWS:,} row or {MAX_COLUMNS} column limit")

    parser.StartDoctypeDeclHandler = reject_doctype
    parser.StartElementHandler = check_element
    parser.Parse(data, True)


def check_archive(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ENTRIES or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
            invalid("Workbook is too large when expanded (20 MiB limit)")
        for member in members:
            if member.flag_bits & 1:
                invalid("Encrypted workbooks are not supported")
            if member.filename.lower().endswith((".xml", ".rels")):
                check_xml(archive.read(member), member.filename.startswith("xl/worksheets/"))


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    result = str(value)
    if len(result) > MAX_CELL_CHARS:
        invalid(f"A cell exceeds the {MAX_CELL_CHARS:,} character limit")
    return result


def checked_rows(source) -> list[list[str]]:
    rows = []
    character_count = 0
    for row in source:
        if len(rows) >= MAX_ROWS:
            invalid(f"Spreadsheet exceeds the {MAX_ROWS:,} row limit")
        if len(row) > MAX_COLUMNS:
            invalid(f"Spreadsheet exceeds the {MAX_COLUMNS} column limit")
        values = [cell_text(value) for value in row]
        character_count += sum(len(value) for value in values)
        if character_count > MAX_PREVIEW_CHARS:
            invalid("Spreadsheet text is too large to preview (2 million character limit)")
        rows.append(values)
    # Trim only trailing empty space; retain interior rows/cells and all text.
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        invalid("The selected sheet contains no text")
    width = max((i + 1 for row in rows for i, value in enumerate(row) if value), default=0)
    return [row[:width] + [""] * max(0, width - len(row)) for row in rows]


def preview_spreadsheet(content: bytes, filename: str, sheet_name: str | None) -> dict:
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Spreadsheet exceeds the 5 MiB upload limit")
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv", ".tsv"}:
        invalid("Choose an .xlsx, .csv or .tsv file; save older .xls workbooks as .xlsx first")
    if not content:
        invalid("The spreadsheet is empty")
    try:
        if suffix != ".xlsx":
            text = content.decode("utf-8-sig")
            delimiter = "\t" if suffix == ".tsv" else ","
            if suffix == ".csv":
                try:
                    delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
                except csv.Error:
                    pass
            name = "Data"
            if sheet_name and sheet_name != name:
                invalid("Selected sheet was not found")
            rows = checked_rows(
                csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            )
            return {"sheets": [name], "sheet_name": name, "rows": rows, "warnings": []}

        check_archive(content)
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheets = [sheet.title for sheet in workbook.worksheets]
            if not sheets:
                invalid("The workbook has no worksheets")
            selected = sheet_name or sheets[0]
            if selected not in sheets:
                invalid("Selected sheet was not found")
            sheet = workbook[selected]
            # Ignore inaccurate/formatting-only dimension metadata. Actual cell
            # coordinates are bounded by the XML check before row allocation.
            sheet.reset_dimensions()
            formula_count = 0

            def values():
                nonlocal formula_count
                for row in sheet.iter_rows():
                    result = []
                    for cell in row:
                        if cell.data_type == "f":
                            formula_count += 1
                            result.append("")
                        else:
                            result.append(cell.value)
                    yield result

            rows = checked_rows(values())
            warnings = []
            if formula_count:
                warnings.append(
                    f"Skipped {formula_count} formula cell(s). Paste their displayed text if needed."
                )
            warnings.append(
                "Preview includes cell text only. Checkboxes, images and page layout must be recreated in the form editor."
            )
            return {"sheets": sheets, "sheet_name": selected, "rows": rows, "warnings": warnings}
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (UnicodeError, csv.Error):
        invalid("Could not read this text file. Save it as UTF-8 CSV or TSV and try again")
    except Exception:  # noqa: BLE001 - untrusted workbooks raise varied parser/library errors
        # Do not disclose parser internals or workbook content in an API error.
        invalid("Could not read this workbook. Save a valid .xlsx file and try again")
