import json
import csv
import io
import html as html_module
import re
import textwrap
import uuid
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Callable, Iterable, Optional, Sequence

try:
    from openpyxl import Workbook  # type: ignore
except ModuleNotFoundError:  # pragma: no cover
    Workbook = None  # type: ignore
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from fastapi import HTTPException
from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import aliased

from app.config import settings
from app.models.audit import AuditLog
from app.models.change_management import (
    ChangeEntry,
    ChangeEntryComponent,
    ChangeEvent,
    Component,
    ComponentRegister,
    IntegrationCheck,
)
from app.models.document import Document
from app.models.jurisdiction import Jurisdiction
from app.models.requirement import Requirement, RequirementStatus
from app.models.review import ReviewCycle, ReviewItem, ReviewItemEvidenceFile
from app.models.user import User
from app.services.review_assurance import (
    assessment_rows,
    assessment_gap_reasons,
    frozen_payload,
    hydrate,
)
from app.services.requirement_scope import current_requirement_condition
from app.services.review_gap_report import render_review_gap_pdf, summary_metrics, owner_actions
from app.services.review_progress import ACTIONABLE_REQUIREMENT_TYPES, COMPLETED_REVIEW_STATUSES


def _spreadsheet_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class _SafeCsvWriter:
    def __init__(self, output):
        self.writer = csv.writer(output)

    def writerow(self, values):
        return self.writer.writerow([_spreadsheet_cell(v) for v in values])

    def writerows(self, rows):
        for row in rows:
            self.writerow(row)


def _append_excel(sheet, values):
    sheet.append([_spreadsheet_cell(value) for value in values])


def _reference_sort_key(reference_id: str) -> list[tuple[int, object]]:
    parts: list[tuple[int, object]] = []
    for part in reference_id.split("."):
        if part.isdigit():
            parts.append((0, int(part)))
        else:
            parts.append((1, part.lower()))
    return parts


def _strip_html(value: str) -> str:
    if not value:
        return ""
    cleaned = re.sub(r"<[^>]+>", " ", value)
    return html_module.unescape(cleaned).strip()


def _stringify(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _join_values(values: Iterable[str]) -> str:
    parts = [value if value is not None else "" for value in values]
    return " ; ".join(parts) if parts else ""


def _latest_requirement_status_subquery():
    return select(
        RequirementStatus.requirement_id.label("requirement_id"),
        RequirementStatus.status.label("status"),
        RequirementStatus.assigned_to.label("assigned_to"),
        RequirementStatus.comment.label("comment"),
        func.row_number()
        .over(
            partition_by=RequirementStatus.requirement_id,
            order_by=(RequirementStatus.changed_at.desc(), RequirementStatus.id.desc()),
        )
        .label("rn"),
    ).subquery()


def _normalize_status(status: Optional[str]) -> str:
    return status or "not_started"


def _is_overdue(deadline: Optional[datetime]) -> bool:
    if deadline is None:
        return False
    if deadline.tzinfo is None:
        return deadline < datetime.now(timezone.utc).replace(tzinfo=None)
    return deadline < datetime.now(deadline.tzinfo)


def _reportable_requirement_condition(status_expr):
    return and_(
        Requirement.requirement_type.in_(ACTIONABLE_REQUIREMENT_TYPES),
        or_(status_expr.is_(None), status_expr != "not_applicable"),
    )


def _build_evidence_download_url(
    review_cycle_id: uuid.UUID,
    review_item_id: uuid.UUID,
    file_id: uuid.UUID,
) -> str:
    path = (
        f"/api/v1/review-cycles/{review_cycle_id}/items/{review_item_id}/files/{file_id}/download"
    )
    base_url = settings.frontend_base_url.rstrip("/")
    if base_url:
        return f"{base_url}{path}"
    return path


def _jurisdiction_header_text(jurisdiction: Jurisdiction) -> str:
    """Return the best available header text for reports."""
    if jurisdiction.report_header_text and jurisdiction.report_header_text.strip():
        return jurisdiction.report_header_text.strip()
    if jurisdiction.regulator_name and jurisdiction.regulator_name.strip():
        return jurisdiction.regulator_name.strip()
    return jurisdiction.name.strip()


def _draw_header_footer_factory(header_text: str):
    def _draw_header_footer(canvas, doc):
        canvas.saveState()
        width, height = A4

        canvas.setFont("Helvetica", 8)
        header_y = height - 0.35 * inch
        if header_text:
            canvas.drawString(
                doc.leftMargin, header_y, f"CAP readiness assessment | Regulator: {header_text}"
            )
        canvas.setLineWidth(0.5)
        canvas.line(doc.leftMargin, header_y - 6, width - doc.rightMargin, header_y - 6)

        footer_y = 0.35 * inch
        canvas.line(doc.leftMargin, footer_y + 8, width - doc.rightMargin, footer_y + 8)
        page_label = f"Page {canvas.getPageNumber()}"
        canvas.drawRightString(width - doc.rightMargin, footer_y, page_label)
        canvas.restoreState()

    return _draw_header_footer


async def _get_jurisdiction_or_404(db: AsyncSession, jurisdiction_id: uuid.UUID) -> Jurisdiction:
    result = await db.execute(select(Jurisdiction).where(Jurisdiction.id == jurisdiction_id))
    jurisdiction = result.scalar_one_or_none()
    if jurisdiction is None:
        raise HTTPException(status_code=404, detail="Jurisdiction not found")
    return jurisdiction


async def _resolve_report_jurisdiction(
    db: AsyncSession,
    organization_id: Optional[uuid.UUID],
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
) -> Jurisdiction:
    if review_cycle_id is not None:
        query = (
            select(Jurisdiction)
            .join(ReviewCycle, ReviewCycle.jurisdiction_id == Jurisdiction.id)
            .where(ReviewCycle.id == review_cycle_id)
        )
        if organization_id is None:
            query = query.where(ReviewCycle.organization_id.is_(None))
        else:
            query = query.where(ReviewCycle.organization_id == organization_id)
        result = await db.execute(query)
        jurisdiction = result.scalar_one_or_none()
        if jurisdiction is None:
            raise HTTPException(status_code=404, detail="Review cycle jurisdiction not found")
        cycle = await db.get(ReviewCycle, review_cycle_id)
        frozen = await frozen_payload(db, cycle)
        return hydrate(Jurisdiction, frozen["jurisdiction"]) if frozen else jurisdiction
    if jurisdiction_id is None:
        raise HTTPException(status_code=400, detail="jurisdiction_id is required")
    return await _get_jurisdiction_or_404(db, jurisdiction_id)


@dataclass(frozen=True)
class RequirementStatusSnapshot:
    status: Optional[str]
    assigned_to: Optional[uuid.UUID]


@dataclass(frozen=True)
class EvidenceAggregate:
    file_ids: list[str]
    filenames: list[str]
    descriptions: list[str]
    uploaded_by: list[str]
    uploaded_at: list[str]
    links: list[str]


@dataclass(frozen=True)
class SoAContext:
    jurisdiction: Jurisdiction
    cycle: ReviewCycle
    item: ReviewItem
    requirement: Requirement
    document: Optional[Document]
    requirement_status: RequirementStatusSnapshot
    evidence: EvidenceAggregate
    reviewer_display: str = ""
    responsible_display: str = ""
    baseline_version: str = ""
    source_page: str = ""
    assessment_rationale: str = ""


@dataclass(frozen=True)
class SoAField:
    key: str
    label: str
    getter: Callable[[SoAContext], str]


def _build_evidence_aggregate(
    cycle_id: uuid.UUID,
    item_id: uuid.UUID,
    files: list[ReviewItemEvidenceFile],
) -> EvidenceAggregate:
    file_ids: list[str] = []
    filenames: list[str] = []
    descriptions: list[str] = []
    uploaded_by: list[str] = []
    uploaded_at: list[str] = []
    links: list[str] = []
    for evidence_file in files:
        file_ids.append(_stringify(evidence_file.id))
        filenames.append(evidence_file.filename)
        descriptions.append(evidence_file.description or "")
        uploaded_by.append(_stringify(evidence_file.uploaded_by))
        uploaded_at.append(_stringify(evidence_file.uploaded_at))
        link = _build_evidence_download_url(cycle_id, item_id, evidence_file.id)
        links.append(f"{evidence_file.filename}: {link}")
    return EvidenceAggregate(
        file_ids=file_ids,
        filenames=filenames,
        descriptions=descriptions,
        uploaded_by=uploaded_by,
        uploaded_at=uploaded_at,
        links=links,
    )


SOA_FIELDS: list[SoAField] = [
    SoAField("jurisdiction_code", "Jurisdiction Code", lambda ctx: ctx.jurisdiction.code),
    SoAField("jurisdiction_name", "Jurisdiction Name", lambda ctx: ctx.jurisdiction.name),
    SoAField(
        "jurisdiction_regulator_name",
        "Jurisdiction Regulator Name",
        lambda ctx: ctx.jurisdiction.regulator_name or "",
    ),
    SoAField("review_cycle_id", "Review Cycle ID", lambda ctx: _stringify(ctx.cycle.id)),
    SoAField("review_cycle_name", "Review Cycle Name", lambda ctx: ctx.cycle.name),
    SoAField(
        "review_cycle_description",
        "Review Cycle Description",
        lambda ctx: ctx.cycle.description or "",
    ),
    SoAField("review_cycle_scope", "Review Cycle Scope", lambda ctx: ctx.cycle.scope),
    SoAField(
        "review_cycle_scope_filter",
        "Review Cycle Scope Filter",
        lambda ctx: ctx.cycle.scope_filter or "",
    ),
    SoAField(
        "review_cycle_deadline",
        "Review Cycle Deadline",
        lambda ctx: _stringify(ctx.cycle.deadline),
    ),
    SoAField("review_cycle_status", "Review Cycle Status", lambda ctx: ctx.cycle.status),
    SoAField(
        "review_cycle_created_by",
        "Review Cycle Created By",
        lambda ctx: _stringify(ctx.cycle.created_by),
    ),
    SoAField(
        "review_cycle_closed_at",
        "Review Cycle Closed At",
        lambda ctx: _stringify(ctx.cycle.closed_at),
    ),
    SoAField(
        "review_cycle_closed_by",
        "Review Cycle Closed By",
        lambda ctx: _stringify(ctx.cycle.closed_by),
    ),
    SoAField(
        "review_cycle_snapshot_id",
        "Review Cycle Snapshot ID",
        lambda ctx: _stringify(ctx.cycle.snapshot_id),
    ),
    SoAField(
        "review_cycle_created_at",
        "Review Cycle Created At",
        lambda ctx: _stringify(ctx.cycle.created_at),
    ),
    SoAField("review_item_id", "Review Item ID", lambda ctx: _stringify(ctx.item.id)),
    SoAField(
        "review_item_review_status",
        "Review Item Review Status",
        lambda ctx: ctx.item.review_status,
    ),
    SoAField(
        "review_item_assigned_reviewer_id",
        "Review Item Assigned Reviewer ID",
        lambda ctx: _stringify(ctx.item.assigned_reviewer_id),
    ),
    SoAField(
        "review_item_reviewer_id",
        "Review Item Reviewer ID",
        lambda ctx: _stringify(ctx.item.reviewer_id),
    ),
    SoAField(
        "review_item_review_evidence",
        "Review Item Review Evidence",
        lambda ctx: _strip_html(ctx.item.review_evidence or ctx.assessment_rationale or ""),
    ),
    SoAField(
        "review_item_jira_issue_key",
        "Review Item Jira Issue Key",
        lambda ctx: ctx.item.jira_issue_key or "",
    ),
    SoAField(
        "review_item_jira_issue_url",
        "Review Item Jira Issue URL",
        lambda ctx: ctx.item.jira_issue_url or "",
    ),
    SoAField(
        "review_item_jira_status",
        "Review Item Jira Status",
        lambda ctx: ctx.item.jira_status or "",
    ),
    SoAField(
        "review_item_jira_summary",
        "Review Item Jira Summary",
        lambda ctx: ctx.item.jira_summary or "",
    ),
    SoAField(
        "review_item_jira_assignee",
        "Review Item Jira Assignee",
        lambda ctx: ctx.item.jira_assignee or "",
    ),
    SoAField(
        "review_item_jira_priority",
        "Review Item Jira Priority",
        lambda ctx: ctx.item.jira_priority or "",
    ),
    SoAField(
        "review_item_jira_updated_at",
        "Review Item Jira Updated At",
        lambda ctx: _stringify(ctx.item.jira_updated_at),
    ),
    SoAField(
        "review_item_jira_synced_at",
        "Review Item Jira Synced At",
        lambda ctx: _stringify(ctx.item.jira_synced_at),
    ),
    SoAField(
        "review_item_jira_sync_error",
        "Review Item Jira Sync Error",
        lambda ctx: ctx.item.jira_sync_error or "",
    ),
    SoAField(
        "review_item_reviewed_at",
        "Review Item Reviewed At",
        lambda ctx: _stringify(ctx.item.reviewed_at),
    ),
    SoAField(
        "review_item_created_at",
        "Review Item Created At",
        lambda ctx: _stringify(ctx.item.created_at),
    ),
    SoAField("requirement_id", "Requirement ID", lambda ctx: _stringify(ctx.requirement.id)),
    SoAField(
        "requirement_document_id",
        "Requirement Document ID",
        lambda ctx: _stringify(ctx.requirement.document_id),
    ),
    SoAField(
        "requirement_reference_id",
        "Requirement Reference ID",
        lambda ctx: ctx.requirement.reference_id,
    ),
    SoAField(
        "requirement_title",
        "Requirement Title",
        lambda ctx: ctx.requirement.title or "",
    ),
    SoAField(
        "requirement_text",
        "Requirement Text",
        lambda ctx: _strip_html(ctx.requirement.text),
    ),
    SoAField(
        "requirement_type",
        "Requirement Type",
        lambda ctx: ctx.requirement.requirement_type,
    ),
    SoAField(
        "requirement_parent_id",
        "Requirement Parent ID",
        lambda ctx: _stringify(ctx.requirement.parent_id),
    ),
    SoAField(
        "requirement_default_owner_id",
        "Requirement Default Owner ID",
        lambda ctx: _stringify(ctx.requirement.default_owner_id),
    ),
    SoAField(
        "requirement_active",
        "Requirement Active",
        lambda ctx: _stringify(ctx.requirement.active),
    ),
    SoAField(
        "requirement_version",
        "Requirement Version",
        lambda ctx: _stringify(ctx.requirement.version),
    ),
    SoAField(
        "requirement_sort_order",
        "Requirement Sort Order",
        lambda ctx: _stringify(ctx.requirement.sort_order),
    ),
    SoAField(
        "requirement_created_at",
        "Requirement Created At",
        lambda ctx: _stringify(ctx.requirement.created_at),
    ),
    SoAField(
        "requirement_current_status",
        "Requirement Current Status",
        lambda ctx: _normalize_status(ctx.requirement_status.status),
    ),
    SoAField(
        "requirement_current_assigned_to",
        "Requirement Current Assigned To",
        lambda ctx: _stringify(ctx.requirement_status.assigned_to),
    ),
    SoAField(
        "document_id",
        "Document ID",
        lambda ctx: _stringify(ctx.document.id) if ctx.document else "",
    ),
    SoAField(
        "document_filename",
        "Document Filename",
        lambda ctx: ctx.document.filename if ctx.document else "",
    ),
    SoAField(
        "document_name",
        "Document Name",
        lambda ctx: ctx.document.name if ctx.document else "",
    ),
    SoAField(
        "document_type",
        "Document Type",
        lambda ctx: ctx.document.document_type if ctx.document else "",
    ),
    SoAField(
        "document_version",
        "Document Version",
        lambda ctx: ctx.document.version if ctx.document else "",
    ),
    SoAField(
        "document_effective_date",
        "Document Effective Date",
        lambda ctx: _stringify(ctx.document.effective_date) if ctx.document else "",
    ),
    SoAField(
        "document_status",
        "Document Status",
        lambda ctx: ctx.document.status if ctx.document else "",
    ),
    SoAField(
        "document_testing_frequency",
        "Document Testing Frequency",
        lambda ctx: ctx.document.testing_frequency if ctx.document else "",
    ),
    SoAField(
        "document_uploaded_by",
        "Document Uploaded By",
        lambda ctx: _stringify(ctx.document.uploaded_by) if ctx.document else "",
    ),
    SoAField(
        "document_approved_by",
        "Document Approved By",
        lambda ctx: _stringify(ctx.document.approved_by) if ctx.document else "",
    ),
    SoAField(
        "document_current_extraction_id",
        "Document Current Extraction ID",
        lambda ctx: _stringify(ctx.document.current_extraction_id) if ctx.document else "",
    ),
    SoAField(
        "document_archived_at",
        "Document Archived At",
        lambda ctx: _stringify(ctx.document.archived_at) if ctx.document else "",
    ),
    SoAField(
        "document_created_at",
        "Document Created At",
        lambda ctx: _stringify(ctx.document.created_at) if ctx.document else "",
    ),
    SoAField(
        "evidence_file_ids",
        "Evidence File IDs",
        lambda ctx: _join_values(ctx.evidence.file_ids),
    ),
    SoAField(
        "evidence_filenames",
        "Evidence Filenames",
        lambda ctx: _join_values(ctx.evidence.filenames),
    ),
    SoAField(
        "evidence_file_descriptions",
        "Evidence File Descriptions",
        lambda ctx: _join_values(ctx.evidence.descriptions),
    ),
    SoAField(
        "evidence_uploaded_by",
        "Evidence Uploaded By",
        lambda ctx: _join_values(ctx.evidence.uploaded_by),
    ),
    SoAField(
        "evidence_uploaded_at",
        "Evidence Uploaded At",
        lambda ctx: _join_values(ctx.evidence.uploaded_at),
    ),
    SoAField(
        "evidence_file_links",
        "Evidence File Links",
        lambda ctx: _join_values(ctx.evidence.links),
    ),
]

SOA_FIELDS.extend(
    [
        SoAField("reviewer_name", "Reviewed By", lambda ctx: ctx.reviewer_display),
        SoAField("responsible_name", "Responsible Person", lambda ctx: ctx.responsible_display),
        SoAField("baseline_version", "Approved Baseline Version", lambda ctx: ctx.baseline_version),
        SoAField("source_page", "Source PDF Page", lambda ctx: ctx.source_page),
    ]
)
READABLE_SOA_FIELDS = [
    "requirement_reference_id",
    "requirement_title",
    "requirement_text",
    "requirement_type",
    "requirement_current_status",
    "review_item_review_status",
    "review_item_review_evidence",
    "responsible_name",
    "reviewer_name",
    "review_item_reviewed_at",
    "evidence_filenames",
    "document_name",
    "document_version",
    "baseline_version",
    "source_page",
    "review_cycle_name",
]

SOA_FIELD_LOOKUP = {field.key: field for field in SOA_FIELDS}


def get_statement_of_applicability_fields(
    selected_keys: Optional[Sequence[str]],
) -> list[SoAField]:
    if selected_keys is None:
        return list(SOA_FIELDS)
    if not selected_keys:
        raise ValueError("fields cannot be empty")
    unknown = [key for key in selected_keys if key not in SOA_FIELD_LOOKUP]
    if unknown:
        raise KeyError(",".join(unknown))
    return [SOA_FIELD_LOOKUP[key] for key in selected_keys]


def _group_soa_fields(
    fields: Sequence[SoAField],
) -> tuple[list[SoAField], list[SoAField], list[SoAField]]:
    review_cycle_fields = [
        field
        for field in fields
        if field.key.startswith("review_cycle_") or field.key.startswith("jurisdiction_")
    ]
    document_fields = [field for field in fields if field.key.startswith("document_")]
    detail_fields = [
        field
        for field in fields
        if field not in review_cycle_fields and field not in document_fields
    ]
    return review_cycle_fields, document_fields, detail_fields


def _build_key_value_table(
    fields: Sequence[SoAField],
    context: SoAContext,
    label_style: ParagraphStyle,
    value_style: ParagraphStyle,
    label_width: float,
    value_width: float,
) -> Table:
    table_data = [
        [
            Paragraph(field.label, label_style),
            Paragraph(
                html_module.escape(_normalize_export_text(field.getter(context) or "-")),
                value_style,
            ),
        ]
        for field in fields
    ]
    table = Table(table_data, colWidths=[label_width, value_width], splitInRow=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
            ]
        )
    )
    return table


def _build_two_col_table(
    rows: Sequence[tuple[str, str]],
    label_style: ParagraphStyle,
    value_style: ParagraphStyle,
    label_width: float,
    value_width: float,
) -> Table:
    table_data = [
        [
            Paragraph(label, label_style),
            Paragraph(html_module.escape(_normalize_export_text(value or "-")), value_style),
        ]
        for label, value in rows
    ]
    table = Table(table_data, colWidths=[label_width, value_width], splitInRow=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
            ]
        )
    )
    return table


def _format_date(value: Optional[datetime]) -> str:
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d")


def _format_user_display(
    user: Optional[User],
    user_id: Optional[uuid.UUID],
    mode: str,
) -> str:
    if mode == "ids":
        return _stringify(user_id)
    if user is None:
        return ""
    if mode == "names":
        return user.full_name
    if mode == "emails":
        return user.email
    return f"{user.full_name} <{user.email}>"


async def _fetch_requirements_with_status(
    db: AsyncSession,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None,
):
    if review_cycle_id:
        cycle = await db.get(ReviewCycle, review_cycle_id)
        if cycle is None or cycle.organization_id != organization_id:
            raise HTTPException(404, "Review cycle not found")
        rows = await assessment_rows(db, cycle)
        return [
            (r.requirement, r.document, r.status, r.owner.id if r.owner else None)
            for r in rows
            if r.requirement.requirement_type in (*ACTIONABLE_REQUIREMENT_TYPES, "not_applicable")
        ]
    latest_status = _latest_requirement_status_subquery()
    join_condition = and_(
        Requirement.id == latest_status.c.requirement_id,
        latest_status.c.rn == 1,
    )
    query = select(
        Requirement,
        Document,
        latest_status.c.status,
        latest_status.c.assigned_to,
    ).outerjoin(latest_status, join_condition)

    if review_cycle_id:
        query = (
            query.join(ReviewItem, ReviewItem.requirement_id == Requirement.id)
            .outerjoin(Document, Document.id == Requirement.document_id)
            .where(ReviewItem.review_cycle_id == review_cycle_id)
        )
        if jurisdiction_id is not None:
            query = query.where(Requirement.jurisdiction_id == jurisdiction_id)
    else:
        if jurisdiction_id is None:
            raise ValueError("jurisdiction_id is required when review_cycle_id is not provided")
        query = query.join(Document, Document.id == Requirement.document_id).where(
            Requirement.active.is_(True),
            current_requirement_condition(),
            Requirement.jurisdiction_id == jurisdiction_id,
            Document.status == "approved",
        )
    if organization_id is None:
        query = query.where(Requirement.organization_id.is_(None))
    else:
        query = query.where(Requirement.organization_id == organization_id)

    query = query.where(_reportable_requirement_condition(latest_status.c.status))

    result = await db.execute(query)
    return result.all()


async def _fetch_statement_contexts(
    db: AsyncSession,
    cycle: ReviewCycle,
) -> list[SoAContext]:
    frozen = await frozen_payload(db, cycle)
    jurisdiction = (
        hydrate(Jurisdiction, frozen["jurisdiction"])
        if frozen
        else await _get_jurisdiction_or_404(db, cycle.jurisdiction_id)
    )
    rows = await assessment_rows(db, cycle)
    items = [
        (r.item, r.requirement, r.document, r.status, r.owner.id if r.owner else None) for r in rows
    ]
    evidence_map = {r.item.id: r.files for r in rows}
    assessment_map = {r.item.id: r for r in rows}
    contexts: list[SoAContext] = []
    for item, requirement, document, status, assigned_to in items:
        evidence = _build_evidence_aggregate(
            cycle.id,
            item.id,
            evidence_map.get(item.id, []),
        )
        contexts.append(
            SoAContext(
                jurisdiction=jurisdiction,
                cycle=cycle,
                item=item,
                requirement=requirement,
                document=document,
                requirement_status=RequirementStatusSnapshot(
                    status=status,
                    assigned_to=assigned_to,
                ),
                evidence=evidence,
                reviewer_display=_format_user_display(
                    assessment_map[item.id].reviewer, item.reviewer_id, "both"
                ),
                responsible_display=_format_user_display(
                    assessment_map[item.id].responsible, item.responsible_user_id, "both"
                ),
                baseline_version=_stringify(assessment_map[item.id].source.get("baseline_version")),
                source_page=_stringify(assessment_map[item.id].source.get("page")),
                assessment_rationale=assessment_map[item.id].status_comment,
            )
        )

    return contexts


async def generate_compliance_summary_pdf(
    db: AsyncSession,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None,
) -> bytes:
    """Generate a PDF compliance summary report."""
    jurisdiction = await _resolve_report_jurisdiction(
        db,
        organization_id,
        review_cycle_id,
        jurisdiction_id,
    )
    header_text = _jurisdiction_header_text(jurisdiction)
    header_footer = _draw_header_footer_factory(header_text)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=18,
        spaceAfter=20,
    )
    story.append(Paragraph("Compliance Summary Report", title_style))
    jurisdiction_label = jurisdiction.name
    if jurisdiction.regulator_name:
        jurisdiction_label = f"{jurisdiction.name} ({jurisdiction.regulator_name})"
    story.append(
        Paragraph(f"Jurisdiction: {html_module.escape(jurisdiction_label)}", styles["Normal"])
    )
    if review_cycle_id:
        cycle = await db.get(ReviewCycle, review_cycle_id)
        if cycle is not None:
            story.append(Paragraph(f"Review: {html_module.escape(cycle.name)}", styles["Normal"]))
    story.append(
        Paragraph(
            (
                "Scope: mandatory and recommended requirements."
                if review_cycle_id
                else "Scope: current approved document sets; mandatory and recommended requirements."
            ),
            styles["Normal"],
        )
    )
    story.append(
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"])
    )
    story.append(Spacer(1, 20))

    # Get data
    rows = await _fetch_requirements_with_status(
        db, review_cycle_id, jurisdiction.id, organization_id
    )
    total = len(rows)
    by_status = {}

    for _, _, status, _assigned_to in rows:
        normalized_status = _normalize_status(status)
        by_status[normalized_status] = by_status.get(normalized_status, 0) + 1

    # Summary table
    story.append(Paragraph("Overall Status", styles["Heading2"]))
    summary_data = [["Status", "Count", "Percentage"]]
    for status, count in sorted(by_status.items()):
        pct = round((count / total * 100) if total > 0 else 0, 1)
        summary_data.append([status.replace("_", " ").title(), str(count), f"{pct}%"])
    summary_data.append(["Total", str(total), "100%" if total else "0%"])

    summary_table = Table(summary_data, colWidths=[2.5 * inch, 1.5 * inch, 1.5 * inch])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#08756E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                ("BACKGROUND", (0, -1), (-1, -1), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]
        )
    )
    story.append(summary_table)

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    buffer.seek(0)
    return buffer.getvalue()


async def generate_detailed_report_pdf(
    db: AsyncSession,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None,
) -> bytes:
    """Generate a detailed PDF report with all requirements."""
    jurisdiction = await _resolve_report_jurisdiction(
        db,
        organization_id,
        review_cycle_id,
        jurisdiction_id,
    )
    header_text = _jurisdiction_header_text(jurisdiction)
    header_footer = _draw_header_footer_factory(header_text)

    buffer = io.BytesIO()
    pdf_doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    story = []

    # Title
    story.append(Paragraph("Detailed Compliance Report", styles["Heading1"]))
    jurisdiction_label = jurisdiction.name
    if jurisdiction.regulator_name:
        jurisdiction_label = f"{jurisdiction.name} ({jurisdiction.regulator_name})"
    story.append(
        Paragraph(f"Jurisdiction: {html_module.escape(jurisdiction_label)}", styles["Normal"])
    )
    if review_cycle_id:
        cycle = await db.get(ReviewCycle, review_cycle_id)
        if cycle is not None:
            story.append(Paragraph(f"Review: {html_module.escape(cycle.name)}", styles["Normal"]))
    story.append(
        Paragraph(
            (
                "Scope: mandatory and recommended requirements."
                if review_cycle_id
                else "Scope: current approved document sets; mandatory and recommended requirements."
            ),
            styles["Normal"],
        )
    )
    story.append(
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"])
    )
    story.append(Spacer(1, 20))

    # Get all requirements grouped by document, ordered by numeric reference hierarchy
    rows = await _fetch_requirements_with_status(
        db, review_cycle_id, jurisdiction.id, organization_id
    )

    by_document: dict[str, dict[str, object]] = {}
    for req, document, status, _assigned_to in rows:
        doc_name = document.name or document.filename if document else "Unknown Document"
        if doc_name not in by_document:
            by_document[doc_name] = {
                "requirements": [],
            }
        by_document[doc_name]["requirements"].append((req, status))

    for doc_name, doc_data in sorted(by_document.items(), key=lambda item: item[0]):
        story.append(Spacer(1, 15))
        story.append(Paragraph(f"Document: {html_module.escape(doc_name)}", styles["Heading2"]))

        requirements = sorted(
            doc_data["requirements"],
            key=lambda item: _reference_sort_key(item[0].reference_id),
        )
        for req, status in requirements:
            normalized_status = _normalize_status(status)
            cleaned_text = _strip_html(req.text)

            story.append(
                Paragraph(
                    f"<b>{req.reference_id}</b> [{normalized_status}]",
                    styles["Normal"],
                )
            )
            story.append(
                Paragraph(
                    cleaned_text[:200] + "..." if len(cleaned_text) > 200 else cleaned_text,
                    styles["Normal"],
                )
            )
            story.append(Spacer(1, 8))

    pdf_doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    buffer.seek(0)
    return buffer.getvalue()


GAP_ANALYSIS_HEADERS = [
    "Requires Evidence",
    "No Evidence Recorded",
    "Unassigned",
    "Review Complete",
    "Requirement Owner Name",
    "Requirement Owner Email",
    "Responsible for Requirement Name",
    "Responsible for Requirement Email",
    "Assigned Reviewer Name",
    "Assigned Reviewer Email",
    "Document",
    "Reference ID",
    "Title",
    "Requirement Text",
    "Requirement Type",
    "Requirement Status",
    "Review Cycle ID",
    "Review Cycle Name",
    "Review Deadline",
    "Overdue",
    "Review Status",
    "Reviewer Name",
    "Reviewer Email",
    "Reviewed At",
    "Review Evidence",
    "Evidence File Count",
    "Evidence File Names",
    "Jira Issue Key",
    "Jira Status",
    "Jira Assignee",
    "Jira Issue URL",
    "Needs Attention",
    "Priority",
    "Gap Reasons / Next Actions",
    "Baseline Version",
    "Latest Comment",
    "Review Link",
]

GAP_ANALYSIS_REQUIREMENT_STATUS_ORDER = [
    "not_started",
    "in_progress",
    "blocked",
    "evidenced",
    "not_applicable",
]

GAP_ANALYSIS_REVIEW_STATUS_ORDER = [
    "pending",
    "in_review",
    "confirmed",
    "updated",
    "escalated",
]


def _bool_label(value: Optional[bool]) -> str:
    if value is None:
        return ""
    return "Yes" if value else "No"


def _humanize_status_label(value: str) -> str:
    if not value:
        return ""
    return value.replace("_", " ").title()


def _user_display(name: str, email: str, default: str = "Unassigned") -> str:
    if name and email:
        return f"{name} <{email}>"
    if name or email:
        return name or email
    return default


def _ordered_status_rows(
    counts: Counter[str],
    preferred_order: Sequence[str],
) -> list[tuple[str, str]]:
    preferred_rank = {status: index for index, status in enumerate(preferred_order)}
    ordered = sorted(
        counts.items(),
        key=lambda item: (
            preferred_rank.get(item[0], len(preferred_order)),
            item[0],
        ),
    )
    return [(_humanize_status_label(status), str(count)) for status, count in ordered]


def _ordered_status_keys(
    statuses: Iterable[str],
    preferred_order: Sequence[str],
) -> list[str]:
    preferred_rank = {status: index for index, status in enumerate(preferred_order)}
    return sorted(
        set(statuses),
        key=lambda status: (
            preferred_rank.get(status, len(preferred_order)),
            status,
        ),
    )


def _gap_analysis_owner_status_summary_rows(
    rows: Sequence["GapAnalysisRow"],
) -> tuple[list[str], list[list[str]]]:
    status_keys = _ordered_status_keys(
        [row.requirement_status for row in rows],
        GAP_ANALYSIS_REQUIREMENT_STATUS_ORDER,
    )
    owner_counts: dict[str, Counter[str]] = {}
    for row in rows:
        owner_counts.setdefault(row.owner_group_label, Counter())[row.requirement_status] += 1

    summary_rows: list[list[str]] = []
    for owner in sorted(owner_counts.keys(), key=str.lower):
        counts = owner_counts[owner]
        summary_rows.append(
            [owner, str(sum(counts.values()))]
            + [str(counts.get(status, 0)) for status in status_keys]
        )
    return status_keys, summary_rows


@dataclass(frozen=True)
class GapAnalysisRow:
    requires_evidence: bool
    no_evidence_recorded: Optional[bool]
    unassigned: bool
    review_complete: Optional[bool]
    requirement_owner_name: str
    requirement_owner_email: str
    responsible_user_name: str
    responsible_user_email: str
    assigned_reviewer_name: str
    assigned_reviewer_email: str
    document_name: str
    reference_id: str
    title: str
    requirement_text: str
    requirement_type: str
    requirement_status: str
    cycle_id: str = ""
    cycle_name: str = ""
    review_deadline: str = ""
    overdue: Optional[bool] = None
    review_status: str = ""
    reviewer_name: str = ""
    reviewer_email: str = ""
    reviewed_at: str = ""
    review_evidence: str = ""
    evidence_file_count: str = ""
    evidence_file_names: str = ""
    jira_issue_key: str = ""
    jira_status: str = ""
    jira_assignee: str = ""
    jira_issue_url: str = ""
    gap_reasons: tuple[str, ...] = ()
    baseline_version: str = ""
    latest_comment: str = ""
    review_link: str = ""

    @property
    def needs_attention(self) -> bool:
        return bool(self.gap_reasons) if self.is_cycle_row else self.requires_evidence

    @property
    def priority(self) -> str:
        if not self.needs_attention:
            return "Ready"
        if self.review_status == "escalated":
            return "Escalated"
        if self.requirement_status == "blocked":
            return "Blocked"
        if self.overdue:
            return "Overdue"
        return "Open"

    @property
    def next_actions(self) -> str:
        return "; ".join(self.gap_reasons) or (
            "Complete evidence and record supporting references" if self.requires_evidence else ""
        )

    @property
    def is_cycle_row(self) -> bool:
        return bool(self.cycle_id)

    @property
    def requires_evidence_label(self) -> str:
        return _bool_label(self.requires_evidence)

    @property
    def no_evidence_recorded_label(self) -> str:
        return _bool_label(self.no_evidence_recorded)

    @property
    def unassigned_label(self) -> str:
        return _bool_label(self.unassigned)

    @property
    def review_complete_label(self) -> str:
        return _bool_label(self.review_complete)

    @property
    def overdue_label(self) -> str:
        return _bool_label(self.overdue)

    @property
    def requirement_owner_label(self) -> str:
        return _user_display(
            self.requirement_owner_name,
            self.requirement_owner_email,
        )

    @property
    def responsible_user_label(self) -> str:
        return _user_display(
            self.responsible_user_name,
            self.responsible_user_email,
        )

    @property
    def assigned_reviewer_label(self) -> str:
        return _user_display(
            self.assigned_reviewer_name,
            self.assigned_reviewer_email,
            default="-",
        )

    @property
    def reviewer_label(self) -> str:
        return _user_display(self.reviewer_name, self.reviewer_email, default="-")

    @property
    def owner_group_label(self) -> str:
        if self.is_cycle_row:
            return self.responsible_user_label
        return self.requirement_owner_label

    def to_export_row(self) -> list[str]:
        return [
            self.requires_evidence_label,
            self.no_evidence_recorded_label,
            self.unassigned_label,
            self.review_complete_label,
            self.requirement_owner_name,
            self.requirement_owner_email,
            self.responsible_user_name,
            self.responsible_user_email,
            self.assigned_reviewer_name,
            self.assigned_reviewer_email,
            self.document_name,
            self.reference_id,
            self.title,
            self.requirement_text,
            self.requirement_type,
            self.requirement_status,
            self.cycle_id,
            self.cycle_name,
            self.review_deadline,
            self.overdue_label,
            self.review_status,
            self.reviewer_name,
            self.reviewer_email,
            self.reviewed_at,
            self.review_evidence,
            self.evidence_file_count,
            self.evidence_file_names,
            self.jira_issue_key,
            self.jira_status,
            self.jira_assignee,
            self.jira_issue_url,
            _bool_label(self.needs_attention),
            self.priority,
            self.next_actions,
            self.baseline_version,
            self.latest_comment,
            self.review_link,
        ]


@dataclass(frozen=True)
class GapAnalysisExportContext:
    jurisdiction: Jurisdiction
    cycle: Optional[ReviewCycle]
    rows: list[GapAnalysisRow]
    informational_count: int = 0
    scope_label: str = "Full review"


def _normalize_export_text(value: Optional[str]) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", value).strip()


def _document_name(document: Optional[Document]) -> str:
    if document is None:
        return "Unknown Document"
    return document.name or document.filename or "Unknown Document"


def _user_name(user: Optional[User]) -> str:
    return user.full_name if user else ""


def _user_email(user: Optional[User]) -> str:
    return user.email if user else ""


def _gap_row_sort_key(row: GapAnalysisRow):
    owner_key = row.owner_group_label.lower()
    return (
        {"Escalated": 0, "Blocked": 1, "Overdue": 2, "Open": 3, "Ready": 4}[row.priority],
        0 if row.requires_evidence else 1,
        0 if row.unassigned else 1,
        owner_key,
        row.document_name.lower(),
        _reference_sort_key(row.reference_id),
    )


def _gap_needs_evidence(requirement_status: str) -> bool:
    return requirement_status not in {"evidenced", "not_applicable"}


def _gap_unassigned_for_cycle(requirement_status: str, responsible_user: Optional[User]) -> bool:
    return requirement_status != "not_applicable" and responsible_user is None


def _gap_unassigned_for_master(
    requirement_status: str,
    requirement_owner_user: Optional[User],
) -> bool:
    return requirement_status != "not_applicable" and requirement_owner_user is None


def _needs_evidence_sort_key(row: GapAnalysisRow):
    return (
        row.owner_group_label.lower(),
        row.document_name.lower(),
        _reference_sort_key(row.reference_id),
    )


async def _build_cycle_gap_analysis_rows(
    db: AsyncSession,
    cycle: ReviewCycle,
    current_user: User,
) -> list[GapAnalysisRow]:
    assessments = await assessment_rows(db, cycle)
    if current_user.role == "assigned_reviewer":
        assessments = [r for r in assessments if r.item.assigned_reviewer_id == current_user.id]
    assessments = [
        r
        for r in assessments
        if r.requirement.requirement_type in (*ACTIONABLE_REQUIREMENT_TYPES, "not_applicable")
    ]
    items = [
        (
            r.item,
            r.requirement,
            r.document,
            r.status,
            r.owner,
            r.assigned,
            r.responsible,
            r.reviewer,
        )
        for r in assessments
    ]
    evidence_names_by_item_id = {r.item.id: [f.filename for f in r.files] for r in assessments}
    rationales = {r.item.id: r.status_comment for r in assessments}
    assessments_by_id = {r.item.id: r for r in assessments}
    rows: list[GapAnalysisRow] = []
    for (
        item,
        requirement,
        document,
        requirement_status,
        requirement_owner_user,
        assigned,
        responsible,
        reviewer,
    ) in items:
        normalized_requirement_status = (
            "not_applicable"
            if requirement.requirement_type == "not_applicable"
            else _normalize_status(requirement_status)
        )
        evidence_names = evidence_names_by_item_id.get(item.id, [])
        review_evidence = _normalize_export_text(
            _strip_html(item.review_evidence or rationales.get(item.id) or "")
        )
        requires_evidence = _gap_needs_evidence(normalized_requirement_status)
        no_evidence_recorded = normalized_requirement_status != "not_applicable" and not (
            review_evidence.strip() or evidence_names
        )
        review_complete = (
            bool(item.reviewer_id and item.reviewed_at)
            and item.review_status in COMPLETED_REVIEW_STATUSES
            and (normalized_requirement_status != "not_applicable" or bool(review_evidence.strip()))
        )
        assessment = assessments_by_id[item.id]
        reasons = tuple(assessment_gap_reasons(assessment))
        overdue = bool(_is_overdue(cycle.deadline) and reasons)
        rows.append(
            GapAnalysisRow(
                requires_evidence=requires_evidence,
                no_evidence_recorded=no_evidence_recorded,
                unassigned=_gap_unassigned_for_cycle(normalized_requirement_status, responsible),
                review_complete=review_complete,
                requirement_owner_name=_user_name(requirement_owner_user),
                requirement_owner_email=_user_email(requirement_owner_user),
                responsible_user_name=_user_name(responsible),
                responsible_user_email=_user_email(responsible),
                assigned_reviewer_name=_user_name(assigned),
                assigned_reviewer_email=_user_email(assigned),
                document_name=_document_name(document),
                reference_id=requirement.reference_id,
                title=requirement.title or "",
                requirement_text=_normalize_export_text(_strip_html(requirement.text)),
                requirement_type=requirement.requirement_type,
                requirement_status=normalized_requirement_status,
                cycle_id=_stringify(cycle.id),
                cycle_name=cycle.name,
                review_deadline=_stringify(cycle.deadline),
                overdue=overdue,
                review_status=item.review_status,
                reviewer_name=_user_name(reviewer),
                reviewer_email=_user_email(reviewer),
                reviewed_at=_stringify(item.reviewed_at),
                review_evidence=review_evidence,
                evidence_file_count=str(len(evidence_names)),
                evidence_file_names=_join_values(evidence_names),
                jira_issue_key=item.jira_issue_key or "",
                jira_status=item.jira_status or "",
                jira_assignee=item.jira_assignee or "",
                jira_issue_url=item.jira_issue_url or "",
                gap_reasons=reasons,
                baseline_version=_stringify(assessment.source.get("baseline_version")),
                latest_comment=(
                    _normalize_export_text(_strip_html(assessment.comments[-1].body))
                    if assessment.comments
                    else ""
                ),
                review_link=f"{settings.frontend_base_url.rstrip('/')}/review-cycles/{cycle.id}?mode=focus&item={item.id}",
            )
        )

    return sorted(rows, key=_gap_row_sort_key)


async def _build_master_gap_analysis_rows(
    db: AsyncSession,
    jurisdiction_id: uuid.UUID,
    current_user: User,
) -> list[GapAnalysisRow]:
    latest_status = _latest_requirement_status_subquery()
    join_condition = and_(
        Requirement.id == latest_status.c.requirement_id,
        latest_status.c.rn == 1,
    )
    requirement_owner = aliased(User)

    requirements_query = (
        select(
            Requirement,
            Document,
            latest_status.c.status,
            requirement_owner,
        )
        .join(Document, Document.id == Requirement.document_id)
        .outerjoin(latest_status, join_condition)
        .outerjoin(requirement_owner, requirement_owner.id == latest_status.c.assigned_to)
        .where(
            Requirement.active.is_(True),
            current_requirement_condition(),
            Requirement.jurisdiction_id == jurisdiction_id,
            Requirement.organization_id == current_user.organization_id,
            Document.organization_id == current_user.organization_id,
            Document.status == "approved",
            _reportable_requirement_condition(latest_status.c.status),
        )
        .order_by(Requirement.document_id, Requirement.sort_order)
    )

    requirements_result = await db.execute(requirements_query)
    rows: list[GapAnalysisRow] = []
    for (
        requirement,
        document,
        requirement_status,
        requirement_owner_user,
    ) in requirements_result.all():
        normalized_requirement_status = _normalize_status(requirement_status)
        rows.append(
            GapAnalysisRow(
                requires_evidence=_gap_needs_evidence(normalized_requirement_status),
                no_evidence_recorded=None,
                unassigned=_gap_unassigned_for_master(
                    normalized_requirement_status,
                    requirement_owner_user,
                ),
                review_complete=None,
                requirement_owner_name=_user_name(requirement_owner_user),
                requirement_owner_email=_user_email(requirement_owner_user),
                responsible_user_name="",
                responsible_user_email="",
                assigned_reviewer_name="",
                assigned_reviewer_email="",
                document_name=_document_name(document),
                reference_id=requirement.reference_id,
                title=requirement.title or "",
                requirement_text=_normalize_export_text(_strip_html(requirement.text)),
                requirement_type=requirement.requirement_type,
                requirement_status=normalized_requirement_status,
            )
        )

    return sorted(rows, key=_gap_row_sort_key)


async def _build_gap_analysis_export_context(
    db: AsyncSession,
    current_user: User,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
) -> GapAnalysisExportContext:
    jurisdiction = await _resolve_report_jurisdiction(
        db,
        current_user.organization_id,
        review_cycle_id,
        jurisdiction_id,
    )
    cycle: Optional[ReviewCycle] = None
    if review_cycle_id is not None:
        cycle_query = select(ReviewCycle).where(ReviewCycle.id == review_cycle_id)
        if current_user.organization_id is None:
            cycle_query = cycle_query.where(ReviewCycle.organization_id.is_(None))
        else:
            cycle_query = cycle_query.where(
                ReviewCycle.organization_id == current_user.organization_id
            )
        cycle_result = await db.execute(cycle_query)
        cycle = cycle_result.scalar_one_or_none()
        if cycle is None:
            raise HTTPException(status_code=404, detail="Review cycle not found")
        rows = await _build_cycle_gap_analysis_rows(db, cycle, current_user)
    else:
        rows = await _build_master_gap_analysis_rows(db, jurisdiction.id, current_user)

    informational_count = 0
    if cycle:
        assessments = await assessment_rows(db, cycle)
        if current_user.role == "assigned_reviewer":
            assessments = [r for r in assessments if r.item.assigned_reviewer_id == current_user.id]
        informational_count = sum(
            r.requirement.requirement_type not in {"mandatory", "recommended", "not_applicable"}
            for r in assessments
        )
    return GapAnalysisExportContext(
        jurisdiction=jurisdiction,
        cycle=cycle,
        rows=rows,
        informational_count=informational_count,
        scope_label=(
            "Assigned requirements only"
            if current_user.role == "assigned_reviewer"
            else ("Full review" if cycle else "Current jurisdiction requirements")
        ),
    )


def _gap_analysis_summary(context: GapAnalysisExportContext):
    requirement_status_counts = Counter(row.requirement_status for row in context.rows)
    requirement_rows = _ordered_status_rows(
        requirement_status_counts,
        GAP_ANALYSIS_REQUIREMENT_STATUS_ORDER,
    )
    review_rows: list[tuple[str, str]] = []
    if context.cycle is not None:
        review_status_counts = Counter(
            row.review_status for row in context.rows if row.review_status
        )
        review_rows = _ordered_status_rows(
            review_status_counts,
            GAP_ANALYSIS_REVIEW_STATUS_ORDER,
        )
    return requirement_rows, review_rows


def _gap_analysis_needs_evidence_rows(context: GapAnalysisExportContext) -> list[GapAnalysisRow]:
    return sorted(
        [row for row in context.rows if row.requires_evidence],
        key=_needs_evidence_sort_key,
    )


def _write_gap_analysis_rows(writer: csv.writer, rows: Sequence[GapAnalysisRow]) -> None:
    writer.writerow(GAP_ANALYSIS_HEADERS)
    for row in rows:
        writer.writerow(row.to_export_row())


async def generate_gap_analysis_csv(
    db: AsyncSession,
    current_user: User,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
) -> str:
    context = await _build_gap_analysis_export_context(
        db,
        current_user,
        review_cycle_id,
        jurisdiction_id,
    )
    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    _write_gap_analysis_rows(writer, context.rows)
    return output.getvalue()


async def generate_gap_analysis_xlsx(
    db: AsyncSession,
    current_user: User,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
) -> bytes:
    if Workbook is None:
        raise HTTPException(
            status_code=500,
            detail="XLSX export unavailable (openpyxl is not installed)",
        )

    context = await _build_gap_analysis_export_context(
        db,
        current_user,
        review_cycle_id,
        jurisdiction_id,
    )
    requirement_rows, review_rows = _gap_analysis_summary(context)
    needs_evidence_rows = _gap_analysis_needs_evidence_rows(context)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _append_excel(summary_sheet, ["Gap Analysis Summary"])
    _append_excel(summary_sheet, ["Jurisdiction", context.jurisdiction.name])
    if context.cycle is not None:
        _append_excel(summary_sheet, ["Review Cycle", context.cycle.name])
    if context.cycle is not None:
        _append_excel(summary_sheet, ["Scope", context.scope_label])
        _append_excel(
            summary_sheet,
            ["Assessment", "Frozen at completion" if context.cycle.closed_at else "Live at export"],
        )
        _append_excel(summary_sheet, ["Deadline", _stringify(context.cycle.deadline)])
        _append_excel(summary_sheet, ["Snapshot ID", _stringify(context.cycle.snapshot_id)])
        _append_excel(summary_sheet, [])
        for label, value in summary_metrics(context):
            _append_excel(summary_sheet, [label, value])
        _append_excel(summary_sheet, [])
        _append_excel(
            summary_sheet,
            [
                "Definitions",
                "Needs attention follows unfinished assessment and review decisions, plus missing not-applicable rationale. No evidence recorded and reviewer attribution gaps are shown separately. Informational sections are excluded.",
            ],
        )
        _append_excel(
            summary_sheet,
            [
                "Priority order",
                "Escalated, Blocked, Overdue, Open, Ready. Priority is a workflow order, not an assessment of regulatory severity.",
            ],
        )
    _append_excel(summary_sheet, ["Generated At", datetime.now(timezone.utc).isoformat()])
    _append_excel(summary_sheet, [])
    _append_excel(summary_sheet, ["Requirement Status", "Count"])
    for label, value in requirement_rows:
        _append_excel(summary_sheet, [label, value])
    if context.cycle is not None:
        _append_excel(summary_sheet, [])
        _append_excel(summary_sheet, ["Review Status", "Count"])
        for label, value in review_rows:
            _append_excel(summary_sheet, [label, value])

    needs_evidence_sheet = workbook.create_sheet("Needs Evidence")
    _append_excel(needs_evidence_sheet, GAP_ANALYSIS_HEADERS)
    for row in needs_evidence_rows:
        _append_excel(needs_evidence_sheet, row.to_export_row())

    open_gaps_sheet = workbook.create_sheet("Open Gaps")
    _append_excel(open_gaps_sheet, GAP_ANALYSIS_HEADERS)
    for row in context.rows:
        if row.needs_attention:
            _append_excel(open_gaps_sheet, row.to_export_row())

    all_items_sheet = workbook.create_sheet("All Items")
    _append_excel(all_items_sheet, GAP_ANALYSIS_HEADERS)
    for row in context.rows:
        _append_excel(all_items_sheet, row.to_export_row())

    if context.cycle is not None:
        owners = workbook.create_sheet("Owner Actions")
        _append_excel(
            owners,
            ["Responsible Person", "Open Gaps", "Evidence Gaps", "Review Gaps", "Overdue Gaps"],
        )
        for row in owner_actions(context):
            _append_excel(owners, row)

    buffer = io.BytesIO()
    _format_workbook(workbook)
    if context.cycle is not None:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False
            if sheet.title == "Summary":
                sheet.auto_filter.ref = None
                sheet.freeze_panes = "B2"
                sheet.column_dimensions["A"].width = 38
                sheet.column_dimensions["B"].width = 95
                for cells in sheet.iter_rows():
                    sheet.row_dimensions[cells[0].row].height = 32
                    for cell in cells:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[1].height = 36
                continue
            sheet.freeze_panes = "D2"
            sheet.print_title_rows = "1:1"
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            headers = {cell.value: cell.column for cell in sheet[1]}
            for name, column in headers.items():
                sheet.column_dimensions[get_column_letter(column)].width = (
                    58
                    if name
                    in {
                        "Requirement Text",
                        "Gap Reasons / Next Actions",
                        "Review Evidence",
                        "Latest Comment",
                    }
                    else (
                        34
                        if "Name" in name or name in {"Document", "Title", "Responsible Person"}
                        else 24
                    )
                )
            for cells in sheet.iter_rows(min_row=2):
                sheet.row_dimensions[cells[0].row].height = 72
                for cell in cells:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                priority_column = headers.get("Priority")
                if priority_column:
                    cell = cells[priority_column - 1]
                    cell.fill = PatternFill(
                        "solid",
                        fgColor={
                            "Escalated": "FCE7E7",
                            "Blocked": "FCE7E7",
                            "Overdue": "FFF1D6",
                            "Ready": "E3F3EA",
                        }.get(cell.value, "EDF3F5"),
                    )
                link_column = headers.get("Review Link")
                if link_column:
                    cell = cells[link_column - 1]
                    if str(cell.value or "").startswith(("https://", "http://")):
                        cell.hyperlink = cell.value
                        cell.font = Font(color="087F79", underline="single")
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


async def generate_gap_analysis_pdf(
    db: AsyncSession,
    current_user: User,
    review_cycle_id: Optional[uuid.UUID] = None,
    jurisdiction_id: Optional[uuid.UUID] = None,
) -> bytes:
    context = await _build_gap_analysis_export_context(
        db,
        current_user,
        review_cycle_id,
        jurisdiction_id,
    )
    if context.cycle is not None:
        return render_review_gap_pdf(context)
    requirement_rows, review_rows = _gap_analysis_summary(context)
    needs_evidence_rows = [row for row in context.rows if row.needs_attention]
    header_footer = _draw_header_footer_factory(_jurisdiction_header_text(context.jurisdiction))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=0.65 * inch,
        bottomMargin=0.6 * inch,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "GapTitle",
        parent=styles["Heading1"],
        fontSize=20,
        fontName="Helvetica-Bold",
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "GapSubtitle",
        parent=styles["Normal"],
        fontSize=8.5,
        textColor=colors.HexColor("#637386"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "GapSection",
        parent=styles["Heading2"],
        fontSize=11,
        fontName="Helvetica-Bold",
        spaceBefore=8,
        spaceAfter=4,
    )
    item_heading_style = ParagraphStyle(
        "GapItemHeading",
        parent=styles["Heading3"],
        fontSize=10.5,
        fontName="Helvetica-Bold",
        spaceAfter=4,
    )
    label_style = ParagraphStyle(
        "GapLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        textColor=colors.HexColor("#637386"),
    )
    value_style = ParagraphStyle(
        "GapValue",
        parent=styles["Normal"],
        fontSize=8.5,
    )
    body_style = ParagraphStyle(
        "GapBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
    )

    subtitle_parts = [f"Jurisdiction: {context.jurisdiction.name}"]
    if context.cycle is not None:
        subtitle_parts.append(f"Review Cycle: {context.cycle.name}")
    subtitle_parts.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    story.append(Paragraph("Gap Analysis", title_style))
    story.append(Paragraph(html_module.escape(" | ".join(subtitle_parts)), subtitle_style))

    section_number = 1
    story.append(Paragraph(f"{section_number}. Requirement Status Summary", section_style))
    if requirement_rows:
        story.append(
            _build_two_col_table(
                requirement_rows,
                label_style,
                value_style,
                doc.width * 0.4,
                doc.width * 0.6,
            )
        )
    else:
        story.append(Paragraph("No requirement statuses available.", value_style))
    story.append(Spacer(1, 10))
    section_number += 1

    if context.cycle is not None:
        story.append(Paragraph(f"{section_number}. Review Status Summary", section_style))
        if review_rows:
            story.append(
                _build_two_col_table(
                    review_rows,
                    label_style,
                    value_style,
                    doc.width * 0.4,
                    doc.width * 0.6,
                )
            )
        else:
            story.append(Paragraph("No review statuses available.", value_style))
        story.append(Spacer(1, 10))
        section_number += 1

    story.append(
        Paragraph(f"{section_number}. Outstanding Items By Responsible Person", section_style)
    )
    if not needs_evidence_rows:
        story.append(Paragraph("No outstanding items by responsible person.", value_style))
    else:
        owner_status_keys, owner_status_rows = _gap_analysis_owner_status_summary_rows(
            needs_evidence_rows
        )
        owner_summary_table_data = [
            [
                Paragraph("Responsible Person", label_style),
                Paragraph("Total Outstanding", label_style),
            ]
            + [
                Paragraph(_humanize_status_label(status), label_style)
                for status in owner_status_keys
            ]
        ]
        for row_values in owner_status_rows:
            owner_summary_table_data.append(
                [Paragraph(html_module.escape(row_values[0]), value_style)]
                + [Paragraph(value, value_style) for value in row_values[1:]]
            )

        owner_col_width = 2.6 * inch
        remaining_width = doc.width - owner_col_width
        status_col_count = 1 + len(owner_status_keys)
        value_col_width = (
            remaining_width / status_col_count if status_col_count else remaining_width
        )
        owner_summary_table = Table(
            owner_summary_table_data,
            repeatRows=1,
            colWidths=[owner_col_width] + [value_col_width] * status_col_count,
        )
        owner_summary_table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(owner_summary_table)
    story.append(Spacer(1, 10))
    section_number += 1

    story.append(
        Paragraph(f"{section_number}. Open Gaps: Evidence and Review Decisions", section_style)
    )
    if not needs_evidence_rows:
        story.append(Paragraph("No evidence or review-decision gaps remain.", value_style))
    else:
        grouped_rows: dict[str, list[GapAnalysisRow]] = {}
        for row in needs_evidence_rows:
            grouped_rows.setdefault(row.owner_group_label, []).append(row)

        for owner in sorted(grouped_rows.keys(), key=str.lower):
            owner_items = sorted(grouped_rows[owner], key=_needs_evidence_sort_key)
            story.append(
                Paragraph(f"{html_module.escape(owner)} ({len(owner_items)})", item_heading_style)
            )
            for row in owner_items:
                title_bits = [row.document_name, row.reference_id]
                if row.title:
                    title_bits.append(row.title)
                story.append(Paragraph(html_module.escape(" - ".join(title_bits)), value_style))
                if row.requirement_text:
                    truncated_text = row.requirement_text
                    if len(truncated_text) > 420:
                        truncated_text = truncated_text[:417] + "..."
                    story.append(Paragraph(html_module.escape(truncated_text), body_style))
                details_rows = [
                    ("Requirement Status", _humanize_status_label(row.requirement_status)),
                    ("Requires Evidence", row.requires_evidence_label),
                    ("Unassigned", row.unassigned_label),
                    (
                        "Responsible for Requirement",
                        row.responsible_user_label if row.is_cycle_row else "-",
                    ),
                    ("Requirement Owner", row.requirement_owner_label),
                    ("Assigned Reviewer", row.assigned_reviewer_label),
                    ("Review Status", _humanize_status_label(row.review_status) or "-"),
                    ("Review Complete", row.review_complete_label or "-"),
                    ("Overdue", row.overdue_label or "-"),
                    ("Reviewed At", row.reviewed_at or "-"),
                    ("Review Evidence", row.review_evidence or "-"),
                    ("Evidence Files", row.evidence_file_names or "-"),
                    ("Jira Issue", row.jira_issue_key or "-"),
                ]
                story.append(
                    _build_two_col_table(
                        details_rows,
                        label_style,
                        value_style,
                        doc.width * 0.3,
                        doc.width * 0.7,
                    )
                )
                story.append(Spacer(1, 8))

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    buffer.seek(0)
    return buffer.getvalue()


async def generate_audit_trail_csv(
    db: AsyncSession,
    from_date: Optional[datetime] = None,
    to_date: Optional[datetime] = None,
    organization_id: Optional[uuid.UUID] = None,
    *,
    current_user: Optional[User] = None,
) -> str:
    """Generate a CSV audit trail report."""
    output = io.StringIO()
    writer = _SafeCsvWriter(output)

    # Header
    writer.writerow(
        [
            "Timestamp",
            "User",
            "Action",
            "Entity Type",
            "Entity ID",
            "Old Value",
            "New Value",
        ]
    )

    query = select(AuditLog).order_by(AuditLog.timestamp.desc())
    if current_user is not None:
        from app.services.access_audit import audit_access_clause

        query = query.where(audit_access_clause(current_user))
    if organization_id is None:
        query = query.where(AuditLog.organization_id.is_(None))
    else:
        query = query.where(AuditLog.organization_id == organization_id)
    if from_date:
        query = query.where(AuditLog.timestamp >= from_date)
    if to_date:
        query = query.where(AuditLog.timestamp <= to_date)

    result = await db.execute(query)
    logs = result.scalars().all()

    for log in logs:
        sensitive = log.entity_type in {"application", "preparation_response", "preparation_case"}
        writer.writerow(
            [
                log.timestamp.isoformat(),
                log.user_name,
                log.action,
                log.entity_type,
                log.entity_id,
                "" if sensitive else log.old_value or "",
                "" if sensitive else log.new_value or "",
            ]
        )

    return output.getvalue()


async def generate_statement_of_applicability_xlsx(
    db: AsyncSession,
    cycle: ReviewCycle,
    fields: Sequence[SoAField],
) -> bytes:
    if Workbook is None:
        raise HTTPException(
            status_code=500,
            detail="XLSX export unavailable (openpyxl is not installed)",
        )
    contexts = await _fetch_statement_contexts(db, cycle)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Statement of Applicability"
    _append_excel(sheet, [field.label for field in fields])

    for context in contexts:
        _append_excel(sheet, [field.getter(context) for field in fields])

    buffer = io.BytesIO()
    _format_workbook(workbook)
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


async def generate_statement_of_applicability_pdf(
    db: AsyncSession,
    cycle: ReviewCycle,
    fields: Sequence[SoAField],
) -> bytes:
    contexts = await _fetch_statement_contexts(db, cycle)
    frozen = await frozen_payload(db, cycle)
    jurisdiction = (
        hydrate(Jurisdiction, frozen["jurisdiction"])
        if frozen
        else await _get_jurisdiction_or_404(db, cycle.jurisdiction_id)
    )
    header_footer = _draw_header_footer_factory(_jurisdiction_header_text(jurisdiction))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "SoATitle",
        parent=styles["Heading1"],
        fontSize=18,
        fontName="Helvetica-Bold",
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "SoASubtitle",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#637386"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "SoASection",
        parent=styles["Heading2"],
        fontSize=12,
        fontName="Helvetica-Bold",
        spaceBefore=8,
        spaceAfter=6,
    )
    item_heading_style = ParagraphStyle(
        "SoAItemHeading",
        parent=styles["Heading3"],
        fontSize=11,
        fontName="Helvetica-Bold",
        spaceAfter=4,
    )
    label_style = ParagraphStyle(
        "SoALabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
    )
    value_style = ParagraphStyle(
        "SoAValue",
        parent=styles["Normal"],
        fontSize=8.5,
    )
    body_style = ParagraphStyle(
        "SoABody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
    )

    review_cycle_fields, document_fields, detail_fields = _group_soa_fields(fields)
    selected_keys = {field.key for field in fields}
    show_requirement_text = "requirement_text" in selected_keys

    story.append(Paragraph("Statement of Applicability", title_style))
    story.append(
        Paragraph(
            f"Jurisdiction: {jurisdiction.name} | Review Cycle: {cycle.name} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            subtitle_style,
        )
    )

    cycle_context = contexts[0] if contexts else None
    story.append(Paragraph("1. Review Cycle Information", section_style))
    if review_cycle_fields and cycle_context:
        story.append(
            _build_key_value_table(
                review_cycle_fields,
                cycle_context,
                label_style,
                value_style,
                doc.width * 0.3,
                doc.width * 0.7,
            )
        )
    elif review_cycle_fields:
        story.append(Paragraph("No review cycle data available.", value_style))
    else:
        story.append(Paragraph("No review cycle fields selected.", value_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. Document Information", section_style))
    if not document_fields:
        story.append(Paragraph("No document fields selected.", value_style))
        story.append(Spacer(1, 10))
    else:
        doc_contexts: list[SoAContext] = []
        seen_docs: set[str] = set()
        for context in contexts:
            if context.document and str(context.document.id) not in seen_docs:
                doc_contexts.append(context)
                seen_docs.add(str(context.document.id))

        if not doc_contexts:
            story.append(Paragraph("No document data available.", value_style))
            story.append(Spacer(1, 10))
        elif len(document_fields) <= 6:
            header_row = [Paragraph(field.label, label_style) for field in document_fields]
            table_data = [header_row]
            for context in doc_contexts:
                table_data.append(
                    [
                        Paragraph(
                            html_module.escape(
                                _normalize_export_text(field.getter(context) or "-")
                            ),
                            value_style,
                        )
                        for field in document_fields
                    ]
                )
            table = Table(table_data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 10))
        else:
            for context in doc_contexts:
                doc_name = context.document.name or context.document.filename
                story.append(Paragraph(html_module.escape(doc_name), item_heading_style))
                story.append(
                    _build_key_value_table(
                        document_fields,
                        context,
                        label_style,
                        value_style,
                        doc.width * 0.3,
                        doc.width * 0.7,
                    )
                )
                story.append(Spacer(1, 8))

    story.append(Paragraph("3. Review Items", section_style))
    for index, context in enumerate(contexts, start=1):
        heading_bits: list[str] = []
        if "requirement_reference_id" in selected_keys and context.requirement.reference_id:
            heading_bits.append(context.requirement.reference_id)
        elif "requirement_id" in selected_keys:
            heading_bits.append(_stringify(context.requirement.id))
        else:
            heading_bits.append(str(index))

        title = context.requirement.title or ""
        if "requirement_title" in selected_keys and title:
            heading_bits.append(f"- {title}")

        heading = "Requirement " + " ".join(heading_bits)

        if "requirement_type" in selected_keys:
            heading = f"{heading} (Type: {context.requirement.requirement_type})"

        story.append(Paragraph(html_module.escape(heading), item_heading_style))

        if show_requirement_text:
            requirement_text = _strip_html(context.requirement.text)
            if requirement_text:
                story.append(Paragraph(html_module.escape(requirement_text), body_style))

        detail_fields_filtered = [
            field for field in detail_fields if field.key != "requirement_text"
        ]
        if detail_fields_filtered:
            story.append(
                _build_key_value_table(
                    detail_fields_filtered,
                    context,
                    label_style,
                    value_style,
                    doc.width * 0.3,
                    doc.width * 0.7,
                )
            )
        story.append(Spacer(1, 10))

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    buffer.seek(0)
    return buffer.getvalue()


async def generate_review_cycle_report_pdf(
    db: AsyncSession,
    cycle: ReviewCycle,
    reviewer_display: str = "both",
) -> bytes:
    frozen = await frozen_payload(db, cycle)
    jurisdiction = (
        hydrate(Jurisdiction, frozen["jurisdiction"])
        if frozen
        else await _get_jurisdiction_or_404(db, cycle.jurisdiction_id)
    )
    header_footer = _draw_header_footer_factory(_jurisdiction_header_text(jurisdiction))

    assessments = await assessment_rows(db, cycle)
    rows = [
        (r.item, r.requirement, r.document, r.status, r.assigned, r.responsible, r.reviewer)
        for r in assessments
    ]
    evidence_map = {r.item.id: r.files for r in assessments}
    sources = {r.item.id: r.source for r in assessments}
    doc_groups: dict[str, dict[str, object]] = {}
    total_actionable = 0
    completed_actionable = 0
    pending_actionable = 0

    for item, requirement, document, current_status, assigned, responsible, reviewer in rows:
        doc_name = document.name or document.filename if document else "Unknown Document"
        group = doc_groups.setdefault(
            doc_name, {"document": document, "items": [], "total": 0, "completed": 0, "pending": 0}
        )
        group["items"].append(
            {
                "item": item,
                "requirement": requirement,
                "document": document,
                "status": current_status,
                "assigned": assigned,
                "responsible": responsible,
                "reviewer": reviewer,
            }
        )
        group["total"] += 1

        actionable = (
            requirement.requirement_type in ACTIONABLE_REQUIREMENT_TYPES
            and current_status != "not_applicable"
        )
        if actionable:
            total_actionable += 1
            if item.review_status in COMPLETED_REVIEW_STATUSES:
                completed_actionable += 1
                group["completed"] += 1
            elif item.review_status not in COMPLETED_REVIEW_STATUSES:
                pending_actionable += 1
                group["pending"] += 1

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=0.65 * inch,
        bottomMargin=0.6 * inch,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    story = []

    palette = {
        "ink": colors.HexColor("#0f172a"),
        "muted": colors.HexColor("#64748b"),
        "line": colors.HexColor("#e2e8f0"),
        "panel": colors.HexColor("#f8fafc"),
        "header": colors.HexColor("#f1f5f9"),
    }

    title_style = ParagraphStyle(
        "RCRTitle",
        parent=styles["Heading1"],
        fontSize=20,
        fontName="Helvetica-Bold",
        textColor=palette["ink"],
        spaceAfter=4,
    )
    cycle_name_style = ParagraphStyle(
        "RCRCycleName",
        parent=styles["Normal"],
        fontSize=11,
        fontName="Helvetica-Bold",
        textColor=palette["ink"],
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "RCRSubtitle",
        parent=styles["Normal"],
        fontSize=8.5,
        textColor=palette["muted"],
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "RCRSection",
        parent=styles["Heading2"],
        fontSize=11,
        fontName="Helvetica-Bold",
        textColor=palette["ink"],
        spaceBefore=8,
        spaceAfter=4,
    )
    item_heading_style = ParagraphStyle(
        "RCRItemHeading",
        parent=styles["Heading3"],
        fontSize=10.5,
        fontName="Helvetica-Bold",
        textColor=palette["ink"],
        spaceAfter=4,
    )
    label_style = ParagraphStyle(
        "RCRLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        textColor=palette["muted"],
    )
    value_style = ParagraphStyle(
        "RCRValue",
        parent=styles["Normal"],
        fontSize=8.5,
        textColor=palette["ink"],
    )
    body_style = ParagraphStyle(
        "RCRBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=palette["ink"],
    )
    stat_value_style = ParagraphStyle(
        "RCRStatValue",
        parent=styles["Normal"],
        fontSize=12,
        fontName="Helvetica-Bold",
        alignment=1,
        textColor=palette["ink"],
    )
    stat_label_style = ParagraphStyle(
        "RCRStatLabel",
        parent=styles["Normal"],
        fontSize=7.5,
        alignment=1,
        textColor=palette["muted"],
    )

    def build_section_rule() -> Table:
        rule = Table([[""]], colWidths=[doc.width])
        rule.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.6, palette["line"])]))
        return rule

    def build_metrics_table() -> Table:
        labels = ["Applicable controls", "Reviewed", "Decisions pending", "Requirement sets"]
        values = [
            str(total_actionable),
            str(completed_actionable),
            str(pending_actionable),
            str(len(doc_groups)),
        ]
        table_data = [
            [Paragraph(value, stat_value_style) for value in values],
            [Paragraph(label, stat_label_style) for label in labels],
        ]
        table = Table(table_data, colWidths=[doc.width / len(values)] * len(values))
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), palette["panel"]),
                    ("BOX", (0, 0), (-1, -1), 0.5, palette["line"]),
                    ("INNERGRID", (0, 0), (-1, -1), 0.25, palette["line"]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return table

    story.append(Paragraph("Readiness Review Report", title_style))
    story.append(
        Paragraph(
            "Supporting assessment and evidence. Check submission and approval requirements with the receiving authority or certification body.",
            styles["Normal"],
        )
    )
    story.append(Paragraph(html_module.escape(cycle.name), cycle_name_style))
    story.append(
        Paragraph(
            f"Jurisdiction: {html_module.escape(jurisdiction.name)} | Status: {html_module.escape(cycle.status.title())} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            subtitle_style,
        )
    )
    story.append(build_metrics_table())
    story.append(
        Paragraph(
            f"Full register: {len(assessments)} sections. Informational and not-applicable entries are retained below.",
            subtitle_style,
        )
    )
    story.append(Spacer(1, 12))

    story.append(Paragraph("Contents", section_style))
    story.append(build_section_rule())
    toc_rows = []
    ordered_items: list[dict[str, object]] = []
    for doc_name in sorted(doc_groups.keys()):
        group = doc_groups[doc_name]
        items = list(group["items"])
        items.sort(
            key=lambda entry: (
                entry["requirement"].sort_order,
                _reference_sort_key(entry["requirement"].reference_id),
            )
        )
        ordered_items.extend(items)

    for entry in ordered_items:
        requirement = entry["requirement"]
        requirement_label = requirement.reference_id or _stringify(requirement.id)
        title = requirement.title or ""
        anchor_name = f"req-{requirement.id}"
        label_link = Paragraph(
            f'<a href="#{anchor_name}">{html_module.escape(requirement_label)}</a>', value_style
        )
        title_text = (
            title
            if title
            else textwrap.shorten(_strip_html(requirement.text), width=120, placeholder="...")
        )
        title_link = Paragraph(
            f'<a href="#{anchor_name}">{html_module.escape(title_text)}</a>', value_style
        )
        toc_rows.append([label_link, title_link])

    if toc_rows:
        toc_table = Table(
            toc_rows,
            colWidths=[doc.width * 0.2, doc.width * 0.8],
        )
        toc_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, palette["panel"]]),
                ]
            )
        )
        story.append(toc_table)
    else:
        story.append(Paragraph("No requirements available.", value_style))
    story.append(Spacer(1, 12))

    summary_rows = [
        ("Name", cycle.name),
        ("Status", cycle.status),
        ("Scope", cycle.scope),
        ("Deadline", _format_date(cycle.deadline) or "-"),
        ("Created At", _format_date(cycle.created_at) or "-"),
        ("Applicable controls", str(total_actionable)),
        ("Completed Items", str(completed_actionable)),
        ("Pending Items", str(pending_actionable)),
    ]
    story.append(Paragraph("1. Review Cycle Summary", section_style))
    story.append(build_section_rule())
    summary_table = _build_two_col_table(
        summary_rows, label_style, value_style, doc.width * 0.3, doc.width * 0.7
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, palette["line"]),
                ("BACKGROUND", (0, 0), (0, -1), palette["panel"]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(summary_table)
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. Document Summary", section_style))
    story.append(build_section_rule())
    if doc_groups:
        table_data = [
            [
                Paragraph("Document", label_style),
                Paragraph("Total Items", label_style),
                Paragraph("Completed", label_style),
                Paragraph("Pending", label_style),
            ]
        ]
        for doc_name in sorted(doc_groups.keys()):
            group = doc_groups[doc_name]
            table_data.append(
                [
                    Paragraph(html_module.escape(doc_name), value_style),
                    Paragraph(str(group["total"]), value_style),
                    Paragraph(str(group["completed"]), value_style),
                    Paragraph(str(group["pending"]), value_style),
                ]
            )
        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[doc.width * 0.52, doc.width * 0.16, doc.width * 0.16, doc.width * 0.16],
        )
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, palette["line"]),
                    ("BACKGROUND", (0, 0), (-1, 0), palette["header"]),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, palette["panel"]]),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("No document data available.", value_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph("3. Review Items", section_style))
    story.append(build_section_rule())
    for doc_name in sorted(doc_groups.keys()):
        group = doc_groups[doc_name]
        doc_header = Table(
            [
                [
                    Paragraph(html_module.escape(doc_name), item_heading_style),
                    Paragraph(
                        f"{group['pending']} pending of {group['total']} requirements",
                        subtitle_style,
                    ),
                ]
            ],
            colWidths=[doc.width * 0.7, doc.width * 0.3],
        )
        doc_header.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), palette["panel"]),
                    ("BOX", (0, 0), (-1, -1), 0.4, palette["line"]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(doc_header)
        story.append(Spacer(1, 6))
        items = group["items"]
        items.sort(
            key=lambda entry: (
                entry["requirement"].sort_order,
                _reference_sort_key(entry["requirement"].reference_id),
            )
        )
        for entry in items:
            item = entry["item"]
            requirement = entry["requirement"]
            status = entry["status"]
            assigned = entry["assigned"]
            responsible = entry["responsible"]
            reviewer = entry["reviewer"]

            requirement_label = requirement.reference_id or _stringify(requirement.id)
            title = requirement.title or ""
            anchor_name = f"req-{requirement.id}"
            heading = f"Requirement {requirement_label}"
            if title:
                heading = f"{heading} — {title}"
            heading = f"{heading} (Type: {requirement.requirement_type.replace('_', ' ').title()})"

            requirement_text = _strip_html(requirement.text)
            detail_heading = Paragraph(
                f'<a name="{anchor_name}"/>{html_module.escape(heading)}', item_heading_style
            )

            evidence_files = evidence_map.get(item.id, [])
            evidence_text = (
                " ; ".join(
                    f"{file.filename} ({_format_date(file.uploaded_at)})" for file in evidence_files
                )
                if evidence_files
                else "-"
            )

            source = sources.get(item.id, {})
            details_rows = [
                (
                    "Baseline / source page",
                    f"Version {source.get('baseline_version') or 'legacy'}; page {source.get('page') or 'not recorded'}",
                ),
                ("Review Status", item.review_status.replace("_", " ").title()),
                ("Requirement Status", _normalize_status(status).replace("_", " ").title()),
                (
                    "Assigned Reviewer",
                    _format_user_display(assigned, item.assigned_reviewer_id, reviewer_display),
                ),
                (
                    "Responsible User",
                    _format_user_display(responsible, item.responsible_user_id, reviewer_display),
                ),
                (
                    "Reviewer",
                    _format_user_display(reviewer, item.reviewer_id, reviewer_display),
                ),
                ("Reviewed At", _format_date(item.reviewed_at) or "-"),
                (
                    (
                        "Applicability rationale"
                        if status == "not_applicable"
                        or requirement.requirement_type == "not_applicable"
                        else "Review Evidence"
                    ),
                    item.review_evidence or "-",
                ),
                ("Jira Issue Key", item.jira_issue_key or "-"),
                ("Jira Status", item.jira_status or "-"),
                ("Jira Summary", item.jira_summary or "-"),
                ("Jira Assignee", item.jira_assignee or "-"),
                ("Jira Priority", item.jira_priority or "-"),
                ("Jira Updated At", _format_date(item.jira_updated_at) or "-"),
                ("Jira Issue URL", item.jira_issue_url or "-"),
                ("Evidence Files", evidence_text),
            ]
            # Keep the assessment readable; integrations with no data add no evidence.
            details_rows = [
                (label, value)
                for label, value in details_rows
                if not label.startswith("Jira ") or value not in ("", "-")
            ]
            if requirement.requirement_type == "informational":
                details_rows = details_rows[:1]
            details_table = _build_two_col_table(
                details_rows, label_style, value_style, doc.width * 0.27, doc.width * 0.73
            )
            details_table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, palette["line"]),
                        ("BACKGROUND", (0, 0), (0, -1), palette["panel"]),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )

            card_flowables = [detail_heading]
            if requirement_text:
                card_flowables.append(Paragraph(html_module.escape(requirement_text), body_style))
            card_flowables.append(details_table)

            story.extend(card_flowables)
            story.append(Spacer(1, 10))

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    buffer.seek(0)
    return buffer.getvalue()


def _apply_register_org_scope(query, organization_id: Optional[uuid.UUID]):
    if organization_id is None:
        return query.where(ComponentRegister.organization_id.is_(None))
    return query.where(ComponentRegister.organization_id == organization_id)


def _export_approved_scope(scope):
    if not scope:
        return None
    # These organisation-wide exports must not expose protected assessment IDs.
    return {
        **{key: value for key, value in scope.items() if key != "blocking_assessment_ids"},
        "blocking_assessment_count": len(scope.get("blocking_assessment_ids", [])),
    }


async def generate_change_management_components_csv(
    db: AsyncSession,
    jurisdiction_id: uuid.UUID,
    organization_id: Optional[uuid.UUID],
) -> str:
    query = (
        select(Component, ComponentRegister)
        .join(ComponentRegister, Component.register_id == ComponentRegister.id)
        .where(ComponentRegister.jurisdiction_id == jurisdiction_id)
        .order_by(Component.component_uid.asc())
    )
    query = _apply_register_org_scope(query, organization_id)

    result = await db.execute(query)
    rows = result.all()

    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    writer.writerow(
        [
            "register id",
            "component id",
            "component uid",
            "definition",
            "version",
            "identifying characteristics",
            "change owner id",
            "change owner name",
            "confidentiality code",
            "integrity code",
            "availability code",
            "accountability code",
            "classification code",
            "checksum/hash",
            "is hardware",
            "geographic location",
            "hosting model",
            "virtualized",
            "public cloud provider",
            "public cloud certification",
            "public cloud independent",
            "public cloud redundancy",
            "status",
            "regulatory scope",
            "responsibility role",
            "created at",
            "updated at",
        ]
    )
    for component, register in rows:
        writer.writerow(
            [
                str(register.id),
                str(component.id),
                component.component_uid,
                component.definition,
                component.version,
                component.identifying_characteristics,
                _stringify(component.change_owner_id),
                _stringify(component.change_owner_name),
                component.confidentiality_code,
                component.integrity_code,
                component.availability_code,
                component.accountability_code,
                component.classification_code,
                _stringify(component.checksum_hash),
                component.is_hardware,
                _stringify(component.geographic_location),
                component.hosting_model,
                component.virtualized,
                _stringify(component.public_cloud_provider),
                component.public_cloud_certification or "",
                component.public_cloud_independent,
                component.public_cloud_redundancy,
                component.status,
                component.regulatory_scope,
                register.responsibility_role,
                _stringify(component.created_at),
                _stringify(component.updated_at),
            ]
        )
    return output.getvalue()


async def generate_change_management_component_history_csv(
    db: AsyncSession,
    component_id: uuid.UUID,
    organization_id: Optional[uuid.UUID],
) -> str:
    component_query = _apply_register_org_scope(
        select(Component, ComponentRegister)
        .join(ComponentRegister, Component.register_id == ComponentRegister.id)
        .where(Component.id == component_id),
        organization_id,
    )
    component_result = await db.execute(component_query)
    component_row = component_result.first()
    if component_row is None:
        raise HTTPException(status_code=404, detail="Component not found")
    component, _ = component_row

    entries_result = await db.execute(
        select(ChangeEntry, ChangeEntryComponent)
        .join(ChangeEntryComponent, ChangeEntryComponent.change_entry_id == ChangeEntry.id)
        .where(ChangeEntryComponent.component_id == component.id)
        .order_by(ChangeEntry.created_at.asc())
    )
    rows = entries_result.all()
    entry_ids = [change_entry.id for change_entry, _ in rows]

    status_timeline_map: dict[uuid.UUID, str] = {entry_id: "" for entry_id in entry_ids}
    decision_notes_map: dict[uuid.UUID, str] = {entry_id: "" for entry_id in entry_ids}
    if entry_ids:
        events_result = await db.execute(
            select(ChangeEvent)
            .where(
                and_(
                    ChangeEvent.change_entry_id.in_(entry_ids),
                    ChangeEvent.deleted_at.is_(None),
                )
            )
            .order_by(ChangeEvent.created_at.asc())
        )
        timeline_parts: dict[uuid.UUID, list[str]] = {entry_id: [] for entry_id in entry_ids}
        decision_parts: dict[uuid.UUID, list[str]] = {entry_id: [] for entry_id in entry_ids}
        for event in events_result.scalars().all():
            if event.event_type == "status_change":
                transition = event.status_to or event.status_from or "status_change"
                timeline_parts[event.change_entry_id].append(transition)
            if event.event_type == "decision" and event.note:
                decision_parts[event.change_entry_id].append(event.note)
        for entry_id in entry_ids:
            status_timeline_map[entry_id] = " -> ".join(timeline_parts[entry_id])
            decision_notes_map[entry_id] = _join_values(decision_parts[entry_id])

    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    writer.writerow(
        [
            "component id",
            "component uid",
            "change id",
            "change title",
            "status",
            "change type",
            "version at proposal",
            "planned version",
            "implemented version",
            "proposed at",
            "approved at",
            "implemented at",
            "verified at",
            "integration related",
            "testing org status",
            "testing org cycle",
            "testing org approved at",
            "approval decision",
            "rejection reason",
            "verification notes",
            "status timeline",
            "decision notes",
            "frozen component snapshot",
            "planned checksum",
            "implemented checksum",
            "approved scope",
            "compliance attestations",
        ]
    )

    for change_entry, link in rows:
        writer.writerow(
            [
                str(component.id),
                (link.frozen_snapshot or {}).get("component_uid", "unresolved historical UID"),
                str(change_entry.id),
                change_entry.title,
                change_entry.status,
                _stringify(change_entry.change_type),
                _stringify(link.version_at_proposal),
                _stringify(link.planned_version),
                _stringify(link.implemented_version),
                _stringify(change_entry.proposed_at),
                _stringify(change_entry.approved_at),
                _stringify(change_entry.implemented_at),
                _stringify(change_entry.verified_at),
                change_entry.integration_related,
                _stringify(change_entry.testing_org_status),
                _stringify(change_entry.testing_org_cycle),
                _stringify(change_entry.testing_org_approved_at),
                _stringify(change_entry.approval_decision),
                _stringify(change_entry.rejection_reason),
                _stringify(change_entry.verification_notes),
                status_timeline_map.get(change_entry.id, ""),
                decision_notes_map.get(change_entry.id, ""),
                json.dumps(link.frozen_snapshot, sort_keys=True),
                _stringify(link.planned_checksum_hash),
                _stringify(link.implemented_checksum_hash),
                json.dumps(_export_approved_scope(change_entry.approved_scope), sort_keys=True),
                json.dumps(change_entry.compliance, sort_keys=True),
            ]
        )
    return output.getvalue()


async def generate_change_management_hardware_locations_csv(
    db: AsyncSession,
    jurisdiction_id: uuid.UUID,
    organization_id: Optional[uuid.UUID],
) -> str:
    query = _apply_register_org_scope(
        select(Component, ComponentRegister)
        .join(ComponentRegister, Component.register_id == ComponentRegister.id)
        .where(
            and_(
                ComponentRegister.jurisdiction_id == jurisdiction_id,
                Component.is_hardware.is_(True),
            )
        )
        .order_by(Component.component_uid.asc()),
        organization_id,
    )
    result = await db.execute(query)
    rows = result.all()

    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    writer.writerow(
        [
            "register id",
            "component id",
            "component uid",
            "definition",
            "hosting model",
            "geographic location",
            "public cloud provider",
            "status",
            "updated at",
        ]
    )
    for component, register in rows:
        writer.writerow(
            [
                str(register.id),
                str(component.id),
                component.component_uid,
                component.definition,
                component.hosting_model,
                _stringify(component.geographic_location),
                _stringify(component.public_cloud_provider),
                component.status,
                _stringify(component.updated_at),
            ]
        )
    return output.getvalue()


async def generate_change_management_verified_changes_csv(
    db: AsyncSession,
    jurisdiction_id: uuid.UUID,
    organization_id: Optional[uuid.UUID],
    period_days: Optional[int],
) -> str:
    query = (
        select(ChangeEntry, ComponentRegister)
        .join(ComponentRegister, ChangeEntry.register_id == ComponentRegister.id)
        .where(
            and_(
                ComponentRegister.jurisdiction_id == jurisdiction_id,
                ChangeEntry.status == "verified",
            )
        )
        .order_by(ChangeEntry.verified_at.desc(), ChangeEntry.created_at.desc())
    )
    query = _apply_register_org_scope(query, organization_id)

    if period_days is not None:
        cutoff = datetime.now(timezone.utc) - timedelta(days=period_days)
        query = query.where(ChangeEntry.verified_at >= cutoff)

    entries_result = await db.execute(query)
    rows = entries_result.all()

    entry_ids = [row[0].id for row in rows]
    components_map: dict[uuid.UUID, list[str]] = {entry_id: [] for entry_id in entry_ids}
    if entry_ids:
        component_rows = await db.execute(
            select(ChangeEntryComponent.change_entry_id, ChangeEntryComponent.frozen_snapshot)
            .join(Component, Component.id == ChangeEntryComponent.component_id)
            .where(ChangeEntryComponent.change_entry_id.in_(entry_ids))
            .order_by(Component.component_uid.asc())
        )
        for entry_id, snapshot in component_rows.all():
            components_map[entry_id].append(
                (snapshot or {}).get("component_uid", "unresolved historical UID")
            )

    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    writer.writerow(
        [
            "register id",
            "change id",
            "title",
            "status",
            "change type",
            "linked components",
            "approval decision",
            "verification notes",
            "testing org required",
            "testing org status",
            "testing org cycle",
            "testing org next due at",
            "testing org approved at",
            "verified at",
            "implemented at",
            "proposed at",
            "approved scope",
            "compliance attestations",
        ]
    )

    for entry, register in rows:
        writer.writerow(
            [
                str(register.id),
                str(entry.id),
                entry.title,
                entry.status,
                _stringify(entry.change_type),
                _join_values(components_map.get(entry.id, [])),
                _stringify(entry.approval_decision),
                _stringify(entry.verification_notes),
                entry.testing_org_required,
                _stringify(entry.testing_org_status),
                _stringify(entry.testing_org_cycle),
                _stringify(entry.testing_org_next_due_at),
                _stringify(entry.testing_org_approved_at),
                _stringify(entry.verified_at),
                _stringify(entry.implemented_at),
                _stringify(entry.proposed_at),
                json.dumps(_export_approved_scope(entry.approved_scope), sort_keys=True),
                json.dumps(entry.compliance, sort_keys=True),
            ]
        )
    return output.getvalue()


async def generate_change_management_integration_changes_csv(
    db: AsyncSession,
    jurisdiction_id: uuid.UUID,
    organization_id: Optional[uuid.UUID],
    period_days: Optional[int],
) -> str:
    query = (
        select(ChangeEntry, ComponentRegister)
        .join(ComponentRegister, ChangeEntry.register_id == ComponentRegister.id)
        .where(
            and_(
                ComponentRegister.jurisdiction_id == jurisdiction_id,
                ChangeEntry.integration_related.is_(True),
            )
        )
        .order_by(ChangeEntry.created_at.desc())
    )
    query = _apply_register_org_scope(query, organization_id)

    if period_days is not None:
        cutoff = datetime.now(timezone.utc) - timedelta(days=period_days)
        query = query.where(ChangeEntry.created_at >= cutoff)

    entries_result = await db.execute(query)
    rows = entries_result.all()
    entry_ids = [row[0].id for row in rows]

    checks_map: dict[uuid.UUID, list[IntegrationCheck]] = {entry_id: [] for entry_id in entry_ids}
    components_map: dict[uuid.UUID, list[str]] = {entry_id: [] for entry_id in entry_ids}
    if entry_ids:
        checks_result = await db.execute(
            select(IntegrationCheck)
            .where(
                IntegrationCheck.change_entry_id.in_(entry_ids),
                IntegrationCheck.deleted_at.is_(None),
            )
            .order_by(IntegrationCheck.created_at.asc())
        )
        for check in checks_result.scalars().all():
            checks_map[check.change_entry_id].append(check)
        component_rows = await db.execute(
            select(ChangeEntryComponent.change_entry_id, ChangeEntryComponent.frozen_snapshot)
            .join(Component, Component.id == ChangeEntryComponent.component_id)
            .where(ChangeEntryComponent.change_entry_id.in_(entry_ids))
            .order_by(Component.component_uid.asc())
        )
        for entry_id, snapshot in component_rows.all():
            components_map[entry_id].append(
                (snapshot or {}).get("component_uid", "unresolved historical UID")
            )

    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    writer.writerow(
        [
            "register id",
            "change id",
            "title",
            "status",
            "change type",
            "linked components",
            "proposed at",
            "implemented at",
            "verified at",
            "testing org status",
            "testing org cycle",
            "integration checks total",
            "integration checks completed",
            "integration actions",
            "integration action references",
            "integration results",
            "integration completed at",
            "integration evidence notes",
            "approved scope",
            "compliance attestations",
        ]
    )

    for entry, register in rows:
        checks = checks_map.get(entry.id, [])
        actions = [check.action for check in checks]
        action_references = [check.action_reference or "" for check in checks]
        results = [check.result for check in checks]
        completed_at_values = [
            check.completed_at.isoformat() if check.completed_at else "" for check in checks
        ]
        evidence_notes = [check.evidence_notes or "" for check in checks]
        completed = [check for check in checks if check.completed_at is not None]

        writer.writerow(
            [
                str(register.id),
                str(entry.id),
                entry.title,
                entry.status,
                _stringify(entry.change_type),
                _join_values(components_map.get(entry.id, [])),
                _stringify(entry.proposed_at),
                _stringify(entry.implemented_at),
                _stringify(entry.verified_at),
                _stringify(entry.testing_org_status),
                _stringify(entry.testing_org_cycle),
                len(checks),
                len(completed),
                _join_values(actions),
                _join_values(action_references),
                _join_values(results),
                _join_values(completed_at_values),
                _join_values(evidence_notes),
                json.dumps(_export_approved_scope(entry.approved_scope), sort_keys=True),
                json.dumps(entry.compliance, sort_keys=True),
            ]
        )
    return output.getvalue()


def _format_workbook(workbook):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for sheet in workbook:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173B45")
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            width = min(60, max(16, max(len(str(c.value or "")) for c in column[:100]) + 2))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
        sheet.sheet_view.showGridLines = False


async def generate_change_management_programme_assurance_csv(db, jurisdiction_id, organization_id):
    from app.services.change_readiness import programme_readiness, denmark_profile

    rows = (
        await db.scalars(
            _apply_register_org_scope(
                select(ComponentRegister).where(
                    ComponentRegister.jurisdiction_id == jurisdiction_id
                ),
                organization_id,
            )
        )
    ).all()
    output = io.StringIO()
    writer = _SafeCsvWriter(output)
    writer.writerow(
        [
            "register id",
            "register name",
            "responsibility role",
            "rule profile",
            "programme assurance",
            "certification due at",
            "report due at",
            "ready",
            "unresolved reasons",
        ]
    )
    for register in rows:
        dk = await denmark_profile(db, register)
        gate = programme_readiness(register) if dk else {}
        writer.writerow(
            [
                str(register.id),
                register.name,
                register.responsibility_role,
                "denmark_scp_3.1" if dk else "unconfigured",
                json.dumps(register.programme_assurance, sort_keys=True),
                gate.get("certification_due_at"),
                gate.get("report_due_at"),
                gate.get("ready"),
                json.dumps(gate.get("reasons", [])),
            ]
        )
    return output.getvalue()
