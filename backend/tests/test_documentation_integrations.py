"""Synthetic workbook persistence, frozen exports and bounded assistance contracts."""
import io
import json
import zipfile
import uuid

from openpyxl import Workbook
import pytest

from tests.test_documentation_processes import review_process  # noqa: F401

from tests.test_documentation_register import (  # noqa: F401
    BASE, documentation_actors, create_register, create_expectation, create_evidence, link_evidence,
)


LONG_REQUIRED_DOCUMENTATION = "  Synthetic governance policy — α\n" + "Retain accountable oversight records. " * 40 + "\nEnd of original column A.  "
EXACT_SOURCE_COMMENT = "  Conditional on activity; confirm applicability\nSource comment stays separate.  "
EXACT_DOCUMENT_NAME = "  Supplier draft policy.docx\nSecond document name  "


def synthetic_workbook():
    workbook = Workbook()
    workbook.create_sheet("Empty cover", 0)
    sheet = workbook.worksheets[1]
    sheet.title = "Games"
    sheet.append(["Synthetic documentation list — banner"])
    sheet.append(["Document / required content", "SCP reference", "Comment/explanation", "Document", "Location"])
    for index in range(63):
        sheet.append([
            LONG_REQUIRED_DOCUMENTATION if index == 0 else f"Synthetic policy {index}",
            f"UNKNOWN-{index}", EXACT_SOURCE_COMMENT if index == 0 else "" if index == 1 else "Conditional on activity; confirm applicability",
            EXACT_DOCUMENT_NAME if index == 0 else "Supplier draft", "Section 2",
        ])
    second = workbook.create_sheet("Records")
    second.append(["Required record / report", "DGA source and location", "Comment/explanation"])
    for index in range(63):
        second.append([f"Synthetic record {index}", f"ANNEX-{index}", "Scoped test period"])
    notes = workbook.create_sheet("Notes and sources")
    notes.append(["Topic / source code", "Explanation / document", "Version / reference"])
    notes.append(["SCP", "Synthetic source alias; not an obligation", "v3"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def test_multisheet_workbook_import_preserves_metadata_and_is_idempotent_across_batches(
    client, documentation_actors, default_jurisdiction
):
    owner = documentation_actors["owner"]
    register = await create_register(client, owner, default_jurisdiction)
    synthetic_bytes = synthetic_workbook()
    preview_response = await client.post(f"{BASE}/registers/{register['id']}/imports/preview",
        headers=owner["headers"], files={"file": ("synthetic-documentation.xlsx", synthetic_bytes)})
    assert preview_response.status_code == 200, preview_response.text
    preview = preview_response.json()
    assert len(preview["rows"]) == 126
    assert [sheet["name"] for sheet in preview["sheets"]] == ["Empty cover", "Games", "Records", "Notes and sources"]
    assert preview["sheets"][0]["included"] is False
    assert preview["sheets"][0]["raw_rows"] == []
    assert preview["sheets"][1]["header_row"] == 2
    assert {row["sheet"] for row in preview["rows"]} == {"Games", "Records"}
    assert preview["rows"][0]["required_documentation"] == LONG_REQUIRED_DOCUMENTATION
    assert preview["rows"][0]["description"] == EXACT_SOURCE_COMMENT
    assert preview["rows"][1]["description"] == ""
    assert preview["rows"][0]["document_name"] == EXACT_DOCUMENT_NAME
    assert preview["rows"][0]["source_reference"] == "UNKNOWN-0"
    assert preview["rows"][0]["metadata"]["Location"] == "Section 2"
    assert (await client.get(f"{BASE}/registers/{register['id']}/expectations", headers=owner["headers"])).json()["total"] == 0
    revision = preview["revision"]
    endpoint = f"{BASE}/registers/{register['id']}/imports/{preview['id']}/confirm"
    for batch in (preview["rows"][:60], preview["rows"][60:]):
        imported = await client.post(endpoint, headers=owner["headers"], json={"revision": revision,
            "choices": [{"row_key": row["key"], "action": "create"} for row in batch]})
        assert imported.status_code == 200, imported.text
        assert imported.json()["created"] == len(batch)
        revision = imported.json()["revision"]
    repeated = await client.post(endpoint, headers=owner["headers"], json={"revision": revision,
        "choices": [{"row_key": row["key"], "action": "create"} for row in preview["rows"]]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["created"] == 0
    assert repeated.json()["repeated"] == 126
    repreview = await client.post(f"{BASE}/registers/{register['id']}/imports/preview", headers=owner["headers"],
        files={"file": ("synthetic-documentation.xlsx", synthetic_bytes)})
    assert repreview.status_code == 200, repreview.text
    reimported = await client.post(f"{BASE}/registers/{register['id']}/imports/{repreview.json()['id']}/confirm",
        headers=owner["headers"], json={"revision": repreview.json()["revision"],
            "choices": [{"row_key": row["key"], "action": "create"} for row in repreview.json()["rows"]]})
    assert reimported.status_code == 200, reimported.text
    assert reimported.json()["created"] == 0
    assert reimported.json()["repeated"] == 126
    sections = await client.get(f"{BASE}/registers/{register['id']}/sections", headers=owner["headers"])
    assert sections.status_code == 200, sections.text
    assert [item["name"] for item in sections.json()["items"]] == ["Games", "Records"]
    sections_by_id = {item["id"]: item["name"] for item in sections.json()["items"]}
    listing = await client.get(f"{BASE}/registers/{register['id']}/expectations?limit=500", headers=owner["headers"])
    assert listing.json()["total"] == 126
    expected_order = [f"UNKNOWN-{index}" for index in range(63)] + [f"ANNEX-{index}" for index in range(63)]
    assert [row["source_reference"] for row in listing.json()["items"]] == expected_order
    boundary = await client.get(f"{BASE}/registers/{register['id']}/expectations", headers=owner["headers"],
        params={"skip": 61, "limit": 4})
    assert boundary.status_code == 200, boundary.text
    assert [row["source_reference"] for row in boundary.json()["items"]] == ["UNKNOWN-61", "UNKNOWN-62", "ANNEX-0", "ANNEX-1"]
    source_only_search = await client.get(f"{BASE}/registers/{register['id']}/expectations", headers=owner["headers"],
        params={"q": "ANNEX-62", "limit": 1})
    assert source_only_search.status_code == 200, source_only_search.text
    assert source_only_search.json()["total"] == 1
    assert [row["source_reference"] for row in source_only_search.json()["items"]] == ["ANNEX-62"]
    by_reference = {row["source_reference"]: row for row in listing.json()["items"]}
    assert by_reference["UNKNOWN-0"]["required_documentation"] == LONG_REQUIRED_DOCUMENTATION
    assert by_reference["UNKNOWN-0"]["description"] == EXACT_SOURCE_COMMENT
    assert by_reference["UNKNOWN-0"]["document_name"] == EXACT_DOCUMENT_NAME
    source_rows = by_reference["UNKNOWN-0"]["source_rows"]
    assert len(source_rows) == 1
    assert source_rows[0]["sheet"] == "Games"
    assert source_rows[0]["row_number"] == 3
    assert source_rows[0]["metadata"] == {"Location": "Section 2"}
    assert "Synthetic source alias; not an obligation" not in json.dumps(source_rows)
    assert by_reference["UNKNOWN-1"]["description"] == ""
    assert sections_by_id[by_reference["UNKNOWN-0"]["section_id"]] == "Games"
    assert sections_by_id[by_reference["ANNEX-0"]["section_id"]] == "Records"
    assert all(not row["requirements"] for row in listing.json()["items"])
    assert all(row["applicability"] == "unknown" and row["coverage_status"] == "unreviewed" for row in listing.json()["items"])
    started = await client.post(f"{BASE}/registers/{register['id']}/suggestions", headers=owner["headers"],
        json={"revision": reimported.json()["revision"], "mode": "local"})
    assert started.status_code == 200, started.text
    run = started.json()
    positions = []
    for _ in range(13):
        resumed = await client.post(f"{BASE}/suggestions/{run['id']}/resume", headers=owner["headers"], json={})
        assert resumed.status_code == 200, resumed.text
        run = resumed.json()
        positions.append(run["cursor"])
    assert positions == [10, 20, 30, 40, 50, 60, 70, 80, 90, 100, 110, 120, 126]
    assert run["status"] == "completed"
    assert len(run["suggestions"]) == 126
    assert all(suggestion["status"] == "pending" for suggestion in run["suggestions"])
    untouched = await client.get(f"{BASE}/registers/{register['id']}/expectations?limit=500", headers=owner["headers"])
    assert all(row["coverage_status"] == "unreviewed" and not row["requirements"] for row in untouched.json()["items"])


async def test_frozen_snapshot_keeps_prior_content_and_requires_current_export_access(
    client, documentation_actors, default_jurisdiction
):
    owner = documentation_actors["owner"]
    register = await create_register(client, owner, default_jurisdiction)
    row = await create_expectation(client, owner, register)
    evidence = await create_evidence(client, owner)
    grouped = await client.post(f"{BASE}/registers/{register['id']}/evidence-groups", headers=owner["headers"],
        json={"title": "Governance policy versions", "evidence_id": evidence["id"]})
    assert grouped.status_code == 201, grouped.text
    row = await link_evidence(client, owner, row, evidence)
    current = await client.get(f"{BASE}/registers/{register['id']}", headers=owner["headers"])
    snapshot = await client.post(f"{BASE}/registers/{register['id']}/snapshots", headers=owner["headers"],
        json={"revision": current.json()["revision"], "evidence_ids": [evidence["id"]]})
    assert snapshot.status_code == 201, snapshot.text
    endpoint = f"{BASE}/snapshots/{snapshot.json()['id']}/download"
    before = await client.get(endpoint, headers=owner["headers"])
    assert before.status_code == 200, before.text
    changed = await client.patch(f"{BASE}/expectations/{row['id']}", headers=owner["headers"],
        json={"revision": row["revision"], "title": "Later replacement title"})
    assert changed.status_code == 200, changed.text
    replacement = await create_evidence(client, owner, body="Later policy version replaces the original note content.")
    evidence_change = await client.post(f"{BASE}/evidence-groups/{grouped.json()['id']}/versions",
        headers=owner["headers"], json={"revision": grouped.json()["revision"], "evidence_id": replacement["id"]})
    assert evidence_change.status_code == 200, evidence_change.text
    after = await client.get(endpoint, headers=owner["headers"])
    assert after.status_code == 200, after.text
    assert after.content == before.content
    with zipfile.ZipFile(io.BytesIO(after.content)) as archive:
        metadata = json.loads(archive.read("documentation.json"))
        assert metadata["expectations"][0]["title"] == row["title"]
        assert metadata["expectations"][0]["coverage_status"] == "unreviewed"
        assert json.loads(archive.read("manifest.json"))["purpose"].startswith("Internal workflow handover")
    for name in ("peer", "foreign"):
        assert (await client.get(endpoint, headers=documentation_actors[name]["headers"])).status_code == 404
    access_path = f"/api/v1/access/documentation_register/{register['id']}"
    access = await client.get(access_path, headers=owner["headers"])
    assert access.status_code == 200, access.text
    granted = await client.put(access_path, headers=owner["headers"], json={
        "expected_revision": access.json()["revision"], "visibility": "secret",
        "grants": [{"subject_type": "user", "subject_id": str(documentation_actors["peer"]["user"].id),
                    "permissions": ["view", "export"]}], "reason": "Synthetic handover review"})
    assert granted.status_code == 200, granted.text
    assert (await client.get(endpoint, headers=documentation_actors["peer"]["headers"])).status_code == 200
    evidence_access_path = f"/api/v1/access/preparation_evidence/{evidence['id']}"
    evidence_access = await client.get(evidence_access_path, headers=owner["headers"])
    restricted = await client.put(evidence_access_path, headers=owner["headers"], json={
        "expected_revision": evidence_access.json()["revision"], "visibility": "secret",
        "grants": [{"subject_type": "user", "subject_id": str(documentation_actors["peer"]["user"].id),
                    "permissions": ["view", "export"]}], "reason": "Limit original source audience"})
    assert restricted.status_code == 200, restricted.text
    assert (await client.get(endpoint, headers=documentation_actors["peer"]["headers"])).status_code == 200
    source_revoked = await client.put(evidence_access_path, headers=owner["headers"], json={
        "expected_revision": restricted.json()["revision"], "visibility": "secret", "grants": [],
        "reason": "Source access withdrawn independently of register access"})
    assert source_revoked.status_code == 200, source_revoked.text
    assert (await client.get(endpoint, headers=documentation_actors["peer"]["headers"])).status_code == 404
    revoked = await client.put(access_path, headers=owner["headers"], json={
        "expected_revision": granted.json()["revision"], "visibility": "secret", "grants": [],
        "reason": "Synthetic handover complete"})
    assert revoked.status_code == 200, revoked.text
    assert (await client.get(endpoint, headers=documentation_actors["peer"]["headers"])).status_code == 404


@pytest.mark.parametrize("invalid", ["expectation", "evidence", "excerpt", "unsupported_claim", "incomplete"])
async def test_ai_transport_rejects_invented_ids_and_unsupported_excerpts(invalid):
    from types import SimpleNamespace
    import httpx
    from app.services.documentation_suggestions import suggest_batch

    payload = {"rows": [{"id": "expected-row", "title": "Oversight policy", "description": "Record board oversight"}],
               "requirements": [], "evidence": [{"id": "selected-evidence", "text": "Board oversight is recorded in section two."}]}
    suggestion = {"expectation_id": "expected-row", "evidence_id": "selected-evidence",
                  "excerpt": "Board oversight", "finding": "possible_support", "explanation": "Possible support in the selected excerpt."}
    if invalid == "expectation": suggestion["expectation_id"] = "invented-row"
    elif invalid == "evidence": suggestion["evidence_id"] = "unselected-evidence"
    elif invalid == "excerpt": suggestion["excerpt"] = "An invented quotation"
    elif invalid == "unsupported_claim":
        suggestion["evidence_id"] = None
        suggestion["excerpt"] = ""
    async def provider(request):
        submitted = json.loads(request.content)
        assert "Record board oversight" in submitted["messages"][0]["content"]
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {
            "content": json.dumps({"suggestions": [] if invalid == "incomplete" else [suggestion]})}}]})
    settings = SimpleNamespace(ai_provider="openai", ai_model="synthetic-model", openai_api_key="synthetic-key",
                               ai_request_timeout_seconds=5)
    with pytest.raises(ValueError, match="invalid expectation|unverified evidence|without a supplied|incomplete"):
        await suggest_batch(payload, settings, "ai", 0, transport=httpx.MockTransport(provider))


async def test_verified_ai_output_remains_pending_and_contains_no_coverage_decision():
    from types import SimpleNamespace
    import httpx
    from app.services.documentation_suggestions import suggest_batch

    payload = {"rows": [{"id": "row", "title": "Policy", "description": "Record oversight"}],
               "requirements": [], "evidence": [{"id": "evidence", "text": "Oversight is recorded."}]}
    async def provider(request):
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {
            "content": json.dumps({"suggestions": [{"expectation_id": "row", "evidence_id": "evidence",
                "excerpt": "Oversight is recorded.", "finding": "possible_support", "explanation": "Human scope review is required."}]})}}]})
    settings = SimpleNamespace(ai_provider="openai", ai_model="synthetic-model", openai_api_key="synthetic-key",
                               ai_request_timeout_seconds=5)
    result = await suggest_batch(payload, settings, "ai", 0, transport=httpx.MockTransport(provider))
    assert result[0]["status"] == "pending"
    assert result[0]["finding"] == "possible_support"
    assert "coverage_status" not in result[0]


async def test_snapshot_deduplicates_selected_bytes_and_detects_frozen_file_tampering(
    client, db_session, documentation_actors, default_jurisdiction
):
    import hashlib
    import uuid
    from app.models.preparation import PreparationEvidence
    from app.services.preparation_files import checked_evidence_path
    from app.services.documentation_export import snapshot_root

    owner = documentation_actors["owner"]
    register = await create_register(client, owner, default_jurisdiction)
    content = b"Synthetic immutable oversight evidence v1"
    evidence = []
    for title in ("Oversight policy", "Copy supplied for another expectation"):
        uploaded = await client.post("/api/v1/preparation/evidence/upload", headers=owner["headers"],
            data={"title": title, "visibility": "organisation"},
            files={"file": ("oversight.txt", content, "text/plain")})
        assert uploaded.status_code == 201, uploaded.text
        evidence.append(uploaded.json())
    snapshot = await client.post(f"{BASE}/registers/{register['id']}/snapshots", headers=owner["headers"],
        json={"revision": register["revision"], "evidence_ids": [row["id"] for row in evidence]})
    assert snapshot.status_code == 201, snapshot.text
    endpoint = f"{BASE}/snapshots/{snapshot.json()['id']}/download"
    archive_response = await client.get(endpoint, headers=owner["headers"])
    assert archive_response.status_code == 200, archive_response.text
    with zipfile.ZipFile(io.BytesIO(archive_response.content)) as archive:
        file_names = [name for name in archive.namelist() if name.startswith("evidence/")]
        assert len(file_names) == 1
        assert archive.read(file_names[0]) == content
        manifest = json.loads(archive.read("manifest.json"))
        file = next(row for row in manifest["files"] if row["path"].startswith("evidence/"))
        assert file["sha256"] == hashlib.sha256(content).hexdigest()
        assert file["size_bytes"] == len(content)
    source = await db_session.get(PreparationEvidence, uuid.UUID(evidence[0]["id"]))
    checked_evidence_path(source).write_bytes(b"Changed live source cannot rewrite the past export")
    repeated = await client.get(endpoint, headers=owner["headers"])
    assert repeated.status_code == 200, repeated.text
    assert repeated.content == archive_response.content
    fresh_snapshot = await client.post(f"{BASE}/registers/{register['id']}/snapshots", headers=owner["headers"],
        json={"revision": register["revision"], "evidence_ids": [row["id"] for row in evidence]})
    assert fresh_snapshot.status_code == 409, fresh_snapshot.text
    (snapshot_root(owner["user"].organization_id) / hashlib.sha256(content).hexdigest()).write_bytes(b"Tampered frozen file")
    corrupted = await client.get(endpoint, headers=owner["headers"])
    assert corrupted.status_code == 409, corrupted.text
    assert "integrity" in corrupted.text.lower(), corrupted.text


@pytest.mark.parametrize("change", ["archive", "settings", "withdraw_candidate"])
async def test_completed_ai_decision_rechecks_frozen_sources_but_allows_earlier_own_decisions(
    client, db_session, documentation_actors, default_jurisdiction, monkeypatch, change
):
    from types import SimpleNamespace
    import httpx
    from app.api import documentation_integrations as integrations
    from app.services.documentation_suggestions import suggest_batch
    from tests.test_preparation_requirements import make_source

    owner = documentation_actors["owner"]
    document, version, _ = await make_source(db_session, owner["user"], default_jurisdiction.id)
    document.status = version.status = "approved"
    await db_session.commit()
    register = await create_register(client, owner, default_jurisdiction)
    first = await create_expectation(client, owner, register, source_reference="1.1")
    second = await create_expectation(client, owner, register, title="Second review", source_reference="1.1")
    third = await create_expectation(client, owner, register, title="Third review", source_reference="1.1")
    evidence = await create_evidence(client, owner)
    runtime = SimpleNamespace(ai_provider="openai", ai_model="synthetic-model", openai_api_key="synthetic-key",
        ai_request_timeout_seconds=5, documentation_ai_revision=0, jev_enabled=False, typesafe_api_key="")

    async def settings(_):
        return runtime

    async def provider(request):
        message = json.loads(request.content)["messages"][0]["content"]
        supplied = json.loads(message.split("\n", 1)[1])
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {
            "content": json.dumps({"suggestions": [{"expectation_id": row["id"], "evidence_id": evidence["id"],
                "excerpt": "board oversight", "finding": "possible_support", "explanation": "Review the supplied scope."}
                for row in supplied["expectations"]]})}}]})

    async def batch(payload, settings, mode, cursor, **kwargs):
        return await suggest_batch(payload, settings, mode, cursor, transport=httpx.MockTransport(provider), **kwargs)

    monkeypatch.setattr(integrations, "assistance_settings", settings)
    monkeypatch.setattr(integrations, "suggest_batch", batch)
    current = (await client.get(f"{BASE}/registers/{register['id']}", headers=owner["headers"])).json()
    body = {"revision": current["revision"], "mode": "ai", "expectation_ids": [first["id"], second["id"], third["id"]],
        "evidence_ids": [evidence["id"]]}
    preview = await client.post(f"{BASE}/registers/{register['id']}/suggestions/preview", headers=owner["headers"], json=body)
    assert preview.status_code == 200, preview.text
    assert preview.json()["requirement_candidates"]
    started = await client.post(f"{BASE}/registers/{register['id']}/suggestions", headers=owner["headers"],
        json={**body, "input_fingerprint": preview.json()["input_fingerprint"], "external_processing_confirmed": True})
    assert started.status_code == 200, started.text
    run_id = started.json()["id"]
    resumed = await client.post(f"{BASE}/suggestions/{run_id}/resume", headers=owner["headers"],
        json={"external_processing_confirmed": True})
    assert resumed.status_code == 200, resumed.text
    assert resumed.json()["status"] == "completed"
    decisions = f"{BASE}/suggestions/{run_id}/decisions"
    applied = await client.post(decisions, headers=owner["headers"], json={"index": 0, "decision": "apply", "rationale": "First manual review"})
    assert applied.status_code == 200, applied.text
    another = await client.post(decisions, headers=owner["headers"], json={"index": 1, "decision": "apply", "rationale": "Second manual review"})
    assert another.status_code == 200, another.text
    if change == "archive":
        archived = await client.patch(f"/api/v1/preparation/evidence/{evidence['id']}", headers=owner["headers"], json={"archived": True})
        assert archived.status_code == 200, archived.text
    elif change == "settings":
        runtime.documentation_ai_revision += 1
    else:
        version.status = "draft"
        await db_session.commit()
    rejected = await client.post(decisions, headers=owner["headers"], json={"index": 2, "decision": "apply", "rationale": "Stale explanation"})
    assert rejected.status_code == 409, rejected.text
    dismissed = await client.post(decisions, headers=owner["headers"], json={"index": 2, "decision": "dismiss"})
    assert dismissed.status_code == 200, dismissed.text
    assert dismissed.json()["suggestions"][2]["status"] == "dismiss"
    target_id = resumed.json()["suggestions"][2]["expectation_id"]
    unchanged = await client.get(f"{BASE}/expectations/{target_id}", headers=owner["headers"])
    assert unchanged.status_code == 200, unchanged.text
    assert unchanged.json()["interpretation"] != "Stale explanation"


async def test_snapshot_freezes_independent_row_reviews_without_copying_collections_and_retains_both_audiences(
    client, db_session, documentation_actors, default_jurisdiction, review_process
):
    from sqlalchemy import func, select
    from app.models.documentation import DocumentationAssessmentEvidence, DocumentationCoverage
    from app.models.program import CertificationProjectRequirementBaseline
    from app.models.review import ReviewItem
    from tests.test_documentation_processes import (
        checked, bind_process, bind_targets, shared_replace, register_current,
        expectation_current, grant,
    )

    process = review_process
    owner, peer = process["owner"], documentation_actors["peer"]
    project = await checked(await client.post("/api/v1/certification-projects", headers=owner["headers"],
        json={"name": "Independent document certification", "jurisdiction_id": str(default_jurisdiction.id)}), 201)
    db_session.add(CertificationProjectRequirementBaseline(project_id=uuid.UUID(project["id"]),
        document_id=process["document"].id, requirement_set_version_id=process["version"].id))
    await db_session.commit()
    cycle = await checked(await client.post("/api/v1/review-cycles", headers=owner["headers"], json={
        "name": "Independent document review", "jurisdiction_id": str(default_jurisdiction.id),
        "certification_project_id": project["id"], "scope": "all"}), 201)
    item = await db_session.scalar(select(ReviewItem).where(ReviewItem.review_cycle_id == uuid.UUID(cycle["id"]),
        ReviewItem.requirement_id == process["requirement"].id))
    register = await create_register(client, owner, default_jurisdiction)
    await bind_process(client, owner, register, review_cycle_id=process["cycle"]["id"])
    first = await create_expectation(client, owner, register, title="Default review document")
    second = await create_expectation(client, owner, register, title="Independent review document")
    alias = await create_expectation(client, owner, register, title="Another scope of the independent document")

    async def choose(row, id_, clear=False):
        current = await expectation_current(client, owner, row)
        return await checked(await client.put(f"{BASE}/expectations/{row['id']}/review-process", headers=owner["headers"],
            json={"revision": current["revision"], "review_cycle_id": id_, "clear_targets": clear}))

    second = await choose(second, cycle["id"])
    # No target or snapshot references exist yet: the row's review pointer itself
    # must prevent destructive deletion, rather than relying on a later FK error.
    deletion = await client.delete(f"/api/v1/review-cycles/{cycle['id']}", headers=owner["headers"])
    assert deletion.status_code == 409, deletion.text
    alias = await choose(alias, cycle["id"])
    first = await bind_targets(client, owner, first, [process["target"]])
    other_target = {**process["source_identity"], "review_item_id": str(item.id)}
    second = await bind_targets(client, owner, second, [other_target])
    alias = await bind_targets(client, owner, alias, [other_target])
    contents = [b"Synthetic default review document bytes", b"Synthetic independent review document bytes"]
    files = []
    for index, content in enumerate(contents):
        files.append(await checked(await client.post("/api/v1/preparation/evidence/upload", headers=owner["headers"],
            data={"title": f"Canonical review document {index}", "visibility": "organisation"},
            files={"file": (f"document-{index}.txt", content, "text/plain")}), 201))
    await shared_replace(client, owner, process, [files[0]["id"]])
    other_process = {"shared_path": f"/api/v1/review-cycles/{cycle['id']}/items/{item.id}/shared-evidence"}
    await shared_replace(client, owner, other_process, [files[1]["id"]])
    for row in (first, second):
        current = await expectation_current(client, owner, row)
        await checked(await client.post(f"{BASE}/expectations/{row['id']}/review", headers=owner["headers"],
            json={"revision": current["revision"], "status": "sufficient", "rationale": "Reviewed this document's own scope"}))
    current = await register_current(client, owner, register)
    snapshot = await checked(await client.post(f"{BASE}/registers/{register['id']}/snapshots", headers=owner["headers"],
        json={"revision": current["revision"]}), 201)
    for project_id, expected_rows in ((process["project"]["id"], {first["id"]}), (project["id"], {second["id"], alias["id"]})):
        scoped = await checked(await client.post(f"{BASE}/registers/{register['id']}/snapshots", headers=owner["headers"],
            json={"revision": current["revision"], "project_id": project_id}), 201)
        scoped_download = await client.get(f"{BASE}/snapshots/{scoped['id']}/download", headers=owner["headers"])
        assert scoped_download.status_code == 200, scoped_download.text
        with zipfile.ZipFile(io.BytesIO(scoped_download.content)) as archive:
            scoped_data = json.loads(archive.read("documentation.json"))
            assert {row["id"] for row in scoped_data["expectations"]} == expected_rows
            assert len(scoped_data["frozen_evidence"]) == 1
    await grant(client, owner, "documentation_register", register["id"], peer["user"], ["view", "export"])
    download = f"{BASE}/snapshots/{snapshot['id']}/download"
    frozen = await client.get(download, headers=peer["headers"])
    assert frozen.status_code == 200, frozen.text
    with zipfile.ZipFile(io.BytesIO(frozen.content)) as archive:
        data = json.loads(archive.read("documentation.json"))
        rows = {row["id"]: row for row in data["expectations"]}
        assert rows[first["id"]]["process_targets"][0]["review_cycle_id"] == process["cycle"]["id"]
        assert rows[second["id"]]["process_targets"][0]["review_cycle_id"] == cycle["id"]
        assert rows[alias["id"]]["process_targets"][0]["review_cycle_id"] == cycle["id"]
        assert rows[first["id"]]["effective_coverage_status"] == "sufficient"
        assert rows[second["id"]]["effective_coverage_status"] == "sufficient"
        assert rows[alias["id"]]["effective_coverage_status"] == "unreviewed"
        assert len(data["frozen_evidence"]) == 2
        for content in contents:
            assert sum(archive.read(name) == content for name in archive.namelist()) == 1
        assert {source["review_cycle_id"] for source in data["process_sources"] if source.get("review_cycle_id")} == {process["cycle"]["id"], cycle["id"]}
    assert await db_session.scalar(select(func.count()).select_from(DocumentationAssessmentEvidence)) == 2
    assert await db_session.scalar(select(func.count()).select_from(DocumentationCoverage)) == 0
    await bind_targets(client, owner, first, [])
    await choose(second, None, True)
    await choose(alias, None, True)
    await bind_process(client, owner, register)
    policy_path = f"/api/v1/access/certification_project/{project['id']}"
    policy = await checked(await client.get(policy_path, headers=owner["headers"]))
    await checked(await client.put(policy_path, headers=owner["headers"], json={"expected_revision": policy["revision"],
        "visibility": "secret", "grants": [], "reason": "Independent original review audience withdrawn"}))
    assert (await client.get(download, headers=peer["headers"])).status_code == 404
    assert (await client.get(download, headers=owner["headers"])).content == frozen.content
