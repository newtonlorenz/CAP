"""Source question import boundaries: privacy, version choice and immutable provenance."""

import io
import uuid
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import Document, ExtractedRequirement, Jurisdiction, Organization, Requirement, User
from app.models.extraction import ExtractionRun
from app.models.preparation import PreparationTemplate
from app.models.program import RequirementSetVersion
from app.services.auth import create_access_token

BASE = "/api/v1/preparation/templates"


def headers(user):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id)})}"}


@pytest.fixture
async def source_actors(db_session):
    organisations = [
        Organization(code="import-source", name="Import source"),
        Organization(code="import-other", name="Other source"),
    ]
    db_session.add_all(organisations)
    await db_session.flush()
    users = {}
    for name, role, organisation in [
        ("manager", "manager", organisations[0]),
        ("contributor", "contributor", organisations[0]),
        ("other", "manager", organisations[1]),
        ("unscoped", "manager", None),
    ]:
        users[name] = User(
            email=f"source-{name}@example.test",
            full_name=name,
            role=role,
            active=True,
            password_hash="unused",
            organization_id=organisation.id if organisation else None,
        )
        db_session.add(users[name])
    await db_session.commit()
    return users


async def make_source(db_session, user, jurisdiction_id, *, extracted=False):
    document = Document(
        organization_id=user.organization_id,
        jurisdiction_id=jurisdiction_id,
        name="Annex A",
        filename="Annex A.pdf",
        file_path="source.pdf",
        document_type="annex",
        status="extracted" if extracted else "draft",
        uploaded_by=user.id,
    )
    db_session.add(document)
    await db_session.flush()
    version = RequirementSetVersion(
        organization_id=user.organization_id,
        document_id=document.id,
        version_number=1,
        status="draft",
        is_current=False,
        created_by=user.id,
    )
    db_session.add(version)
    await db_session.flush()
    if extracted:
        run = ExtractionRun(
            document_id=document.id,
            status="completed",
            ai_provider="deterministic",
            ai_model="rule_based",
        )
        db_session.add(run)
        await db_session.flush()
        document.current_extraction_id = run.id
        await db_session.commit()
        return document, version, run
    requirements = [
        Requirement(
            organization_id=user.organization_id,
            document_id=document.id,
            jurisdiction_id=jurisdiction_id,
            requirement_set_version_id=version.id,
            reference_id="1",
            title="Applicant",
            text="Applicant information",
            requirement_type="informational",
            sort_order=0,
        ),
        Requirement(
            organization_id=user.organization_id,
            document_id=document.id,
            jurisdiction_id=jurisdiction_id,
            requirement_set_version_id=version.id,
            reference_id="1.1",
            title="Company name",
            text="Enter your registered name",
            requirement_type="mandatory",
            sort_order=1,
        ),
    ]
    db_session.add_all(requirements)
    await db_session.flush()
    requirements[1].parent_id = requirements[0].id
    await db_session.commit()
    return document, version, requirements


async def preview(client, user, document, jurisdiction_id, **params):
    return await client.get(
        f"{BASE}/requirements-preview",
        headers=headers(user),
        params={"document_id": str(document.id), "jurisdiction_id": str(jurisdiction_id), **params},
    )


def template_body(fields):
    return {"name": "Imported Annex A", "kind": "licence_application", "fields": fields}


async def test_draft_questions_preview_save_and_case_keep_source_snapshot(
    client, db_session, source_actors, default_jurisdiction
):
    user = source_actors["manager"]
    document, version, requirements = await make_source(db_session, user, default_jurisdiction.id)
    result = await preview(client, user, document, default_jurisdiction.id)
    assert result.status_code == 200, result.text
    body = result.json()
    assert body["source"]["version_status"] == "draft"
    assert body["total"] == 2
    assert any("not been approved" in warning for warning in body["warnings"])
    assert [field["label"] for field in body["fields"]] == [
        "1 Applicant information",
        "1.1 Enter your registered name",
    ]
    assert [field["required"] for field in body["fields"]] == [False, True]
    assert all(field["type"] == "multiline" for field in body["fields"])
    assert body["fields"][1]["section"] == "1 Applicant"
    assert not any("value" in field or "current_status" in field for field in body["fields"])
    assert await db_session.scalar(select(func.count()).select_from(PreparationTemplate)) == 0
    # Preview metadata cannot be used to fabricate a different source identity.
    body["fields"][1]["source"]["document_name"] = "Forged source name"
    saved = await client.post(BASE, headers=headers(user), json=template_body(body["fields"]))
    assert saved.status_code == 201, saved.text
    saved_body = saved.json()
    assert saved_body["fields"][1]["source"]["document_name"] == "Annex A"
    assert saved_body["fields"][1]["source"]["requirement_id"] == str(requirements[1].id)
    assert saved_body["fields"][1]["source"]["requirement_set_version_id"] == str(version.id)
    case = await client.post(
        "/api/v1/preparation/cases",
        headers=headers(user),
        json={
            "template_id": saved_body["id"],
            "jurisdiction_id": str(default_jurisdiction.id),
            "name": "Annex A response",
        },
    )
    assert case.status_code == 201, case.text
    snapshot = case.json()["fields"]
    requirements[1].text = "Changed original wording"
    await db_session.commit()
    stale = await client.post(BASE, headers=headers(user), json=template_body(saved_body["fields"]))
    assert stale.status_code == 409
    # Editing an existing template preserves its old source record, while cases
    # remain pinned to their complete immutable field snapshot.
    edited = await client.patch(
        f"{BASE}/{saved_body['id']}",
        headers=headers(user),
        json={
            "expected_revision": 1,
            "fields": saved_body["fields"],
            "name": "Template renamed",
        },
    )
    assert edited.status_code == 200, edited.text
    reloaded = await client.get(
        f"/api/v1/preparation/cases/{case.json()['id']}", headers=headers(user)
    )
    assert reloaded.json()["fields"] == snapshot


async def test_source_privacy_role_jurisdiction_and_version_boundaries(
    client, db_session, source_actors, default_jurisdiction
):
    user = source_actors["manager"]
    document, _, _ = await make_source(db_session, user, default_jurisdiction.id)
    for name, expected in [("other", 404), ("unscoped", 404), ("contributor", 403)]:
        assert (
            await preview(client, source_actors[name], document, default_jurisdiction.id)
        ).status_code == expected
    assert (
        await client.get(
            f"{BASE}/requirements-preview",
            params={
                "document_id": str(document.id),
                "jurisdiction_id": str(default_jurisdiction.id),
            },
        )
    ).status_code == 401
    assert (await preview(client, user, document, uuid.uuid4())).status_code == 422
    other_document, other_version, _ = await make_source(db_session, user, default_jurisdiction.id)
    assert (
        await preview(
            client, user, document, default_jurisdiction.id, version_id=str(other_version.id)
        )
    ).status_code == 404
    visible = (await preview(client, user, document, default_jurisdiction.id)).json()
    hidden = await client.post(
        BASE, headers=headers(source_actors["other"]), json=template_body(visible["fields"])
    )
    assert hidden.status_code == 404
    forged = deepcopy(visible["fields"])
    forged[1]["source"]["document_id"] = str(other_document.id)
    assert (
        await client.post(BASE, headers=headers(user), json=template_body(forged))
    ).status_code == 404


async def test_selected_approved_and_submitted_versions_are_kept_separate(
    client, db_session, source_actors, default_jurisdiction
):
    user = source_actors["manager"]
    document, approved, _ = await make_source(db_session, user, default_jurisdiction.id)
    approved.status = "approved"
    approved.is_current = True
    submitted = RequirementSetVersion(
        organization_id=user.organization_id,
        document_id=document.id,
        version_number=2,
        status="pending_approval",
        created_by=user.id,
    )
    db_session.add(submitted)
    await db_session.flush()
    db_session.add(
        Requirement(
            organization_id=user.organization_id,
            document_id=document.id,
            jurisdiction_id=default_jurisdiction.id,
            requirement_set_version_id=submitted.id,
            reference_id="2.1",
            text="New submitted question",
            requirement_type="mandatory",
        )
    )
    await db_session.commit()
    latest = (await preview(client, user, document, default_jurisdiction.id)).json()
    assert latest["source"]["version_status"] == "pending_approval"
    assert latest["total"] == 1
    assert latest["fields"][0]["label"] == "2.1 New submitted question"
    older = (
        await preview(client, user, document, default_jurisdiction.id, version_id=str(approved.id))
    ).json()
    assert older["source"]["version_status"] == "approved"
    assert older["total"] == 2
    assert older["fields"][1]["label"] == "1.1 Enter your registered name"
    assert not any("not been approved" in warning for warning in older["warnings"])


@pytest.mark.parametrize("target", ["standalone", "application"])
async def test_source_jurisdiction_cannot_be_bypassed_when_instantiating_a_form(
    client, db_session, source_actors, default_jurisdiction, target
):
    user = source_actors["manager"]
    document, _, _ = await make_source(db_session, user, default_jurisdiction.id)
    source = (await preview(client, user, document, default_jurisdiction.id)).json()
    template = await client.post(BASE, headers=headers(user), json=template_body(source["fields"]))
    assert template.status_code == 201, template.text
    other_jurisdiction = Jurisdiction(code="se", name="Sweden", active=True)
    db_session.add(other_jurisdiction)
    await db_session.commit()
    if target == "standalone":
        response = await client.post(
            "/api/v1/preparation/cases",
            headers=headers(user),
            json={
                "template_id": template.json()["id"],
                "jurisdiction_id": str(other_jurisdiction.id),
                "name": "Wrong space",
            },
        )
    else:
        application = await client.post(
            "/api/v1/applications",
            headers=headers(user),
            json={
                "name": "Swedish application",
                "scope": "licence",
                "jurisdiction_id": str(other_jurisdiction.id),
            },
        )
        assert application.status_code == 201, application.text
        response = await client.post(
            f"/api/v1/applications/{application.json()['id']}/components",
            headers=headers(user),
            json={
                "expected_revision": 1,
                "name": "Wrong space",
                "kind": "form",
                "template_id": template.json()["id"],
            },
        )
    assert response.status_code == 422, response.text
    assert "jurisdiction differs" in response.json()["detail"]


async def test_completed_extraction_with_no_final_requirements_supports_paged_subset(
    client, db_session, source_actors, default_jurisdiction
):
    user = source_actors["manager"]
    document, version, run = await make_source(
        db_session, user, default_jurisdiction.id, extracted=True
    )
    old_run = ExtractionRun(
        document_id=document.id, status="completed", ai_provider="deterministic", ai_model="old"
    )
    db_session.add(old_run)
    await db_session.flush()
    rows = [
        ExtractedRequirement(
            document_id=document.id,
            extraction_run_id=run.id,
            reference_id=f"A.{index + 1}",
            text=f"Form question {index + 1}",
            original_text="Original PDF text",
            status="pending",
            requirement_type="mandatory",
            page_number=1,
            sort_order=index,
        )
        for index in range(205)
    ]
    rows.extend(
        [
            ExtractedRequirement(
                document_id=document.id,
                extraction_run_id=run.id,
                reference_id="rejected",
                text="Rejected fragment",
                original_text="Rejected fragment",
                status="rejected",
                page_number=1,
            ),
            ExtractedRequirement(
                document_id=document.id,
                extraction_run_id=old_run.id,
                reference_id="old",
                text="Old extraction",
                original_text="Old extraction",
                status="pending",
                page_number=1,
            ),
        ]
    )
    db_session.add_all(rows)
    await db_session.commit()
    first = await preview(client, user, document, default_jurisdiction.id)
    assert first.status_code == 200, first.text
    first_body = first.json()
    assert first_body["total"] == 205
    assert len(first_body["fields"]) == 100
    assert first_body["source"]["kind"] == "extracted_requirements"
    assert first_body["source"]["version_status"] is None
    assert first_body["source"]["extraction_status"] == "completed"
    assert any("before requirement-set approval" in warning for warning in first_body["warnings"])
    # Choosing a specific empty version must not substitute different PDF text.
    explicit = (
        await preview(client, user, document, default_jurisdiction.id, version_id=str(version.id))
    ).json()
    assert explicit["total"] == 0
    assert explicit["fields"] == []
    assert explicit["source"]["kind"] == "requirements"
    assert explicit["source"]["version_id"] == str(version.id)
    last = (await preview(client, user, document, default_jurisdiction.id, skip=200)).json()
    assert [field["label"] for field in last["fields"]] == [
        f"A.{index} Form question {index}" for index in range(201, 206)
    ]
    source = last["fields"][0]["source"]
    assert source["kind"] == "extracted_requirement"
    assert source["requirement_id"] is None
    assert source["extracted_requirement_id"] == str(rows[200].id)
    assert source["extraction_run_id"] == str(run.id)
    assert source["extraction_review_status"] == "pending"
    saved = await client.post(BASE, headers=headers(user), json=template_body(last["fields"]))
    assert saved.status_code == 201, saved.text
    assert len(saved.json()["fields"]) == 5
    rows[200].status = "rejected"
    await db_session.commit()
    assert (
        await client.post(BASE, headers=headers(user), json=template_body(last["fields"]))
    ).status_code == 409


@pytest.mark.parametrize("status", ["running", "failed", "cancelled"])
async def test_unfinished_extraction_is_not_an_importable_source(
    client, db_session, source_actors, default_jurisdiction, status
):
    user = source_actors["manager"]
    document, _, run = await make_source(db_session, user, default_jurisdiction.id, extracted=True)
    run.status = status
    await db_session.commit()
    result = await preview(client, user, document, default_jurisdiction.id)
    assert result.status_code == 409
    assert "finish" in result.json()["detail"]


async def test_long_source_text_is_preserved_or_explicitly_unavailable(
    client, db_session, source_actors, default_jurisdiction
):
    user = source_actors["manager"]
    document, _, requirements = await make_source(db_session, user, default_jurisdiction.id)
    requirements[0].text = "Original wording\n" * 300
    requirements[1].text = "x" * 10001
    await db_session.commit()
    result = (await preview(client, user, document, default_jurisdiction.id)).json()
    assert result["total"] == 2
    assert len(result["fields"]) == 1
    assert result["fields"][0]["label"] == "1 Applicant"
    assert result["fields"][0]["help_text"] == "Original wording\n" * 300
    assert result["unavailable"][0]["requirement_id"] == str(requirements[1].id)
    assert "10,000" in result["unavailable"][0]["reason"]
    saved = await client.post(BASE, headers=headers(user), json=template_body(result["fields"]))
    assert saved.status_code == 201, saved.text


async def test_import_template_download_is_a_usable_workbook_with_import_role_boundary(
    client, source_actors
):
    response = await client.get(
        f"{BASE}/import-template", headers=headers(source_actors["manager"])
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        'filename="preparation-question-template.xlsx"' in response.headers["content-disposition"]
    )
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert workbook.sheetnames[0] == "Questions"
        assert list(next(workbook["Questions"].values)) == [
            "Section",
            "Question number",
            "Question text",
            "Answer type",
            "Choices",
            "Required",
            "Instructions",
        ]
    finally:
        workbook.close()
    assert (
        await client.get(f"{BASE}/import-template", headers=headers(source_actors["contributor"]))
    ).status_code == 403
