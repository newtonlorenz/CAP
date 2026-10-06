import csv
import io
import re
import uuid
from datetime import datetime, timedelta, timezone

import pytest
import pdfplumber

try:
    from openpyxl import load_workbook  # type: ignore
except ModuleNotFoundError:  # pragma: no cover
    load_workbook = None  # type: ignore

from app.models import (
    AuditLog,
    Document,
    Requirement,
    RequirementStatus,
    ReviewCycle,
    ReviewItem,
    ReviewItemEvidenceFile,
    User,
)
from app.services.auth import create_access_token, hash_password


@pytest.fixture
async def test_user(db_session):
    user = User(
        email="test@example.com",
        password_hash=hash_password("testpass"),
        full_name="Test User",
        role="admin",
    )
    db_session.add(user)
    await db_session.commit()
    await db_session.refresh(user)
    return user


@pytest.fixture
async def test_document(db_session, test_user, default_jurisdiction):
    doc = Document(
        jurisdiction_id=default_jurisdiction.id,
        filename="policy.pdf",
        name="Test Policy",
        document_type="policy",
        file_path="/tmp/policy.pdf",
        status="approved",
        uploaded_by=test_user.id,
    )
    db_session.add(doc)
    await db_session.commit()
    await db_session.refresh(doc)
    return doc


@pytest.fixture
async def test_requirement(db_session, test_user, test_document, default_jurisdiction):
    req = Requirement(
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-001",
        text="Test requirement",
        requirement_type="mandatory",
        active=True,
        version=1,
        sort_order=0,
    )
    db_session.add(req)
    await db_session.commit()
    await db_session.refresh(req)
    return req


@pytest.fixture
async def review_cycle(db_session, test_user, default_jurisdiction):
    cycle = ReviewCycle(
        name="Test Cycle",
        jurisdiction_id=default_jurisdiction.id,
        scope="all",
        created_by=test_user.id,
    )
    db_session.add(cycle)
    await db_session.commit()
    await db_session.refresh(cycle)
    return cycle


@pytest.fixture
async def review_item(db_session, review_cycle, test_requirement):
    item = ReviewItem(
        review_cycle_id=review_cycle.id,
        requirement_id=test_requirement.id,
        review_status="pending",
    )
    db_session.add(item)
    await db_session.commit()
    await db_session.refresh(item)
    return item


@pytest.fixture
async def review_item_evidence(db_session, review_item, test_user):
    evidence = ReviewItemEvidenceFile(
        review_item_id=review_item.id,
        filename="evidence.txt",
        file_path="/tmp/evidence.txt",
        description="Evidence file",
        uploaded_by=test_user.id,
    )
    db_session.add(evidence)
    await db_session.commit()
    await db_session.refresh(evidence)
    return evidence


@pytest.fixture
async def requirement_status(db_session, test_requirement, test_user):
    status = RequirementStatus(
        requirement_id=test_requirement.id,
        status="in_progress",
        assigned_to=test_user.id,
        comment="Evidence noted",
        changed_by=test_user.id,
    )
    db_session.add(status)
    await db_session.commit()
    await db_session.refresh(status)
    return status


@pytest.fixture
async def audit_logs(db_session, test_user):
    for i in range(5):
        log = AuditLog(
            user_id=test_user.id,
            user_name=test_user.full_name,
            action="status_change",
            entity_type="requirement",
            entity_id=str(uuid.uuid4()),
            old_value="not_started",
            new_value="in_progress",
        )
        db_session.add(log)
    await db_session.commit()


@pytest.fixture
def auth_headers(test_user):
    token = create_access_token({"sub": str(test_user.id)})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
async def contributor_headers(db_session):
    contributor = User(
        email="contributor@example.com",
        password_hash=hash_password("testpass"),
        full_name="Contributor User",
        role="contributor",
    )
    db_session.add(contributor)
    await db_session.commit()
    token = create_access_token({"sub": str(contributor.id)})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
async def contributor_user(db_session):
    contributor = User(
        email="report-contributor@example.com",
        password_hash=hash_password("testpass"),
        full_name="Contributor User",
        role="contributor",
    )
    db_session.add(contributor)
    await db_session.commit()
    await db_session.refresh(contributor)
    return contributor


@pytest.fixture
def contributor_user_headers(contributor_user):
    token = create_access_token({"sub": str(contributor_user.id)})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
async def assigned_reviewer_user(db_session):
    reviewer = User(
        email="assigned-reviewer@example.com",
        password_hash=hash_password("testpass"),
        full_name="Assigned Reviewer",
        role="assigned_reviewer",
    )
    db_session.add(reviewer)
    await db_session.commit()
    await db_session.refresh(reviewer)
    return reviewer


@pytest.fixture
def assigned_reviewer_headers(assigned_reviewer_user):
    token = create_access_token({"sub": str(assigned_reviewer_user.id)})
    return {"Authorization": f"Bearer {token}"}


async def _create_requirement(
    db_session,
    *,
    jurisdiction_id,
    document_id,
    reference_id: str,
    text: str,
    requirement_type: str,
    sort_order: int,
    title=None,
):
    requirement = Requirement(
        jurisdiction_id=jurisdiction_id,
        document_id=document_id,
        reference_id=reference_id,
        title=title,
        text=text,
        requirement_type=requirement_type,
        active=True,
        version=1,
        sort_order=sort_order,
    )
    db_session.add(requirement)
    await db_session.flush()
    return requirement


async def _create_requirement_status(
    db_session,
    *,
    requirement_id,
    user_id,
    status: str,
    assigned_to=None,
    comment: str = "",
):
    requirement_status = RequirementStatus(
        requirement_id=requirement_id,
        status=status,
        assigned_to=assigned_to,
        comment=comment,
        changed_by=user_id,
    )
    db_session.add(requirement_status)
    await db_session.flush()
    return requirement_status


async def _create_review_item(
    db_session,
    *,
    cycle_id,
    requirement_id,
    review_status: str,
    assessment_status=None,
    assigned_reviewer_id=None,
    responsible_user_id=None,
    reviewer_id=None,
    review_evidence=None,
    reviewed_at=None,
    jira_issue_key=None,
    jira_status=None,
    jira_assignee=None,
    jira_issue_url=None,
):
    review_item = ReviewItem(
        review_cycle_id=cycle_id,
        requirement_id=requirement_id,
        review_status=review_status,
        assessment_status=assessment_status,
        assigned_reviewer_id=assigned_reviewer_id,
        responsible_user_id=responsible_user_id,
        reviewer_id=reviewer_id,
        review_evidence=review_evidence,
        reviewed_at=reviewed_at,
        jira_issue_key=jira_issue_key,
        jira_status=jira_status,
        jira_assignee=jira_assignee,
        jira_issue_url=jira_issue_url,
    )
    db_session.add(review_item)
    await db_session.flush()
    return review_item


async def _create_evidence_file(
    db_session,
    *,
    review_item_id,
    uploaded_by,
    filename: str,
):
    evidence = ReviewItemEvidenceFile(
        review_item_id=review_item_id,
        filename=filename,
        file_path=f"/tmp/{filename}",
        description="Evidence file",
        uploaded_by=uploaded_by,
    )
    db_session.add(evidence)
    await db_session.flush()
    return evidence


def _read_csv_rows(content: str) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(content)))


def _worksheet_rows(sheet) -> list[dict[str, str]]:
    headers = [
        str(cell.value) if cell.value is not None else ""
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    ]
    rows: list[dict[str, str]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = ["" if value is None else str(value) for value in row]
        rows.append(dict(zip(headers, values)))
    return rows


def _summary_table_rows(sheet, heading: str) -> dict[str, str]:
    rows: dict[str, str] = {}
    in_table = False
    for row in sheet.iter_rows(values_only=True):
        first = "" if not row or row[0] is None else str(row[0])
        second = ""
        if row and len(row) > 1 and row[1] is not None:
            second = str(row[1])
        if first == heading and second == "Count":
            in_table = True
            continue
        if not in_table:
            continue
        if not first:
            break
        rows[first] = second
    return rows


@pytest.fixture
async def action_gap_cycle_data(
    db_session,
    test_user,
    contributor_user,
    assigned_reviewer_user,
    test_document,
    default_jurisdiction,
):
    cycle = ReviewCycle(
        name="Operational Gap Cycle",
        jurisdiction_id=default_jurisdiction.id,
        scope="all",
        created_by=test_user.id,
        deadline=datetime.now(timezone.utc) - timedelta(days=1),
    )
    db_session.add(cycle)
    await db_session.flush()

    req_no_action = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-100",
        title="Completed control",
        text="Completed requirement text",
        requirement_type="mandatory",
        sort_order=0,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_no_action.id,
        user_id=test_user.id,
        status="evidenced",
        assigned_to=test_user.id,
        comment="Completed",
    )
    item_no_action = await _create_review_item(
        db_session,
        cycle_id=cycle.id,
        requirement_id=req_no_action.id,
        review_status="confirmed",
        assessment_status="evidenced",
        assigned_reviewer_id=test_user.id,
        reviewer_id=test_user.id,
        review_evidence="Evidence recorded",
        reviewed_at=datetime.now(timezone.utc),
    )
    await _create_evidence_file(
        db_session,
        review_item_id=item_no_action.id,
        uploaded_by=test_user.id,
        filename="completed-evidence.txt",
    )

    req_unassigned = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-110",
        title="Needs assignment",
        text="Unassigned requirement text",
        requirement_type="mandatory",
        sort_order=1,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_unassigned.id,
        user_id=test_user.id,
        status="not_started",
    )
    await _create_review_item(
        db_session,
        cycle_id=cycle.id,
        requirement_id=req_unassigned.id,
        review_status="pending",
        assessment_status="not_started",
        assigned_reviewer_id=test_user.id,
    )

    req_blocked = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-120",
        title="Blocked control",
        text="Blocked requirement text",
        requirement_type="mandatory",
        sort_order=2,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_blocked.id,
        user_id=test_user.id,
        status="blocked",
        assigned_to=contributor_user.id,
        comment="Waiting on dependency",
    )
    await _create_review_item(
        db_session,
        cycle_id=cycle.id,
        requirement_id=req_blocked.id,
        review_status="escalated",
        assessment_status="blocked",
        responsible_user_id=test_user.id,
    )

    req_missing_evidence = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-130",
        title="Needs evidence",
        text="Missing evidence requirement text",
        requirement_type="mandatory",
        sort_order=3,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_missing_evidence.id,
        user_id=test_user.id,
        status="in_progress",
        assigned_to=contributor_user.id,
        comment="Owner assigned",
    )
    await _create_review_item(
        db_session,
        cycle_id=cycle.id,
        requirement_id=req_missing_evidence.id,
        review_status="pending",
        assessment_status="in_progress",
        assigned_reviewer_id=assigned_reviewer_user.id,
        responsible_user_id=contributor_user.id,
    )

    req_pending_review = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-140",
        title="Review still open",
        text="Pending review requirement text",
        requirement_type="mandatory",
        sort_order=4,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_pending_review.id,
        user_id=test_user.id,
        status="in_progress",
        assigned_to=contributor_user.id,
        comment="Owner assigned",
    )
    await _create_review_item(
        db_session,
        cycle_id=cycle.id,
        requirement_id=req_pending_review.id,
        review_status="pending",
        assessment_status="in_progress",
        review_evidence="Draft evidence note",
        jira_issue_key="OPS-14",
        jira_status="In Progress",
        jira_assignee=contributor_user.full_name,
        jira_issue_url="https://jira.example.com/browse/OPS-14",
    )

    req_info = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-150",
        title="Informational control",
        text="Informational requirement text",
        requirement_type="informational",
        sort_order=5,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_info.id,
        user_id=test_user.id,
        status="not_applicable",
    )
    await _create_review_item(
        db_session,
        cycle_id=cycle.id,
        requirement_id=req_info.id,
        review_status="confirmed",
        assessment_status="not_applicable",
    )

    await db_session.commit()
    await db_session.refresh(cycle)
    return {
        "cycle": cycle,
        "assigned_reviewer_reference": req_missing_evidence.reference_id,
        "other_reference": req_pending_review.reference_id,
    }


@pytest.fixture
async def master_gap_data(
    db_session,
    test_user,
    contributor_user,
    test_document,
    default_jurisdiction,
):
    req_done = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-M1",
        title="Master complete",
        text="Complete master requirement",
        requirement_type="mandatory",
        sort_order=10,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_done.id,
        user_id=test_user.id,
        status="evidenced",
        assigned_to=test_user.id,
        comment="Done",
    )

    req_unassigned = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-M2",
        title="Master unassigned",
        text="Unassigned master requirement",
        requirement_type="mandatory",
        sort_order=11,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_unassigned.id,
        user_id=test_user.id,
        status="not_started",
        comment="Waiting for owner",
    )

    req_blocked = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-M3",
        title="Master blocked",
        text="Blocked master requirement",
        requirement_type="mandatory",
        sort_order=12,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_blocked.id,
        user_id=test_user.id,
        status="blocked",
        assigned_to=contributor_user.id,
        comment="Blocked",
    )

    req_info = await _create_requirement(
        db_session,
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        reference_id="REQ-M4",
        title="Master informational",
        text="Informational master requirement",
        requirement_type="informational",
        sort_order=13,
    )
    await _create_requirement_status(
        db_session,
        requirement_id=req_info.id,
        user_id=test_user.id,
        status="not_applicable",
    )

    await db_session.commit()
    return {
        "done": req_done.reference_id,
        "unassigned": req_unassigned.reference_id,
        "blocked": req_blocked.reference_id,
        "informational": req_info.reference_id,
    }


async def test_download_compliance_summary(
    client, auth_headers, test_requirement, default_jurisdiction
):
    response = await client.get(
        f"/api/v1/reports/compliance-summary?jurisdiction_id={default_jurisdiction.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "compliance-summary" in response.headers.get("content-disposition", "")


async def test_download_detailed_report(
    client, auth_headers, test_requirement, default_jurisdiction
):
    response = await client.get(
        f"/api/v1/reports/detailed?jurisdiction_id={default_jurisdiction.id}", headers=auth_headers
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"


async def test_download_gap_analysis(client, auth_headers, test_requirement, default_jurisdiction):
    response = await client.get(
        f"/api/v1/reports/gap-analysis?jurisdiction_id={default_jurisdiction.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    content = response.text
    assert "Requires Evidence" in content
    assert "Responsible for Requirement Name" in content
    assert "REQ-001" in content


async def test_cycle_scoped_compliance_summary(client, auth_headers, review_cycle, review_item):
    response = await client.get(
        f"/api/v1/reports/compliance-summary?review_cycle_id={review_cycle.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"

    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert f"Review: {review_cycle.name}" in text
    assert "Scope: mandatory and recommended requirements." in text


async def test_cycle_scoped_detailed_report(client, auth_headers, review_cycle, review_item):
    response = await client.get(
        f"/api/v1/reports/detailed?review_cycle_id={review_cycle.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"


async def test_cycle_scoped_gap_analysis(client, auth_headers, review_cycle, review_item):
    response = await client.get(
        f"/api/v1/reports/gap-analysis?review_cycle_id={review_cycle.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"


async def test_cycle_scoped_gap_analysis_rules(
    client,
    auth_headers,
    action_gap_cycle_data,
):
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        f"/api/v1/reports/gap-analysis?review_cycle_id={cycle.id}&format=csv",
        headers=auth_headers,
    )
    assert response.status_code == 200
    rows = _read_csv_rows(response.text)
    rows_by_ref = {row["Reference ID"]: row for row in rows}
    assert set(rows_by_ref) == {"REQ-100", "REQ-110", "REQ-120", "REQ-130", "REQ-140"}
    assert rows_by_ref["REQ-100"]["Requires Evidence"] == "No"
    assert rows_by_ref["REQ-100"]["Review Complete"] == "Yes"

    assert rows_by_ref["REQ-110"]["Requires Evidence"] == "Yes"
    assert rows_by_ref["REQ-110"]["Unassigned"] == "Yes"
    assert rows_by_ref["REQ-110"]["Responsible for Requirement Name"] == ""
    assert rows_by_ref["REQ-110"]["Assigned Reviewer Name"] == "Test User"
    assert rows_by_ref["REQ-110"]["Review Complete"] == "No"

    assert rows_by_ref["REQ-120"]["Requires Evidence"] == "Yes"
    assert rows_by_ref["REQ-120"]["Unassigned"] == "No"
    assert rows_by_ref["REQ-120"]["Responsible for Requirement Name"] == "Test User"
    assert rows_by_ref["REQ-120"]["Review Status"] == "escalated"
    assert rows_by_ref["REQ-120"]["Review Complete"] == "No"

    assert rows_by_ref["REQ-130"]["Requires Evidence"] == "Yes"
    assert rows_by_ref["REQ-130"]["Unassigned"] == "No"
    assert rows_by_ref["REQ-130"]["Responsible for Requirement Name"] == "Contributor User"
    assert rows_by_ref["REQ-130"]["Assigned Reviewer Name"] == "Assigned Reviewer"

    assert rows_by_ref["REQ-140"]["Requires Evidence"] == "Yes"
    assert rows_by_ref["REQ-140"]["Review Complete"] == "No"
    assert rows_by_ref["REQ-140"]["Review Evidence"] == "Draft evidence note"

    assert rows_by_ref["REQ-110"]["Overdue"] == "Yes"
    assert "REQ-150" not in rows_by_ref


async def test_master_gap_analysis_rules(
    client, auth_headers, default_jurisdiction, master_gap_data
):
    response = await client.get(
        f"/api/v1/reports/gap-analysis?jurisdiction_id={default_jurisdiction.id}&format=csv",
        headers=auth_headers,
    )
    assert response.status_code == 200
    rows = _read_csv_rows(response.text)
    rows_by_ref = {row["Reference ID"]: row for row in rows}
    assert set(rows_by_ref) == {
        master_gap_data["done"],
        master_gap_data["unassigned"],
        master_gap_data["blocked"],
    }
    assert rows_by_ref[master_gap_data["done"]]["Requires Evidence"] == "No"
    assert rows_by_ref[master_gap_data["done"]]["Unassigned"] == "No"
    assert rows_by_ref[master_gap_data["unassigned"]]["Requires Evidence"] == "Yes"
    assert rows_by_ref[master_gap_data["unassigned"]]["Unassigned"] == "Yes"
    assert rows_by_ref[master_gap_data["blocked"]]["Requires Evidence"] == "Yes"
    assert rows_by_ref[master_gap_data["blocked"]]["Unassigned"] == "No"
    assert master_gap_data["informational"] not in rows_by_ref
    assert all(not row["Review Status"] for row in rows)
    assert all(not row["Review Complete"] for row in rows)


async def test_cycle_gap_analysis_assigned_reviewer_visibility(
    client,
    assigned_reviewer_headers,
    action_gap_cycle_data,
):
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        f"/api/v1/reports/gap-analysis?review_cycle_id={cycle.id}&format=csv",
        headers=assigned_reviewer_headers,
    )
    assert response.status_code == 200
    rows = _read_csv_rows(response.text)
    assert [row["Reference ID"] for row in rows] == [
        action_gap_cycle_data["assigned_reviewer_reference"]
    ]


async def test_cycle_gap_analysis_allowed_for_contributor(
    client,
    contributor_user_headers,
    action_gap_cycle_data,
):
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        f"/api/v1/reports/gap-analysis?review_cycle_id={cycle.id}&format=csv",
        headers=contributor_user_headers,
    )
    assert response.status_code == 200
    assert "REQ-140" in response.text


async def test_master_gap_analysis_forbidden_for_contributor(
    client,
    contributor_headers,
    default_jurisdiction,
):
    response = await client.get(
        f"/api/v1/reports/gap-analysis?jurisdiction_id={default_jurisdiction.id}&format=csv",
        headers=contributor_headers,
    )
    assert response.status_code == 403


async def test_gap_analysis_jurisdiction_xlsx(
    client,
    auth_headers,
    default_jurisdiction,
    master_gap_data,
):
    if load_workbook is None:
        pytest.skip("openpyxl is not installed")
    response = await client.get(
        f"/api/v1/reports/gap-analysis?jurisdiction_id={default_jurisdiction.id}&format=xlsx",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "Needs Evidence", "Open Gaps", "All Items"]
    summary_sheet = workbook["Summary"]
    requirement_summary = _summary_table_rows(summary_sheet, "Requirement Status")
    assert requirement_summary == {
        "Not Started": "1",
        "Blocked": "1",
        "Evidenced": "1",
    }
    assert _summary_table_rows(summary_sheet, "Review Status") == {}
    all_rows = _worksheet_rows(workbook["All Items"])
    rows_by_ref = {row["Reference ID"]: row for row in all_rows}
    assert rows_by_ref[master_gap_data["blocked"]]["Requires Evidence"] == "Yes"
    needs_rows = _worksheet_rows(workbook["Needs Evidence"])
    assert {row["Reference ID"] for row in needs_rows} == {
        master_gap_data["unassigned"],
        master_gap_data["blocked"],
    }


async def test_gap_analysis_cycle_xlsx(client, auth_headers, action_gap_cycle_data):
    if load_workbook is None:
        pytest.skip("openpyxl is not installed")
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        f"/api/v1/reports/gap-analysis?review_cycle_id={cycle.id}&format=xlsx",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "Needs Evidence", "Open Gaps", "All Items", "Owner Actions"]
    summary_sheet = workbook["Summary"]
    assert _summary_table_rows(summary_sheet, "Requirement Status") == {
        "Not Started": "1",
        "In Progress": "2",
        "Blocked": "1",
        "Evidenced": "1",
    }
    assert _summary_table_rows(summary_sheet, "Review Status") == {
        "Pending": "3",
        "Confirmed": "1",
        "Escalated": "1",
    }
    needs_rows = _worksheet_rows(workbook["Needs Evidence"])
    assert {row["Reference ID"] for row in needs_rows} == {
        "REQ-110",
        "REQ-120",
        "REQ-130",
        "REQ-140",
    }
    assert all(row["Requires Evidence"] == "Yes" for row in needs_rows)


    open_rows = _worksheet_rows(workbook["Open Gaps"])
    assert all(row["Needs Attention"] == "Yes" and row["Gap Reasons / Next Actions"] for row in open_rows)
    assert workbook["Open Gaps"].freeze_panes == "D2"
    assert workbook["Open Gaps"].auto_filter.ref


async def test_gap_analysis_jurisdiction_pdf(
    client,
    auth_headers,
    default_jurisdiction,
    master_gap_data,
):
    response = await client.get(
        f"/api/v1/reports/gap-analysis?jurisdiction_id={default_jurisdiction.id}&format=pdf",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Gap Analysis" in text
    assert "Requirement Status Summary" in text
    assert "Outstanding Items By Responsible Person" in text
    assert "Open Gaps: Evidence and Review Decisions" in text
    assert "Review Status Summary" not in text
    assert "Action Items" not in text
    assert master_gap_data["informational"] not in text
    assert master_gap_data["blocked"] in text
    owner_summary = text.split("Outstanding Items By Responsible Person", 1)[1].split(
        "Needs Evidence", 1
    )[0]
    assert "Responsible Person" in owner_summary
    assert "Total Outstanding" in owner_summary
    assert "Not Started" in owner_summary
    assert "Blocked" in owner_summary
    assert "Unassigned" in owner_summary
    assert "Contributor User" in owner_summary
    assert re.search(r"Unassigned\s+1\s+1\s+0", owner_summary)


async def test_gap_analysis_cycle_pdf(client, auth_headers, action_gap_cycle_data):
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        f"/api/v1/reports/gap-analysis?review_cycle_id={cycle.id}&format=pdf",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Review gap analysis" in text
    assert "Outstanding work by responsible person" in text
    assert "Outstanding actions" in text
    assert "Live assessment at export" in text
    assert "Full review" in text
    assert "REQ-150" not in text
    assert "REQ-130" in text
    assert "Action: Evidence is not complete" in text
    assert "Contributor User" in text
    assert "Unassigned" in text
    assert "Page 1" in text


async def test_compliance_summary_excludes_non_actionable_requirements(
    client,
    auth_headers,
    default_jurisdiction,
    master_gap_data,
):
    response = await client.get(
        f"/api/v1/reports/compliance-summary?jurisdiction_id={default_jurisdiction.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Overall Status" in text
    assert "Blocked" in text
    assert "Evidenced" in text
    assert "Not Applicable" not in text


async def test_detailed_report_excludes_non_actionable_requirements(
    client,
    auth_headers,
    default_jurisdiction,
    master_gap_data,
):
    response = await client.get(
        f"/api/v1/reports/detailed?jurisdiction_id={default_jurisdiction.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert master_gap_data["blocked"] in text
    assert master_gap_data["informational"] not in text


async def test_statement_of_applicability_retains_informational_requirements(
    client,
    auth_headers,
    action_gap_cycle_data,
):
    if load_workbook is None:
        pytest.skip("openpyxl is not installed")
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        "/api/v1/reports/statement-of-applicability"
        f"?review_cycle_id={cycle.id}"
        "&format=xlsx"
        "&fields=requirement_reference_id",
        headers=auth_headers,
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    refs = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert set(refs) == {"REQ-100", "REQ-110", "REQ-120", "REQ-130", "REQ-140", "REQ-150"}
    assert "REQ-150" in refs


async def test_statement_of_applicability_xlsx(
    client,
    auth_headers,
    review_cycle,
    review_item,
    review_item_evidence,
    requirement_status,
):
    if load_workbook is None:
        pytest.skip("openpyxl is not installed")
    response = await client.get(
        "/api/v1/reports/statement-of-applicability"
        f"?review_cycle_id={review_cycle.id}"
        "&format=xlsx"
        "&fields=review_cycle_id,requirement_reference_id,review_item_review_status",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers == [
        "Review Cycle ID",
        "Requirement Reference ID",
        "Review Item Review Status",
    ]


async def test_statement_of_applicability_pdf(
    client,
    auth_headers,
    review_cycle,
    review_item,
    review_item_evidence,
    requirement_status,
):
    response = await client.get(
        f"/api/v1/reports/statement-of-applicability?review_cycle_id={review_cycle.id}&format=pdf",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Statement of Applicability" in text
    assert "Review Cycle Information" in text


async def test_statement_of_applicability_fields_empty(client, auth_headers, review_cycle):
    response = await client.get(
        f"/api/v1/reports/statement-of-applicability?review_cycle_id={review_cycle.id}&fields=",
        headers=auth_headers,
    )
    assert response.status_code == 400


async def test_statement_of_applicability_fields_invalid(client, auth_headers, review_cycle):
    response = await client.get(
        f"/api/v1/reports/statement-of-applicability?review_cycle_id={review_cycle.id}&fields=bad_field",
        headers=auth_headers,
    )
    assert response.status_code == 400


async def test_review_cycle_report_pdf(client, auth_headers, review_cycle, review_item):
    response = await client.get(
        f"/api/v1/reports/review-cycle-report?review_cycle_id={review_cycle.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Readiness Review Report" in text
    assert "Review Cycle Summary" in text


async def test_review_cycle_report_retains_informational_requirements(
    client,
    auth_headers,
    action_gap_cycle_data,
):
    cycle = action_gap_cycle_data["cycle"]
    response = await client.get(
        f"/api/v1/reports/review-cycle-report?review_cycle_id={cycle.id}",
        headers=auth_headers,
    )
    assert response.status_code == 200
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "REQ-140" in text
    assert "REQ-150" in text


async def test_download_audit_trail(client, auth_headers, audit_logs):
    response = await client.get("/api/v1/reports/audit-trail", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"


async def test_audit_trail_with_date_filter(client, auth_headers, audit_logs):
    response = await client.get(
        "/api/v1/reports/audit-trail?from_date=2020-01-01&to_date=2030-12-31",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"


async def test_reports_unauthorized(client):
    response = await client.get("/api/v1/reports/compliance-summary")
    assert response.status_code == 401


async def test_reports_forbidden_for_contributor(client, contributor_headers, test_requirement):
    response = await client.get("/api/v1/reports/compliance-summary", headers=contributor_headers)
    assert response.status_code == 403


async def test_master_reports_require_jurisdiction(client, auth_headers):
    response = await client.get("/api/v1/reports/compliance-summary", headers=auth_headers)
    assert response.status_code == 400


async def test_global_exports_use_current_baseline_and_preserve_review_scope(
    db_session,
    test_user,
    test_document,
    test_requirement,
    review_item,
    review_cycle,
    default_jurisdiction,
):
    from app.models.program import RequirementSetVersion
    from app.services.reports import (
        _fetch_requirements_with_status,
        _build_master_gap_analysis_rows,
    )

    old = RequirementSetVersion(
        document_id=test_document.id,
        version_number=1,
        status="approved",
        is_current=False,
        created_by=test_user.id,
    )
    current = RequirementSetVersion(
        document_id=test_document.id,
        version_number=2,
        status="approved",
        is_current=True,
        created_by=test_user.id,
    )
    db_session.add_all([old, current])
    await db_session.flush()
    test_requirement.requirement_set_version_id = old.id
    current_requirement = Requirement(
        jurisdiction_id=default_jurisdiction.id,
        document_id=test_document.id,
        requirement_set_version_id=current.id,
        reference_id="CURRENT-1",
        text="Current control",
        requirement_type="mandatory",
        active=True,
    )
    db_session.add(current_requirement)
    await db_session.commit()
    global_rows = await _fetch_requirements_with_status(
        db_session, jurisdiction_id=default_jurisdiction.id
    )
    assert [row[0].id for row in global_rows] == [current_requirement.id]
    review_rows = await _fetch_requirements_with_status(
        db_session, review_cycle_id=review_cycle.id, jurisdiction_id=default_jurisdiction.id
    )
    assert [row[0].id for row in review_rows] == [test_requirement.id]
    gap_rows = await _build_master_gap_analysis_rows(db_session, default_jurisdiction.id, test_user)
    assert len(gap_rows) == 1
    assert gap_rows[0].reference_id == "CURRENT-1"


async def test_audit_calendar_dates_include_full_day_but_timestamps_keep_precision(client, auth_headers, db_session, test_user):
    for day, hour, marker in [(20, 0, "at-midnight"), (20, 16, "late-afternoon"), (21, 0, "next-day")]:
        db_session.add(AuditLog(user_id=test_user.id, user_name=test_user.full_name, action="update", entity_type="test", entity_id=marker, timestamp=datetime(2026, 9, day, hour, tzinfo=timezone.utc)))
    await db_session.commit()
    response = await client.get("/api/v1/reports/audit-trail", params={"from_date": "2026-09-20", "to_date": "2026-09-20"}, headers=auth_headers)
    assert response.status_code == 200
    assert "late-afternoon" in response.text and "at-midnight" in response.text
    assert "next-day" not in response.text
    timestamp = await client.get("/api/v1/reports/audit-trail", params={"from_date": "2026-09-20", "to_date": "2026-09-20T00:00:00Z"}, headers=auth_headers)
    assert timestamp.status_code == 200
    assert "at-midnight" in timestamp.text and "late-afternoon" not in timestamp.text
    for params in [{"from_date": "2026-09-21", "to_date": "2026-09-20"}, {"to_date": "invalid"}, {"to_date": "2026-02-30"}]:
        assert (await client.get("/api/v1/reports/audit-trail", params=params, headers=auth_headers)).status_code == 422
