from app.services.reports import get_statement_of_applicability_fields
import io
import json
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from app.services.document_reader import default_reader
from reportlab.pdfgen.canvas import Canvas

from app.config import settings
from app.models import (
    Document,
    Requirement,
    RequirementStatus,
    ReviewCycle,
    ReviewItem,
    Snapshot,
    User,
)
from app.services.auth import create_access_token, hash_password
from app.services.review_assurance import assessment_rows
from app.services.reports import (
    generate_review_cycle_report_pdf,
    generate_statement_of_applicability_xlsx,
)
import csv


@pytest.fixture
async def assessment(db_session, default_jurisdiction, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    user = User(
        email="assessor@example.test",
        full_name="Compliance Lead",
        role="admin",
        password_hash=hash_password("fixture-password"),
    )
    db_session.add(user)
    await db_session.flush()
    source = tmp_path / "source.pdf"
    pdf = Canvas(str(source))
    pdf.drawString(40, 700, "Official source fixture")
    pdf.save()
    document = Document(
        jurisdiction_id=default_jurisdiction.id,
        filename="source.pdf",
        name="Source <baseline> & requirements",
        document_type="standard",
        status="approved",
        file_path=str(source),
        uploaded_by=user.id,
    )
    db_session.add(document)
    await db_session.flush()
    cycle = ReviewCycle(
        name="Assessment <operator> & suppliers",
        jurisdiction_id=default_jurisdiction.id,
        created_by=user.id,
    )
    db_session.add(cycle)
    await db_session.flush()
    items = []
    for index, kind in enumerate(["mandatory", "mandatory", "informational"]):
        requirement = Requirement(
            jurisdiction_id=default_jurisdiction.id,
            document_id=document.id,
            reference_id=f"3.{index + 1}",
            title='=HYPERLINK("https://example.test")',
            text="The operator shall retain evidence.",
            requirement_type=kind,
            sort_order=index,
        )
        db_session.add(requirement)
        await db_session.flush()
        item = ReviewItem(
            review_cycle_id=cycle.id,
            requirement_id=requirement.id,
            review_status="pending",
            assessment_status="not_started",
        )
        db_session.add(item)
        items.append(item)
    await db_session.commit()
    headers = {"Authorization": "Bearer " + create_access_token({"sub": str(user.id)})}
    return user, cycle, items, headers, source


async def finish(client, assessment):
    user, cycle, items, headers, _ = assessment
    for item, state, evidence in [
        (items[0], "evidenced", "Control tested on 2026-09-20. Evidence reference: E-001"),
        (
            items[1],
            "not_applicable",
            "Operator does not offer this product; scope approved by compliance.",
        ),
    ]:
        response = await client.put(
            f"/api/v1/review-cycles/{cycle.id}/items/{item.id}",
            headers=headers,
            json={
                "review_status": "confirmed",
                "requirement_status": state,
                "review_evidence": evidence,
            },
        )
        assert response.status_code == 200, response.text
    return await client.post(f"/api/v1/review-cycles/{cycle.id}/close", headers=headers)


async def test_missing_cycle_assessment_does_not_inherit_global_evidence(
    client, db_session, assessment,
):
    user, cycle, items, headers, _ = assessment
    first = items[0]
    first.assessment_status = None
    first.review_status = "confirmed"
    first.review_evidence = "Cycle-local note"
    db_session.add(RequirementStatus(
        requirement_id=first.requirement_id, status="evidenced",
        comment="Other workflow evidence", changed_by=user.id,
    ))
    await db_session.commit()
    rows = await assessment_rows(db_session, cycle)
    assert rows[0].status == "not_started"
    assert (await client.post(f"/api/v1/review-cycles/{cycle.id}/close", headers=headers)).status_code == 409


async def test_closure_gates_and_frozen_reports_survive_changes(client, db_session, assessment):
    user, cycle, items, headers, source = assessment
    base = f"/api/v1/review-cycles/{cycle.id}"
    rejected = await client.post(base + "/close", headers=headers)
    assert rejected.status_code == 409
    assert rejected.json()["detail"]["blocker_count"] == 2
    assert (await finish(client, assessment)).status_code == 200
    await db_session.refresh(cycle)
    old = await assessment_rows(db_session, cycle)
    assert len(old) == 3
    # Change the global master status and source after closure: the review remains frozen.
    req = await db_session.get(Requirement, items[0].requirement_id)
    req.text = "CHANGED SOURCE TEXT"
    db_session.add(
        RequirementStatus(
            requirement_id=req.id,
            status="blocked",
            comment="Changed after closure",
            changed_by=user.id,
        )
    )
    source.write_bytes(b"changed original source")
    await db_session.commit()
    response = await client.get(base, headers=headers)
    assert response.status_code == 200, response.text
    row = response.json()["items"][0]
    assert row["requirement_current_status"] == "evidenced"
    assert row["requirement"]["text"] == "The operator shall retain evidence."
    assert response.json()["readiness"]["can_close"] is True
    exported = await client.get(
        "/api/v1/reports/review-package", headers=headers, params={"review_cycle_id": str(cycle.id)}
    )
    assert exported.status_code == 200, exported.text[:200] if exported.status_code != 200 else ""
    archive = zipfile.ZipFile(io.BytesIO(exported.content))
    manifest = json.loads(archive.read("manifest.json"))
    assert any(entry["kind"] == "sources" for entry in manifest["files"])
    import hashlib

    for entry in manifest["files"]:
        assert hashlib.sha256(archive.read(entry["path"])).hexdigest() == entry["sha256"]
    workbook = load_workbook(io.BytesIO(archive.read("statement-of-applicability.xlsx")))
    assert workbook.active.max_row == 4  # Includes informational and not-applicable controls.
    assert all(cell.data_type != "f" for row in workbook.active for cell in row)
    text = "\n".join(
        p["text"] for p in default_reader.read_pages(archive.read("readiness-review.pdf"))
    )
    assert "not apply" not in text or "applicable" in text
    assert "not applicable" in text.lower()
    assert "CHANGED SOURCE TEXT" not in text


async def test_source_free_review_closes_without_source_snapshot(client, db_session, assessment):
    _, cycle, items, headers, _ = assessment
    requirement = await db_session.get(Requirement, items[0].requirement_id)
    document = await db_session.get(Document, requirement.document_id)
    document.filename = None
    document.file_path = None
    await db_session.commit()

    closed = await finish(client, assessment)
    assert closed.status_code == 200, closed.text
    snapshot = await db_session.get(Snapshot, uuid.UUID(closed.json()["snapshot_id"]))
    manifest = json.loads(snapshot.data_json)["files"]
    assert not any(entry["kind"] == "sources" for entry in manifest)


@pytest.mark.parametrize("status", ["evidenced", "not_applicable"])
async def test_empty_evidence_is_allowed_for_evidenced_but_not_applicability(client, assessment, status):
    _, cycle, items, headers, _ = assessment
    result = await client.put(
        f"/api/v1/review-cycles/{cycle.id}/items/{items[0].id}",
        headers=headers,
        json={
            "review_status": "confirmed",
            "requirement_status": status,
            "review_evidence": "<p><br></p>",
        },
    )
    assert result.status_code == (200 if status == "evidenced" else 422)


async def test_review_status_is_local_and_invalid_states_rejected(client, db_session, assessment):
    user, cycle, items, headers, _ = assessment
    other = ReviewCycle(
        name="Separate assessment", jurisdiction_id=cycle.jurisdiction_id, created_by=user.id
    )
    db_session.add(other)
    await db_session.flush()
    db_session.add(
        ReviewItem(
            review_cycle_id=other.id,
            requirement_id=items[0].requirement_id,
            assessment_status="blocked",
        )
    )
    await db_session.commit()
    assert (await finish(client, assessment)).status_code == 200
    assert (await assessment_rows(db_session, other))[0].status == "blocked"
    result = await client.patch(
        f"/api/v1/review-cycles/{other.id}/items/bulk",
        headers=headers,
        json={"item_ids": [str(items[0].id)], "review_status": "banana"},
    )
    assert result.status_code == 422


async def test_export_detects_tampered_frozen_evidence(client, db_session, assessment):
    user, cycle, items, headers, source = assessment
    assert (await finish(client, assessment)).status_code == 200
    await db_session.refresh(cycle)
    snapshot = await db_session.get(Snapshot, cycle.snapshot_id)
    payload = json.loads(snapshot.data_json)
    frozen = (
        Path(settings.upload_dir)
        / "review-snapshots"
        / str(snapshot.id)
        / payload["files"][0]["path"]
    )
    frozen.write_bytes(b"tampered")
    response = await client.get(
        "/api/v1/reports/review-package", headers=headers, params={"review_cycle_id": str(cycle.id)}
    )
    assert response.status_code == 409
    assert "integrity" in response.text


async def test_long_evidence_can_span_pages(db_session, assessment):
    _, cycle, items, _, _ = assessment
    items[0].review_evidence = "Long evidence record & verification details. " * 800
    await db_session.commit()
    content = await generate_review_cycle_report_pdf(db_session, cycle)
    assert len(default_reader.read_pages(content)) > 2
    workbook = load_workbook(
        io.BytesIO(
            await generate_statement_of_applicability_xlsx(
                db_session, cycle, get_statement_of_applicability_fields(None)
            )
        )
    )
    assert workbook.active.max_row == 4
    assert all(c.data_type != "f" for row in workbook.active for c in row)


async def test_source_pdf_is_authenticated_and_scoped(client, db_session, assessment):
    user, cycle, items, headers, source = assessment
    req = await db_session.get(Requirement, items[0].requirement_id)
    route = f"/api/v1/documents/{req.document_id}/source"
    assert (await client.get(route)).status_code == 401
    response = await client.get(route, headers=headers)
    assert response.status_code == 200
    assert response.content == source.read_bytes()
    assert response.headers["content-disposition"].startswith("inline")
    other = User(email="different-company@example.test", full_name="Other company", role="admin", organization_id=__import__("uuid").uuid4(), password_hash="not-a-login")
    db_session.add(other)
    await db_session.commit()
    assert (await client.get(route, headers={"Authorization": "Bearer " + create_access_token({"sub": str(other.id)})})).status_code == 404


async def test_closed_evidence_download_is_frozen_and_completed_review_is_retained(client, db_session, assessment):
    _, cycle, items, headers, _ = assessment
    base = f"/api/v1/review-cycles/{cycle.id}"
    uploaded = await client.post(base + f"/items/{items[0].id}/files", headers=headers,
                                 files={"file": ("test-evidence.txt", b"verified original evidence", "text/plain")})
    assert uploaded.status_code == 201, uploaded.text
    file_id = uploaded.json()["id"]
    assert (await finish(client, assessment)).status_code == 200
    from app.models.review import ReviewItemEvidenceFile
    import uuid
    evidence = await db_session.get(ReviewItemEvidenceFile, uuid.UUID(file_id))
    Path(evidence.file_path).write_bytes(b"changed original evidence")
    download = base + f"/items/{items[0].id}/files/{file_id}/download"
    response = await client.get(download, headers=headers)
    assert response.status_code == 200
    assert response.content == b"verified original evidence"
    rejected = await client.delete(base, headers=headers)
    assert rejected.status_code == 409
    assert "Archive" in rejected.text


async def test_confirmed_evidenced_item_without_recorded_evidence_is_not_a_completion_blocker(client, db_session, assessment):
    user, cycle, items, headers, _ = assessment
    db_session.add(RequirementStatus(requirement_id=items[0].requirement_id, status="evidenced",
                                   comment="Old master evidence from a different review", changed_by=user.id))
    items[0].assessment_status = "evidenced"
    items[0].review_status = "confirmed"
    await db_session.commit()
    rows = await assessment_rows(db_session, cycle)
    assert rows[0].status_comment == ""
    response = await client.post(f"/api/v1/review-cycles/{cycle.id}/close", headers=headers)
    assert response.status_code == 409
    assert all(b["item_id"] != str(items[0].id) for b in response.json()["detail"]["blockers"])


async def test_evidenced_confirmed_without_recorded_evidence_can_close(client, assessment):
    _, cycle, items, headers, _ = assessment
    base = f"/api/v1/review-cycles/{cycle.id}"
    first = await client.put(base + f"/items/{items[0].id}", headers=headers, json={
        "requirement_status": "evidenced", "review_status": "confirmed",
    })
    assert first.status_code == 200, first.text
    second = await client.put(base + f"/items/{items[1].id}", headers=headers, json={
        "requirement_status": "not_applicable", "review_status": "confirmed",
        "review_evidence": "Outside scope",
    })
    assert second.status_code == 200, second.text
    view = await client.get(base, headers=headers)
    assert view.status_code == 200
    assert view.json()["readiness"]["blockers"] == []
    assert view.json()["items"][0]["reviewer_name"] == "Compliance Lead"
    exported = await client.get(
        "/api/v1/reports/gap-analysis", headers=headers,
        params={"review_cycle_id": str(cycle.id), "format": "csv"},
    )
    assert exported.status_code == 200
    first_row = next(row for row in csv.DictReader(io.StringIO(exported.text)) if row["Reference ID"] == "3.1")
    assert first_row["No Evidence Recorded"] == "Yes"
    assert first_row["Needs Attention"] == "No"
    assert (await client.post(base + "/close", headers=headers)).status_code == 200


async def test_evidence_edit_after_confirmation_is_disclosed_without_reopening(client, assessment):
    _, cycle, items, headers, _ = assessment
    base = f"/api/v1/review-cycles/{cycle.id}"
    assert (await client.put(base + f"/items/{items[0].id}", headers=headers, json={
        "requirement_status": "evidenced", "review_status": "confirmed",
        "review_evidence": "Initial reference",
    })).status_code == 200
    assert (await client.put(base + f"/items/{items[0].id}", headers=headers, json={
        "review_evidence": "Revised reference",
    })).status_code == 200
    view = await client.get(base, headers=headers)
    row = next(item for item in view.json()["items"] if item["id"] == str(items[0].id))
    assert row["evidence_changed_since_review"] is True
    assert all(b["item_id"] != str(items[0].id) for b in view.json()["readiness"]["blockers"])
    assert (await client.put(base + f"/items/{items[0].id}", headers=headers, json={
        "review_status": "confirmed",
    })).status_code == 200
    refreshed = await client.get(base, headers=headers)
    reviewed_row = next(item for item in refreshed.json()["items"] if item["id"] == str(items[0].id))
    assert reviewed_row["evidence_changed_since_review"] is False


async def test_source_corrections_and_feedback_cannot_cross_company_boundary(client, db_session, assessment):
    from app.models.requirement import ExtractedRequirement
    from app.models.extraction import ExtractionRun
    import uuid
    user, cycle, items, headers, _ = assessment
    req = await db_session.get(Requirement, items[0].requirement_id)
    run = ExtractionRun(document_id=req.document_id, status="completed", total_pages=1, current_page=1, ai_provider="local", ai_model="test")
    db_session.add(run)
    await db_session.flush()
    extracted = ExtractedRequirement(document_id=req.document_id, extraction_run_id=run.id, reference_id="3.1", text="Original source", original_text="Original source", page_number=1)
    other = User(email="source-editor-other@example.test", full_name="Other company", role="admin", organization_id=uuid.uuid4(), password_hash="not-a-login")
    db_session.add_all([extracted, other])
    await db_session.commit()
    other_headers = {"Authorization": "Bearer " + create_access_token({"sub": str(other.id)})}
    route = f"/api/v1/documents/{req.document_id}/extractions/{extracted.id}"
    assert (await client.put(route, headers=other_headers, json={"text":"Cross-company overwrite"})).status_code == 404
    assert (await client.post(route + "/feedback", headers=other_headers, json={"action":"accept"})).status_code == 404
    await db_session.refresh(extracted)
    assert extracted.text == "Original source"


async def test_archiving_source_preserves_provenance_and_completed_review(client, db_session, assessment):
    _, cycle, items, headers, source = assessment
    assert (await finish(client, assessment)).status_code == 200
    req = await db_session.get(Requirement, items[0].requirement_id)
    rejected = await client.delete(f"/api/v1/requirements/sets/{req.document_id}", headers=headers)
    assert rejected.status_code == 409
    assert (await client.delete(f"/api/v1/documents/{req.document_id}", headers=headers)).status_code == 204
    assert source.is_file()
    assert (await client.get("/api/v1/reports/review-package", headers=headers, params={"review_cycle_id":str(cycle.id)})).status_code == 200


async def test_evidence_only_save_preserves_existing_review_decision(client, assessment):
    user, cycle, items, headers, _ = assessment
    target = f'/api/v1/review-cycles/{cycle.id}/items/{items[0].id}'
    first = await client.put(target, headers=headers, json={'review_evidence': 'Initial source evidence'})
    assert first.status_code == 200
    assert first.json()['review_status'] == 'pending'
    assert first.json()['reviewer_id'] is None
    decision = await client.put(target, headers=headers, json={'review_status': 'confirmed'})
    assert decision.status_code == 200
    updated = await client.put(target, headers=headers, json={'review_evidence': 'Additional source reference'})
    assert updated.status_code == 200
    assert updated.json()['review_status'] == 'confirmed'
    assert updated.json()['reviewer_id'] == decision.json()['reviewer_id']
    assert updated.json()['reviewed_at'] == decision.json()['reviewed_at']
    assert updated.json()['review_evidence'] == 'Additional source reference'
    invalid = await client.put(target, headers=headers, json={'review_status': 'typo'})
    assert invalid.status_code == 422


@pytest.mark.parametrize("status", ["evidenced", "not_applicable"])
async def test_evidence_removal_preserves_current_assessment_and_applicability_rule(
    client, db_session, assessment, status
):
    _, cycle, items, headers, _ = assessment
    target = f"/api/v1/review-cycles/{cycle.id}/items/{items[0].id}"
    original = await client.put(
        target,
        headers=headers,
        json={"review_status": "confirmed", "requirement_status": status,
              "review_evidence": "Original assessment rationale"},
    )
    assert original.status_code == 200, original.text
    removal = await client.put(target, headers=headers, json={"review_evidence": "<p><br></p>"})
    assert removal.status_code == (200 if status == "evidenced" else 422)
    await db_session.refresh(items[0])
    assert items[0].review_evidence == ("<p><br></p>" if status == "evidenced" else "Original assessment rationale")

    replacement = await client.put(target, headers=headers, json={"review_evidence": "Current assessment rationale"})
    assert replacement.status_code == 200, replacement.text
    assert replacement.json()["reviewer_id"] == original.json()["reviewer_id"]
    assert replacement.json()["reviewed_at"] == original.json()["reviewed_at"]
    await db_session.refresh(items[0])
    assert items[0].assessment_rationale == "Current assessment rationale"
    rows = await assessment_rows(db_session, cycle)
    assert rows[0].status_comment == "Current assessment rationale"
    current_export = load_workbook(io.BytesIO(await generate_statement_of_applicability_xlsx(
        db_session, cycle, get_statement_of_applicability_fields(None)
    )))
    current_cells = [str(cell.value) for row in current_export.active for cell in row if cell.value]
    assert "Current assessment rationale" in current_cells
    assert "Original assessment rationale" not in current_cells

    other = await client.put(
        f"/api/v1/review-cycles/{cycle.id}/items/{items[1].id}",
        headers=headers,
        json={"review_status": "confirmed", "requirement_status": "not_applicable",
              "review_evidence": "Outside review scope"},
    )
    assert other.status_code == 200, other.text
    closed = await client.post(f"/api/v1/review-cycles/{cycle.id}/close", headers=headers)
    assert closed.status_code == 200, closed.text
    await db_session.refresh(cycle)
    frozen = await assessment_rows(db_session, cycle)
    assert frozen[0].status_comment == "Current assessment rationale"
    frozen_export = load_workbook(io.BytesIO(await generate_statement_of_applicability_xlsx(
        db_session, cycle, get_statement_of_applicability_fields(None)
    )))
    frozen_cells = [str(cell.value) for row in frozen_export.active for cell in row if cell.value]
    assert "Current assessment rationale" in frozen_cells
    assert "Original assessment rationale" not in frozen_cells
    assert "Original assessment rationale" not in json.dumps(json.loads(
        (await db_session.get(Snapshot, cycle.snapshot_id)).data_json
    ))


async def test_stale_rationale_is_not_reused_after_text_removal(client, db_session, assessment):
    _, cycle, items, headers, _ = assessment
    first = items[0]
    first.assessment_status = "evidenced"
    first.assessment_rationale = "Superseded evidence text"
    first.review_evidence = None
    first.review_status = "confirmed"
    first.reviewer_id = assessment[0].id
    first.reviewed_at = datetime.now(timezone.utc)
    await db_session.commit()
    other = await client.put(
        f"/api/v1/review-cycles/{cycle.id}/items/{items[1].id}",
        headers=headers,
        json={"review_status": "confirmed", "requirement_status": "not_applicable",
              "review_evidence": "Outside review scope"},
    )
    assert other.status_code == 200, other.text
    rows = await assessment_rows(db_session, cycle)
    assert rows[0].status_comment == ""
    closed = await client.post(f"/api/v1/review-cycles/{cycle.id}/close", headers=headers)
    assert closed.status_code == 200, closed.text
    frozen = await assessment_rows(db_session, cycle)
    assert frozen[0].status_comment == ""


@pytest.mark.parametrize("status", ["evidenced", "not_applicable"])
async def test_legacy_evidence_edit_captures_status_without_global_comment_fallback(
    client, db_session, assessment, status
):
    user, cycle, items, headers, _ = assessment
    first = items[0]
    first.assessment_status = None
    first.review_evidence = "Initial local text"
    db_session.add(RequirementStatus(
        requirement_id=first.requirement_id,
        status=status,
        comment="Superseded global rationale",
        changed_by=user.id,
    ))
    await db_session.commit()
    target = f"/api/v1/review-cycles/{cycle.id}/items/{first.id}"
    removal = await client.put(target, headers=headers, json={"review_evidence": None})
    assert removal.status_code == (200 if status == "evidenced" else 422)
    await db_session.refresh(first)
    assert first.assessment_status == ("evidenced" if status == "evidenced" else None)
    assert first.review_evidence == (None if status == "evidenced" else "Initial local text")

    replacement = await client.put(target, headers=headers, json={"review_evidence": "Current local text"})
    assert replacement.status_code == 200, replacement.text
    await db_session.refresh(first)
    assert first.assessment_status == status
    assert first.assessment_rationale == "Current local text"
    rows = await assessment_rows(db_session, cycle)
    assert rows[0].status_comment == "Current local text"
    assert (await client.put(target, headers=headers, json={"review_evidence": "<p><br></p>"})).status_code == (200 if status == "evidenced" else 422)


async def test_bulk_evidence_edit_uses_same_assessment_invariant(client, db_session, assessment):
    _, cycle, items, headers, _ = assessment
    target = f"/api/v1/review-cycles/{cycle.id}/items/{items[0].id}"
    original = await client.put(
        target, headers=headers,
        json={"review_status": "confirmed", "requirement_status": "evidenced",
              "review_evidence": "Original control proof"},
    )
    assert original.status_code == 200, original.text
    bulk = f"/api/v1/review-cycles/{cycle.id}/items/bulk"
    cleared = await client.patch(
        bulk, headers=headers,
        json={"item_ids": [str(items[0].id)], "review_evidence": ""},
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["updated_count"] == 1
    assert cleared.json()["failed_count"] == 0
    accepted = await client.patch(
        bulk, headers=headers,
        json={"item_ids": [str(items[0].id)], "review_evidence": "Replacement control proof"},
    )
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["updated_count"] == 1
    await db_session.refresh(items[0])
    assert items[0].assessment_rationale == "Replacement control proof"
    assert items[0].review_status == "confirmed"
    assert str(items[0].reviewer_id) == original.json()["reviewer_id"]
    assert items[0].reviewed_at.isoformat() == datetime.fromisoformat(original.json()["reviewed_at"]).isoformat()


async def test_last_evidence_file_deletion_is_recorded_without_reopening(client, db_session, assessment):
    _, cycle, items, headers, _ = assessment
    target = f"/api/v1/review-cycles/{cycle.id}/items/{items[0].id}"
    files = target + "/files"
    first = await client.post(files, headers=headers,
                              files={"file": ("first.txt", b"first proof", "text/plain")})
    assert first.status_code == 201, first.text
    marked = await client.put(target, headers=headers,
                              json={"review_status": "confirmed", "requirement_status": "evidenced"})
    assert marked.status_code == 200, marked.text
    assert marked.json()["review_evidence"] is None
    first_url = files + "/" + first.json()["id"]
    assert (await client.delete(first_url, headers=headers)).status_code == 204
    view = await client.get(f"/api/v1/review-cycles/{cycle.id}", headers=headers)
    row = next(item for item in view.json()["items"] if item["id"] == str(items[0].id))
    assert row["evidence_files"] == []
    assert row["evidence_changed_since_review"] is True
    assert all(b["item_id"] != str(items[0].id) for b in view.json()["readiness"]["blockers"])
    await db_session.refresh(items[0])
    assert items[0].assessment_status == "evidenced"
    assert items[0].assessment_rationale is None
    rows = await assessment_rows(db_session, cycle)
    assert rows[0].files == []
    assert rows[0].status_comment == ""


async def test_comment_attachments_survive_closure_with_original_bytes(client, db_session, assessment):
    _, cycle, items, headers, _ = assessment
    base = f"/api/v1/review-cycles/{cycle.id}"
    payloads = {f"attachment.{ext}": f"synthetic {ext}".encode()
                for ext in ("jpg", "jpeg", "png", "pdf", "docx", "txt", "md")}
    response = await client.post(
        base + f"/items/{items[0].id}/comments/attachments", headers=headers,
        data={"body": "Supporting documents"},
        files=[("files", (name, content, "application/octet-stream"))
               for name, content in payloads.items()],
    )
    assert response.status_code == 201, response.text
    comment_id = response.json()["id"]
    view = (await client.get(base, headers=headers)).json()
    row = next(item for item in view["items"] if item["id"] == str(items[0].id))
    assert {file["filename"] for file in row["evidence_files"]} == set(payloads)
    assert all(file["comment_id"] == comment_id for file in row["evidence_files"])
    assert (await finish(client, assessment)).status_code == 200
    view = (await client.get(base, headers=headers)).json()
    row = next(item for item in view["items"] if item["id"] == str(items[0].id))
    assert any(comment["id"] == comment_id for comment in row["comments"])
    from app.models import ReviewItemEvidenceFile
    for file in row["evidence_files"]:
        assert file["comment_id"] == comment_id
        stored = await db_session.get(ReviewItemEvidenceFile, uuid.UUID(file["id"]))
        Path(stored.file_path).write_bytes(b"changed live file")
        downloaded = await client.get(
            base + f"/items/{items[0].id}/files/{file['id']}/download", headers=headers,
        )
        assert downloaded.status_code == 200
        assert downloaded.content == payloads[file["filename"]]
    rejected = await client.post(
        base + f"/items/{items[0].id}/comments/attachments", headers=headers,
        data={"body": "Too late"}, files={"files": ("late.txt", b"late")},
    )
    assert rejected.status_code == 409


@pytest.mark.parametrize("failure", ["unsupported", "oversized", "empty", "too_many"])
async def test_failed_comment_upload_leaves_no_comment_or_files(
    client, db_session, assessment, monkeypatch, failure,
):
    from sqlalchemy import select
    from app.models import ReviewItemComment, ReviewItemEvidenceFile
    _, cycle, items, headers, _ = assessment
    monkeypatch.setattr(settings, "max_evidence_upload_mb", 1)
    files = [("files", ("first.txt", b"must be cleaned up"))]
    if failure == "unsupported":
        files.append(("files", ("script.html", b"<script>")))
    elif failure == "oversized":
        files.append(("files", ("huge.txt", b"x" * (1024 * 1024 + 1))))
    elif failure == "empty":
        files.append(("files", ("empty.txt", b"")))
    else:
        files *= 11
    response = await client.post(
        f"/api/v1/review-cycles/{cycle.id}/items/{items[0].id}/comments/attachments",
        headers=headers, data={"body": "Do not save a partial comment"}, files=files,
    )
    assert response.status_code == (413 if failure == "oversized" else 400), response.text
    assert not (await db_session.execute(select(ReviewItemComment))).scalars().all()
    assert not (await db_session.execute(select(ReviewItemEvidenceFile))).scalars().all()
    await db_session.refresh(items[0])
    assert items[0].review_comment is None
    folder = Path(settings.upload_dir) / "review-item-evidence"
    assert not folder.exists() or not list(folder.iterdir())


def test_comment_attachment_schema_upgrade_preserves_existing_files(tmp_path):
    import importlib.util
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    migration_path = Path(__file__).parents[1] / "alembic/versions/n20260930_comment_attachments.py"
    spec = importlib.util.spec_from_file_location("comment_attachments_migration", migration_path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = sa.create_engine(f"sqlite:///{tmp_path / 'migration.db'}")
    with engine.begin() as connection:
        connection.exec_driver_sql("PRAGMA foreign_keys=ON")
        connection.exec_driver_sql("CREATE TABLE review_item_comments (id CHAR(32) PRIMARY KEY)")
        connection.exec_driver_sql(
            "CREATE TABLE review_item_evidence_files (id CHAR(32) PRIMARY KEY, filename TEXT NOT NULL)"
        )
        file_id, comment_id = uuid.uuid4().hex, uuid.uuid4().hex
        connection.execute(sa.text("INSERT INTO review_item_evidence_files VALUES (:id, 'retained.pdf')"),
                           {"id": file_id})
        migration.op = Operations(MigrationContext.configure(connection))
        migration.upgrade()
        assert connection.exec_driver_sql(
            "SELECT filename, comment_id FROM review_item_evidence_files"
        ).one() == ("retained.pdf", None)
        connection.execute(sa.text("INSERT INTO review_item_comments VALUES (:id)"), {"id": comment_id})
        connection.execute(sa.text("UPDATE review_item_evidence_files SET comment_id=:id"), {"id": comment_id})
        connection.execute(sa.text("DELETE FROM review_item_comments WHERE id=:id"), {"id": comment_id})
        assert connection.exec_driver_sql("SELECT comment_id FROM review_item_evidence_files").scalar() is None
        migration.downgrade()
        assert connection.exec_driver_sql("SELECT filename FROM review_item_evidence_files").scalar() == "retained.pdf"
    engine.dispose()
